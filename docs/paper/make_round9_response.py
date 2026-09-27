"""Build the answers to the round-9 technical information request.

Reads every question in ``R9_review_Questions.docx`` in document order and
answers each underneath it twice: once in plain words, once as technical facts
with the artefact, file, function or commit that supports it. Numbers are read
from ``round9_analysis.json`` (no model calls) and earlier recorded artefacts;
claims about the data are asserted before they are printed.

    uv run --with tokenizers --with gguf python -m tests.eval.extraction.round9_analysis
    uv run python docs/paper/make_round9_response.py
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

PAPER = Path(__file__).resolve().parent
sys.path.insert(0, str(PAPER))

import docx  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from make_round2_response import AQUA, BLUE, GRID, INK, ORANGE, QUOTE, YELLOW, Doc, pct  # noqa: E402
from make_round3_response import GREY  # noqa: E402

ROOT = PAPER.parents[1]
RESULTS = ROOT / "tests" / "eval" / "extraction" / "results"
QUESTIONS = PAPER / "R9_review_Questions.docx"
FIG = PAPER / "figures" / "round9"
OUT = PAPER / "Round9_Technical_Response.docx"
FIELDS = ["domain", "modality", "prompt_level", "outcome"]
MODELS = ["llama3.2:latest", "qwen2.5:3b", "qwen2.5:7b", "qwen3:14b"]
SHORT = {"llama3.2:latest": "Llama 3.2", "qwen2.5:3b": "Qwen 2.5 3B", "qwen2.5:7b": "Qwen 2.5 7B", "qwen3:14b": "Qwen 3 14B"}
PLAIN_FILL, PLAIN_EDGE = "EAF7F1", "1BAF7A"


def load(name: str) -> Any:
    return json.loads((RESULTS / name).read_text())


def day(iso: str) -> str:
    return datetime.fromisoformat(iso).strftime("%-d %b %Y")


def frac(h: int, n: int, sparse: bool = False) -> str:
    if not n:
        return "0/0 (no opportunity)"
    return f"{h}/{n} ({100 * h / n:.1f}%)" + (" sparse" if sparse else "")


# ---------------------------------------------------------------------------
# Questions, verbatim and in order
# ---------------------------------------------------------------------------

def read_items() -> list[dict[str, Any]]:
    src = docx.Document(str(QUESTIONS))
    items: list[dict[str, Any]] = []
    sub = None
    for el in src.element.body.iterchildren():
        if el.tag.split("}")[1] != "p":
            continue
        p = Paragraph(el, src)
        t = p.text.strip()
        if not t:
            continue
        if p.style.name == "Heading 1":
            items.append({"title": t, "why": [], "questions": [], "other": []})
            sub = None
        elif not items:
            continue
        elif p.style.name == "Heading 2":
            sub = t
        elif p.style.name == "List Number":
            items[-1]["questions"].append(t)
        elif sub == "Why this information is needed":
            items[-1]["why"].append(t)
        else:
            items[-1]["other"].append((sub, t))
    return [i for i in items if i["title"].startswith("R9-")]


def question(b: Doc, n: int, text: str) -> None:
    para = b.doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.keep_with_next = True
    r = para.add_run(f"Q{n}. ")
    r.bold = True
    r.font.color.rgb = QUOTE
    b._runs(para, text, colour=QUOTE)


def plain(b: Doc, text: str) -> None:
    b.boxed("In plain words", [text], fill=PLAIN_FILL, edge=PLAIN_EDGE)


def facts(b: Doc, lines: list[str]) -> None:
    b.label("The facts")
    b.bullets(lines)


# ---------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------

def fig_tokens(pk: dict) -> Path:
    rows = sorted(pk["packets"].items(), key=lambda kv: kv[1]["tokens_full_input"])
    fig, ax = plt.subplots(figsize=(10, 3.9))
    for i, (c, v) in enumerate(rows):
        t = v["tokens_full_input"]
        ax.barh(i, t, color=BLUE, height=0.62)
        ax.barh(i, 8192 - t, left=t, color=YELLOW, height=0.62)
        ax.text(8192 + 60, i, f"{t:,} in · {8192 - t} left", va="center", fontsize=7.4)
    ax.axvline(8192, color=INK, lw=1)
    ax.text(8192, len(rows) - 0.2, " 8,192-token window", fontsize=7.6, va="bottom")
    ax.set_yticks(range(len(rows)), [c for c, _ in rows], fontsize=7.8)
    ax.set_xlim(0, 10400)
    ax.set_xlabel("tokens")
    ax.legend(handles=[plt.Rectangle((0, 0), 1, 1, color=BLUE), plt.Rectangle((0, 0), 1, 1, color=YELLOW)],
              labels=["model-visible input (system + packet)", "room left for the report"], frameon=False,
              fontsize=7.6, ncol=2, loc="upper center", bbox_to_anchor=(0.45, 1.13))
    ax.grid(axis="x", color=GRID, lw=0.6)
    ax.set_axisbelow(True)
    fig.tight_layout()
    p = FIG / "r9_fig1_reporter_tokens.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


def fig_gold_nulls(fu: dict) -> Path:
    fig, ax = plt.subplots(figsize=(9, 3.3))
    sets = [("main_batch_330", "330 main batch", BLUE), ("frozen_132", "132 frozen", GREY),
            ("overlap_excluded_121", "121 post-audit", AQUA)]
    for j, (k, lab, col) in enumerate(sets):
        xs = [i + (j - 1) * 0.27 for i in range(4)]
        vals = [fu[k]["gold_null_by_field"][f] for f in FIELDS]
        ax.bar(xs, vals, width=0.25, color=col, label=lab)
        for x, v in zip(xs, vals):
            ax.text(x, v + 1.5, str(v), ha="center", fontsize=7.4)
    ax.set_xticks(range(4), [f.replace("_", " ") for f in FIELDS])
    ax.set_ylabel("reference-null slots\n(chances to make an unsupported claim)")
    ax.legend(frameon=False, fontsize=8)
    ax.grid(axis="y", color=GRID, lw=0.6)
    ax.set_axisbelow(True)
    fig.tight_layout()
    p = FIG / "r9_fig2_gold_nulls.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


def fig_toy(toy: dict) -> Path:
    cols = {"play": BLUE, "request": ORANGE, "other": GREY, "response": AQUA}
    eps, sel = toy["episodes_in_order"], set(toy["selected_indices"])
    fig, ax = plt.subplots(figsize=(10, 1.9))
    for i, t in enumerate(eps):
        ax.add_patch(plt.Rectangle((i, 0), 0.9, 0.9, color=cols[t], alpha=1.0 if i in sel else 0.22))
        if i in sel:
            ax.text(i + 0.45, 0.45, "✓", ha="center", va="center", color="white", fontsize=9, weight="bold")
    ax.set_xlim(-0.2, len(eps) + 0.2)
    ax.set_ylim(-0.6, 1.4)
    ax.axis("off")
    ax.legend(handles=[plt.Rectangle((0, 0), 1, 1, color=c) for c in cols.values()], labels=list(cols),
              frameon=False, fontsize=7.6, ncol=4, loc="upper center", bbox_to_anchor=(0.5, 1.25))
    ax.text(0, -0.45, "episodes in transcript order →   (solid ✓ = kept, faded = not kept)", fontsize=7.6)
    fig.tight_layout()
    p = FIG / "r9_fig3_toy_selection.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


# ---------------------------------------------------------------------------
# R9-1
# ---------------------------------------------------------------------------

def r9_1(b: Doc, qs: list[str], ctx) -> None:
    ch = ctx["r9"]["r9_1_chronology"]
    batch, q3, cb, pf = ch["three_generator_batch"], ch["qwen3_14b"], ch["codebook_history"], ch["prompt_file_history"]
    assert batch["generated_at"] < cb["first_date"] or day(batch["generated_at"]) < day(cb["first_date"])
    assert pf["last_date"] < batch["generated_at"][:10] + "T99" and not ch["prompt_changed_after_batch"]
    assert not any(v for r in ch["codebook_rule_terms_in_prompt"].values() for v in r.values())
    assert ch["prototype_default_model"] == "llama3.2:latest"

    b.table("Chronology (R9-1 preferred format). Sources: git history; `round9_analysis.json` (r9_1_chronology).",
            ["Model / role", "Prompt version", "Tool / schema", "Generated", "Codebook rules in prompt?",
             "Current prototype same?", "Evidence"],
            [["Llama 3.2, Qwen 2.5 3B, Qwen 2.5 7B — the 3 generators (330 batch)", f"prompts.py {ch['prompt_version_constant']} "
              f"(last changed {day(pf['last_date'])}, {pf['last_commit']})", "EXTRACTION_TOOL_SCHEMA_OPENAI, same file",
              f"{day(batch['generated_at'])} ({batch['first_commit']})", "No (none of the 3)", "Same prompt file, unchanged",
              "results/annotation_batch.json"],
             ["Qwen 3 14B — 4th generator, added later", "same prompts.py 1.0.0", "same", f"{day(q3['augmented_at'])} "
              f"({q3['last_commit']})", "No", "Same", "results/annotation_batch_augmented.json"],
             ["Human codebook (the 3 clarifications)", "—", "—", f"{day(cb['first_date'])} ({cb['first_commit']})",
              "—", "—", "docs/annotation_codebook.md"],
             ["Prototype request path (today)", "same prompts.py 1.0.0", "same", "—", "No",
              f"Default model is {ch['prototype_default_model']}, not Qwen 2.5 7B", "src/backend/config.py"]],
            [1.5, 1.0, 0.9, 0.9, 0.8, 0.9, 1.1])

    question(b, 1, qs[0])
    plain(b, "Yes. The three models wrote their answers on 7 September 2026. The human codebook with the three extra "
             "rules was written on 15 September, eight days later. The models could not have seen it.")
    facts(b, [
        f"Generator outputs: `tests/eval/extraction/results/annotation_batch.json`, `generated_at` "
        f"{batch['generated_at']}, committed {batch['first_commit']} ({day(batch['first_date'])}).",
        f"Codebook: `docs/annotation_codebook.md`, first and only commit {cb['first_commit']} ({day(cb['first_date'])}).",
        f"Prompt: `src/backend/ml/extraction/prompts.py`, `PROMPT_VERSION = \"{ch['prompt_version_constant']}\"`, "
        f"last changed {pf['last_commit']} ({day(pf['last_date'])}); category lists in "
        f"`src/backend/schemas/observation_schema.py`, unchanged since {day(ch['observation_enum_history']['first_date'])}.",
        "The batch file does **not** record a prompt version per output (the replay code says so: "
        "`mars_branch_replay.py`, `prompt_version=\"replay-unrecorded\"`). The version is identified by chronology: the "
        "prompt file did not change between its last commit and the generation date.",
        "Limit: git shows committed code. It cannot rule out an uncommitted local edit on the day of generation.",
    ])

    question(b, 2, qs[1])
    plain(b, "No to all three. The prompt only lists the allowed values for each field and says \"use null if the "
             "text does not clearly support a value; never guess\". It says nothing about gestures versus handling "
             "objects, nothing about needing a target before calling an answer correct or incorrect, and nothing "
             "about prompts versus opportunities.")
    facts(b, [
        "(a) Communicative gesture vs object manipulation: **No.** The words \"communicative\", \"object\" and "
        "\"manipulat…\" do not occur in the system or user prompt.",
        "(b) Identifiable target for correct/incorrect: **No.** \"target\" does not occur.",
        "(c) Prompt vs antecedent/opportunity: **No.** \"antecedent\" and \"opportunity\" do not occur.",
        "What the prompt does say about these fields (verbatim): \"Use null for any field not clearly supported by the "
        "source text\"; \"NEVER fabricate or infer beyond what the text explicitly states\"; then the value lists, "
        "e.g. \"prompt_level: independent, gestural, verbal, model, partial_physical, full_physical, unknown\". Tool "
        "schema field descriptions: modality \"Communication modality used by the child.\", prompt_level \"Level of "
        "prompting provided.\", outcome \"Result of the child's behavioral attempt.\"",
        "Related fact: the codebook's `outcome = none` (\"no target\") was written in the sheet as the word \"none\", "
        "meaning no value, and is stored as null in the gold labels. The models' outcome list had no \"none\"; "
        "they could return null or a value such as \"unknown\".",
    ])

    question(b, 3, qs[2])
    plain(b, "Qwen 3 14B got exactly the same instructions and the same answer form as the other three. The only "
             "differences are the model itself and the date: it was run on 16–17 September, after the humans had "
             "finished annotating and after the codebook existed — but the codebook was not added to its prompt.")
    facts(b, [
        "Script: `tests/eval/extraction/add_generator.py`, which uses the same `OllamaExtractionProvider` "
        "(`src/backend/ml/extraction/ollama_llm.py`) as the original batch builder. That provider sends "
        "`build_system_prompt()`, `build_user_prompt_chat_segment()` and `EXTRACTION_TOOL_SCHEMA_OPENAI` from "
        "`prompts.py`, with a forced tool call; no temperature or other sampling option is set, for any generator.",
        f"`prompts.py` and `observation_schema.py` did not change between 7 and 17 September; Qwen 3 14B outputs "
        f"stamped `augmented_at` {q3['augmented_at']}, committed {q3['last_commit']}.",
        "Differences relevant to the scored fields: none in prompt, tool schema or allowed values. Procedural "
        "differences only: run after annotation; results written to a separate file; a resume step skips segments "
        "already recorded for that model.",
    ])

    question(b, 4, qs[3])
    plain(b, "The prototype's instructions are still the original ones; nothing from the codebook was added. One "
             "correction to the question: the prototype is not set up to use Qwen 2.5 7B. Its configured default "
             "model is Llama 3.2. Qwen 2.5 7B would be used only if someone changes that setting.")
    facts(b, [
        "`src/backend/config.py`: `ollama_model: str = \"llama3.2:latest\"`; `.env.example` has "
        "`OLLAMA_MODEL=llama3.2:latest`; there is no `.env` in the repository overriding it. "
        "`docs/paper/clarification_answers.md` (round-2 Q6) records the same.",
        "Whatever model is configured, the request path uses `OllamaExtractionProvider` with `prompts.py` "
        f"{ch['prompt_version_constant']}, unchanged since {pf['last_commit']} ({day(pf['last_date'])}).",
        "Factual differences between the preserved-evaluation prompt and today's prototype prompt: **none**.",
        "If the manuscript says the prototype runs Qwen 2.5 7B, that statement does not match the code.",
    ])

    question(b, 5, qs[4])
    plain(b, "Confirmed: no. There is no newer prompt, so nothing could have been recomputed with one.")
    facts(b, ["No commit touches `prompts.py` after 28 Aug 2026, and no experiment artefact records a different prompt."])


# ---------------------------------------------------------------------------
# R9-2
# ---------------------------------------------------------------------------

def r9_2(b: Doc, qs: list[str], ctx) -> None:
    pk = ctx["r9"]["r9_2_3_4_packets"]
    rg = ctx["rg"]
    errs = [e for p in rg["per_child"] for e in p["validation_errors"]]
    assert len(rg["per_child"]) == 11 and all(not p["validation_passed"] for p in rg["per_child"])
    assert all(len(p["validation_errors"]) == 2 for p in rg["per_child"])
    assert set(errs) == {"Abstained KPI 'cross_setting_comparison' not acknowledged in report",
                         "Abstained KPI 'cross_contributor_comparison' not acknowledged in report"}
    assert all(set(v["withheld_kpis"]) == {"cross_setting_comparison", "cross_contributor_comparison"}
               for v in pk["packets"].values())
    ex = pk["packets"]["brett"]["withheld_kpis"]

    question(b, 1, qs[0])
    plain(b, "The reporter got one fixed instruction text (quoted in full in the appendix) plus, as its message, "
             "the child's whole evidence packet written out as JSON. Nothing else was added.")
    facts(b, [
        "System prompt: `REPORTING_SYSTEM_PROMPT` in `src/backend/ml/reporting_client.py` at commit 9ecc330 "
        "(17 Sep 2026), the commit that recorded the 11 reports. Verbatim text in Appendix A.",
        "User message: `json.dumps(evidence_packet.model_dump(mode=\"json\"), indent=2)` — the packet only.",
    ])

    question(b, 2, qs[1])
    plain(b, "Not as strongly as that. The instruction says to keep abstention messages and, if a figure was withheld "
             "because of insufficient evidence, to say so. It does not say \"list every withheld analysis\". Also, "
             "the two analyses that were actually withheld were withheld because everything came from one setting "
             "and one type of contributor — not worded as \"insufficient evidence\". A separate line asks the "
             "writer to note when observations come from a single contributor or setting.")
    facts(b, [
        "Rule 3 (verbatim): \"Preserve all abstention messages. If a metric was withheld due to insufficient "
        "evidence, state that clearly.\"",
        "Rule 4: \"If evidence is insufficient or limited, explicitly state this. Do not convert limited evidence "
        "into confident conclusions.\"",
        "Under \"You SHOULD\": \"Explicitly state evidence limitations and gaps\"; \"Note when observations come from "
        "a single contributor or setting\".",
        f"Withheld reasons in the packets (example, brett): cross_setting_comparison — \"{ex['cross_setting_comparison']}\"; "
        f"cross_contributor_comparison — \"{ex['cross_contributor_comparison']}\".",
        "No instruction names the withheld KPIs or requires each one to be listed.",
    ])

    question(b, 3, qs[2])
    plain(b, "The answer form had no special place for withheld analyses. Each section had a title, text, a list of "
             "record IDs, and an optional \"abstention note\" that could be left empty.")
    facts(b, [
        "Schema `ReportingClient.REPORT_SCHEMA` (Appendix A), sent as Ollama's `format` (constrained decoding): object "
        "with required `sections` (array, any length); each item: `heading` string (required), `body` string "
        "(required), `evidence_refs` array of strings (required), `abstention_note` string or null (**optional**, at "
        "most one per section).",
        "**No** mandatory field or structure for withheld analyses. The four required section names appear in the "
        "prompt text only; the schema does not enforce them.",
    ])

    question(b, 4, qs[3])
    plain(b, "The checker is a simple word search in the section text. A withheld analysis counted as mentioned if "
             "any section's main text contained its name written with spaces (\"cross setting comparison\"), or the "
             "word \"withheld\", or the word \"insufficient\". It did not look in the titles or in the abstention-note "
             "field. So a single use of \"withheld\" anywhere would have passed both; a mention written as "
             "\"cross-setting\", or placed only in the abstention note, would not have counted.")
    facts(b, [
        "Function `validate_report()` in `src/backend/ml/reporting_client.py`, logic unchanged since 75eb1d5 "
        "(25 Aug 2026). For each KPI with `abstained` true and a reason: acknowledged if, for any section, "
        "`kpi.name.replace(\"_\", \" \")` or `\"withheld\"` or `\"insufficient\"` is a substring of `section.body.lower()`.",
        "Not examined: `heading`, `abstention_note`, `evidence_refs`. Hyphenated or underscored names do not match.",
        "The same function also checks the prohibited-content patterns and, only when adequacy is INSUFFICIENT, an "
        "acknowledgement of insufficiency (all 11 packets were LIMITED, so that check did not apply).",
    ])

    question(b, 5, qs[4])
    plain(b, "Yes. All 11 reports failed for the same two reasons and nothing else: they did not mention the "
             "withheld cross-setting comparison or the withheld cross-contributor comparison in a way the checker "
             "recognises. Those two were withheld in every one of the 11 packets. We cannot re-read the reports to "
             "see what they did say, because the report text itself was not saved.")
    facts(b, [
        "`tests/eval/reporting/results/reporting_groundedness.json` @ 9ecc330: 11/11 `validation_passed: false`; each "
        "has exactly the two errors \"Abstained KPI 'cross_setting_comparison' not acknowledged in report\" and "
        "\"… 'cross_contributor_comparison' …\". No prohibited-content error in any report.",
        "Withheld KPI families expected in all 11 packets: cross_setting_comparison and cross_contributor_comparison "
        "(every child's records come from one corpus, hence one setting and one contributor role).",
        "The result file stores scores, not the generated text; which words the model used cannot be recovered.",
    ])

    question(b, 6, qs[5])
    plain(b, "Everything is in two code files and one results file, all at the same commit.")
    b.table("Artefacts for the recorded 11-report run.", ["What", "Path", "Commit"], [
        ["Prompt, schema, validator", "src/backend/ml/reporting_client.py", "9ecc330 (validator logic since 75eb1d5)"],
        ["Context setting (num_ctx 8192)", "src/backend/config.py", "9ecc330"],
        ["Packet builder and scorer", "tests/eval/reporting/groundedness.py", "9ecc330"],
        ["Recorded output", "tests/eval/reporting/results/reporting_groundedness.json", "9ecc330"],
    ], [1.8, 3.2, 1.7])


# ---------------------------------------------------------------------------
# R9-3
# ---------------------------------------------------------------------------

def r9_3(b: Doc, qs: list[str], ctx) -> None:
    pk = ctx["r9"]["r9_2_3_4_packets"]
    ts, om = pk["token_summary"], pk["ollama_model_info"]
    rows = pk["packets"]
    assert all(v["tokens_full_input"] < 8192 for v in rows.values())
    assert all(v["withheld_names_present_in_input"] for v in rows.values())
    room = {c: 8192 - v["tokens_full_input"] for c, v in rows.items()}

    question(b, 1, qs[0])
    plain(b, "The exact counts the server saw were not saved. We recounted them with the same tokenizer the model "
             "uses (checked to be identical to the model file on this machine), on the exact text the server builds "
             "from the instruction and the packet. Each input is about 7,400–8,000 tokens.")
    facts(b, [
        f"Method: {pk['tokenizer']['tokenizer']}; {pk['tokenizer']['check']}. Text counted = the qwen2.5 Ollama chat "
        "template applied to the system prompt and the packet JSON (no tools), including its special tokens.",
        "Not included: the JSON schema passed as `format`. The model's chat template has no slot for it "
        f"(template_has_format_placeholder = {om['template_has_format_placeholder']}); whether the server adds it "
        "elsewhere cannot be verified from preserved artefacts.",
        "API-reported counts (`prompt_eval_count`) were not recorded.",
    ])
    b.table("Model-visible input per child and room left in the 8,192-token window. Source: `round9_analysis.json` "
            "(r9_2_3_4_packets.packets).",
            ["Child", "Input tokens", "Left for output", "Packet characters"],
            [[c, f"{v['tokens_full_input']:,}", str(room[c]), f"{v['chars_user']:,}"]
             for c, v in sorted(rows.items(), key=lambda kv: kv[1]["tokens_full_input"])],
            [1.8, 1.4, 1.4, 1.6])

    question(b, 2, qs[1])
    plain(b, f"Smallest {ts['min']:,}, middle {ts['median']:,}, largest {ts['max']:,} tokens ({ts['max_child']}).")
    facts(b, [f"Minimum {ts['min']:,}; median {ts['median']:,}; maximum {ts['max']:,} ({ts['max_child']})."])
    b.figure(fig_tokens(pk), "Each packet nearly fills the 8,192-token window; the yellow part is all the room the "
             "model had to write its report.")

    question(b, 3, qs[2])
    plain(b, "The whole packet, including both withheld-analysis entries, fits inside the window for all 11 children, "
             "so the input was not cut when it was sent. But see the next answer: the room left for the report itself "
             "was small, and we cannot show from what was saved that the report stayed within it.")
    facts(b, [
        "Input < 8,192 tokens for all 11; both withheld KPI names are present in every serialized packet.",
        "What the server does when a conversation plus its reply outgrows the window depends on the Ollama version, "
        f"which was not recorded for the run (installed now: {om['ollama_version_now']}). We do not state that behaviour.",
    ])

    question(b, 4, qs[3])
    plain(b, f"No limit was set on the length of the report. After the input, only {min(room.values())} to "
             f"{max(room.values())} tokens of the window remained ({min(room, key=room.get)} had the least). The "
             "length of the reports was not saved, so we cannot show that input plus report fitted. This is a gap in "
             "the record, not something we can close without re-running.")
    facts(b, [
        "No `num_predict` (output limit) was sent; `options` contained only `num_ctx: 8192`.",
        f"Remaining window per child: {', '.join(f'{c} {r}' for c, r in sorted(room.items(), key=lambda kv: kv[1]))}.",
        "Output token counts and report text were not preserved.",
    ])

    question(b, 5, qs[4])
    plain(b, "Only a few settings were set on purpose: the model, the answer form, the 8,192 window, and one "
             "report per child. Randomness settings were left at the server's defaults and no seed was fixed, so a "
             "re-run would not reproduce the same text exactly.")
    b.table("Reporter call as executed (`_call_llm_ollama_native`, reporting_client.py @ 9ecc330).",
            ["Setting", "Value"], [
                ["Model", "qwen2.5:7b (passed by groundedness.py, DEFAULT_MODEL)"],
                ["Endpoint", "Ollama native /api/chat; on an exception it falls back to the /v1 endpoint, which does "
                             "not enforce the schema. Whether that fallback ever fired was not recorded"],
                ["Structured output", "format = REPORT_SCHEMA (JSON schema)"],
                ["Context length", "num_ctx 8192"], ["Stream", "false"],
                ["Generations per child", "1 (template fallback only if the call raised; 0 fallbacks in the run)"],
                ["Timeout", "600 s"],
                ["temperature, top_p, top_k, num_predict, repeat_penalty, seed", "not set"],
            ], [2.2, 4.5])

    question(b, 6, qs[5])
    plain(b, "Everything not set came from the server's built-in defaults. The model file adds none of its own. We "
             "do not list the default values, because they depend on the server version, which was not recorded.")
    facts(b, [f"`/api/show` for qwen2.5:7b: Modelfile parameters = {om['modelfile_parameters']} (none set)."])

    question(b, 7, qs[6])
    plain(b, "The model file on this machine was downloaded on 7 September, before the run, and has not changed "
             "since, so it is almost certainly the one used. The server version at the time is unknown.")
    facts(b, [
        f"qwen2.5:7b digest {om['digest'][:12]}… ({om['quantization']}), local copy modified {om['modified_at'][:19]} "
        "(before the 16–17 Sep run). Run-time digest was not written into the result file.",
        f"Ollama version at run time: not recorded (now {om['ollama_version_now']}).",
        "Prompt/schema: reporting_client.py @ 9ecc330. Output: reporting_groundedness.json @ 9ecc330.",
    ])


# ---------------------------------------------------------------------------
# R9-4
# ---------------------------------------------------------------------------

def r9_4(b: Doc, qs: list[str], ctx) -> None:
    rows = ctx["r9"]["r9_2_3_4_packets"]["packets"]
    assert all(v["goals"] == [] for v in rows.values())
    assert all(v["adequacy_level"] == "limited" for v in rows.values())
    checks = list(next(iter(rows.values()))["adequacy_checks"])
    tally = {k: sum(v["adequacy_checks"][k]["passed"] for v in rows.values()) for k in checks}
    assert tally["goal_coverage"] == 11 and tally["contributor_representation"] == 0 and tally["setting_representation"] == 0

    question(b, 1, qs[0])
    plain(b, "There were no goals. Every one of the 11 packets had an empty goal list.")
    facts(b, ["`tests/eval/reporting/groundedness.py`, `build_packet()`: `goals=[]` for every child; the router is called "
              "with `goals=[]`."])
    question(b, 2, qs[1])
    plain(b, "Not applicable — no goal was created, so there is no goal to show.")
    facts(b, ["No goal object exists in any packet."])
    question(b, 3, qs[2])
    plain(b, "No record was linked to any goal. With no goals, the goal check simply passes and says \"No active goals "
             "defined\"; nothing is counted.")
    facts(b, ["`AdequacyRouter._check_goal_coverage()` (`src/backend/ml/adequacy_router.py`): `if not goals: return "
              "passed=True, detail=\"No active goals defined\"`. (When goals exist, the rule is: goal domains that "
              "appear among record domains ÷ goal domains ≥ 0.5.)"])
    question(b, 4, qs[3])
    plain(b, "Every child's review period starts on 5 January 2026. Each session of that child is given a pretend date "
             "one week after the previous one, and the review period ends on the last of those dates. The same rule "
             "was used for all 11 children.")
    facts(b, [
        "`load_records()`: sessions of a child's 30 records sorted by session ID (text order); session i gets "
        "`2026-01-05 + 7·i days`; every record takes its session's date.",
        "`build_packet()`: review start = earliest record date, end = latest record date.",
        "Resulting ends: " + "; ".join(f"{c} {v['review_period'][1]} ({v['n_sessions']} sessions)" for c, v in sorted(rows.items())),
    ])
    question(b, 5, qs[4])
    plain(b, "No. Every packet failed only the two known checks (one setting, one contributor type) and passed all "
             "others. The goal check could not fail because there were no goals. Two failed checks are enough for "
             "LIMITED.")
    b.table("Adequacy checks across the 11 packets. Source: `round9_analysis.json` (adequacy_checks).",
            ["Check", "Passed", "Example detail (brett)"],
            [[k.replace("_", " "), f"{tally[k]}/11", rows["brett"]["adequacy_checks"][k]["detail"]] for k in checks],
            [1.9, 0.7, 4.1])
    facts(b, ["`_compute_level()`: any failed check prevents ADEQUATE; no critical check (coverage, field completeness) "
              "failed, so the level is LIMITED. The temporal-distribution check uses the pretend weekly dates."])
    question(b, 6, qs[5])
    plain(b, "All of it is built in one test script, using the real adequacy and analytics code.")
    facts(b, ["`tests/eval/reporting/groundedness.py` — `load_records()` (dates, setting, contributor role), "
              "`build_packet()` (goals, review period); `src/backend/ml/adequacy_router.py`; thresholds in "
              "`src/backend/config.py` (`adequacy_*`)."])


# ---------------------------------------------------------------------------
# R9-5
# ---------------------------------------------------------------------------

def r9_5(b: Doc, qs: list[str], ctx) -> None:
    fu = ctx["r9"]["r9_5_field_unsupported"]
    g = {k: fu[k]["gold_null_by_field"] for k in fu}
    assert g["overlap_excluded_121"]["domain"] == 0 and g["overlap_excluded_121"]["prompt_level"] == 0

    question(b, 1, qs[0])
    plain(b, "An \"opportunity\" is a field where the human reference says \"nothing here\". Almost all of them are "
             "in modality and outcome. Domain and prompt level have almost none — and none at all in the 121 set.")
    b.table("Reference-null opportunities by field. Source: `round9_analysis.json` (r9_5_field_unsupported).",
            ["Set", "domain", "modality", "prompt_level", "outcome", "Total"],
            [[lab, *[str(g[k][f]) for f in FIELDS], str(sum(g[k].values()))]
             for k, lab in (("main_batch_330", "330 main batch"), ("frozen_132", "132 frozen"), ("overlap_excluded_121", "121 post-audit"))],
            [1.6, 0.9, 0.9, 1.0, 0.9, 0.8])
    b.figure(fig_gold_nulls(fu), "Where the chances to make an unsupported claim are: almost entirely modality and outcome.")

    def per_field_table(key: str, lab: str) -> None:
        pm = fu[key]["per_model"]
        b.table(f"{lab}: unsupported assertions ÷ opportunities (%), per field and pooled. 'sparse' = fewer than 20 "
                "opportunities (our flag, percentages shown anyway).",
                ["Generator", *FIELDS, "Pooled", "Of which 'unknown'"],
                [[SHORT[m], *[frac(pm[m]["per_field"][f]["hallucinated"], pm[m]["per_field"][f]["gold_null"],
                                   pm[m]["per_field"][f]["sparse"] and pm[m]["per_field"][f]["gold_null"] > 0) for f in FIELDS],
                  frac(pm[m]["pooled"]["hallucinated"], pm[m]["pooled"]["gold_null"]),
                  str(sum(pm[m]["per_field"][f]["hallucinated_value_unknown"] for f in FIELDS))] for m in MODELS],
                [0.9, 0.95, 1.0, 0.95, 1.0, 1.0, 0.7])

    question(b, 2, qs[1])
    plain(b, "The table shows, for each model and field, how many times it filled in a value where the reference had "
             "none, out of how many chances.")
    per_field_table("main_batch_330", "330 main batch")
    question(b, 3, qs[2])
    plain(b, "Same table for the 121 set. Domain and prompt level have no chances at all here, so the 121 unsupported "
             "rate is really a modality-and-outcome rate.")
    per_field_table("overlap_excluded_121", "121 post-audit set")
    per_field_table("frozen_132", "132 frozen set (for reference)")

    question(b, 4, qs[3])
    pm = fu["main_batch_330"]["per_model"]
    plain(b, "Confirmed. The reported figure adds up all unsupported claims across the four fields and divides by all "
             "chances across the four fields. It is not the average of four percentages; the two can differ a lot "
             f"(Llama 3.2 on the 330: {100 * pm['llama3.2:latest']['pooled']['rate']:.1f}% pooled vs "
             f"{100 * pm['llama3.2:latest']['unweighted_mean_of_fields']:.1f}% as an average of fields).")
    facts(b, ["`unsupported()` in `tests/eval/extraction/round3_analysis.py`: hallucinated ÷ (hallucinated + "
              "correctly_null) summed over all four fields. The pooled values here reproduce `round4_analysis.json` "
              "exactly (asserted)."])
    question(b, 5, qs[4])
    plain(b, "Done: every field keeps its raw count, and fields with fewer than 20 chances are marked 'sparse' instead "
             "of being hidden.")
    facts(b, ["Sparse flag: fewer than 20 opportunities. The threshold is ours, for flagging only; change it if you prefer."])
    question(b, 6, qs[5])
    plain(b, "The counts come from files already on disk; no model was run.")
    facts(b, [
        "330: `results/scored_against_gold_augmented.json` (made by `score_against_gold.py`); 132/121: "
        "`results/scored_holdout.json`; counting script `tests/eval/extraction/round9_analysis.py`.",
        "Scoring rule (`judge_field()`): reference empty + model empty = correctly_null; reference empty + model has any "
        "value = unsupported (hallucinated). The value \"unknown\" counts as a value; the last column shows how many "
        "of each model's unsupported assertions were literally \"unknown\".",
    ])


# ---------------------------------------------------------------------------
# R9-7
# ---------------------------------------------------------------------------

def r9_7(b: Doc, qs: list[str], ctx) -> None:
    sel = ctx["r9"]["r9_7_selection"]
    c273, c455 = sel["caps"]["3000"], sel["caps"]["5000"]
    toy = sel["toy_example"]
    trunc = [c for c, v in c273["per_child"].items() if v["truncated_by_cap"]]
    below = {k: sum(v["n_selected"] < sel["caps"][k]["per_child_cap"] for v in sel["caps"][k]["per_child"].values()) for k in ("3000", "5000")}

    question(b, 1, qs[0])
    plain(b, "For each child: count how many kinds of episode the child has. Give each kind the same quota (the cap "
             "divided by the number of kinds, rounded up). Take the earliest episodes of each kind up to its quota. "
             "If that adds up to more than the cap, cut the extra off the end. That's all — nothing is moved from one "
             "kind to another.")
    b.code([
        "select_segments(all_episodes, max_segments):       # 3000 -> 2,442 universe; 5000 -> 3,371",
        "    group episodes by child, keeping pipeline order",
        "    C = ceil(max_segments / number_of_children)     # 11 children -> 273 or 455",
        "    for each child: selected += select_for_child(child_episodes, C)",
        "    if len(selected) > max_segments: keep the first max_segments   # never triggered here",
        "",
        "select_for_child(E, C):",
        "    types = episode types present in E, in order of first appearance",
        "    q = ceil(C / len(types))",
        "    S = []",
        "    for t in types:  S += first min(q, count(t)) episodes of type t, in pipeline order",
        "    if len(S) > C:   S = S[:C]                      # drops from the end (last type)",
        "    return S",
        "",
        "pipeline order = sessions in sorted file-path order, then transcript order within a session",
    ])
    facts(b, ["`data/pipeline/select_segments.py`: `select_segments()`, `_select_for_child()`; called with "
              "`max_segments=3000` and `5000` (`tests/eval/extraction/round4_analysis.py::universe`)."])

    question(b, 2, qs[1])
    plain(b, "Quota = cap ÷ number of kinds this child actually has, rounded up. Kinds the child does not have are "
             "not counted.")
    facts(b, ["`per_type = ceil(C / num_types)`, `num_types = len(by_type)` for that child's episodes only. Flusberg "
              "children have 2 or 3 types (play, other, sometimes initiation); Rollins children have 7."])

    question(b, 3, qs[2])
    plain(b, "There is no remainder step. Because quotas are rounded up they can add up to more than the cap; the "
             "extra is simply cut from the end of the list, i.e. from the kind that appeared last in the transcripts.")
    facts(b, [f"At cap 273 this cut happened for {', '.join(trunc)} (274 → 273, one episode removed from the "
              "second type). At cap 455 it never happened."])

    question(b, 4, qs[3])
    plain(b, "No. If a kind has fewer episodes than its quota, the spare places are lost, not given to other kinds. "
             f"That is why most children end up below the cap ({below['3000']} of 11 at 273, {below['5000']} of 11 at 455).")
    facts(b, [f"Example brett at cap 455: quota {c455['per_child']['brett']['quota_per_type']}; play has "
              f"{c455['per_child']['brett']['available']['play']} (all taken), other {c455['per_child']['brett']['selected']['other']} "
              f"→ {c455['per_child']['brett']['n_selected']} kept; the unused places are not reallocated."])

    question(b, 5, qs[4])
    plain(b, "Within a kind, episodes are taken in the order they occur: first by session file name, then by position "
             "in the transcript. There is no separate sorting key, so there are no ties to break.")
    facts(b, ["Sessions come from `sorted(corpus_dir.glob(\"**/*.cha\"))` (`chat_parser.parse_corpus_directory`); "
              "episodes within a session in transcript order (`segment_episodes`). Python list slicing keeps that order."])

    question(b, 6, qs[5])
    plain(b, "A child ends up below the cap whenever their kinds, each capped at the quota, do not add up to the cap.")
    b.table("Per child: quota, selected and available by type. Source: `round9_analysis.json` (r9_7_selection).",
            ["Child", "Cap 273: quota → selected", "Cap 455: quota → selected", "Available by type"],
            [[c, f"{v['quota_per_type']} → {v['n_selected']}" + (" (cut)" if v["truncated_by_cap"] else ""),
              f"{c455['per_child'][c]['quota_per_type']} → {c455['per_child'][c]['n_selected']}",
              ", ".join(f"{t} {n}" for t, n in v["available"].items())] for c, v in c273["per_child"].items()],
            [1.2, 1.4, 1.4, 2.7])

    question(b, 7, qs[6])
    plain(b, "Inside each child, episodes of an under-sampled kind count a bit more and episodes of an over-sampled "
             "kind count a bit less, so that the child's 121-set episodes match the kind mix of the pool they came "
             "from. Each child's total weight stays the same, so children are not reweighted against each other.")
    b.code([
        "for each child c:",
        "    N[c,k] = episodes of type k in the frame (analysed or intended frame)",
        "    n[c,k] = episodes of type k among c's episodes in the 121",
        "    covered = sum of N[c,k] over types k that have at least one sampled episode",
        "    weight of each sampled episode of type k:  w = (N[c,k] / covered) / (n[c,k] / n[c])",
        "    # weights within a child sum to n[c]; no normalisation across children",
        "",
        "merged      = sum_x w_x * (acceptable fields of x / 4)            / sum_x w_x",
        "whole       = sum_x w_x * [all 4 fields acceptable]              / sum_x w_x",
        "unsupported = sum_x w_x * hallucinated_x / sum_x w_x * (hallucinated_x + correctly_null_x)",
        "reported frame effect = estimate(intended frame) - estimate(analysed frame)",
    ])
    facts(b, ["`tests/eval/extraction/round4_analysis.py`, `same_universe()` → `poststrat()`; two type groupings: the "
              "7 episode types, and 'other' vs any typed. Frame types with no sampled episode are dropped and their "
              "share reported as `unaccounted_intended_share_by_child`. Unit = episode."])

    question(b, 8, qs[7])
    plain(b, f"Toy child: {len(toy['episodes_in_order'])} episodes of 4 kinds, cap {toy['cap']}. Quota = ceil(10/4) = 3. "
             "Play gives 3, request only has 2, other gives 3, response gives 3 — 11 in total, one too many, so the "
             "last one is cut. Result: 10 episodes.")
    b.figure(fig_toy(toy), "Toy example run through the real `_select_for_child` function (no transcript content).")
    facts(b, [f"Selected positions: {toy['selected_indices']}; types {toy['selected_types']}."])


# ---------------------------------------------------------------------------
# R9-8
# ---------------------------------------------------------------------------

def r9_8(b: Doc, qs: list[str], ctx) -> None:
    o = ctx["r9"]["r9_8_oracle"]
    base = o["three_generators_all_outputs_330"]
    assert (base["numerator"], base["denominator"]) == (20, 330)
    st = o["status_counts_3gen"]

    question(b, 1, qs[0])
    plain(b, "All 330 main-batch segments.")
    facts(b, ["Denominator = 330 (every segment with an adjudicated reference)."])
    question(b, 2, qs[1])
    plain(b, "Every saved answer from the three original models, whatever its status — including ones the model "
             "itself marked as abstained. No filtering at all.")
    facts(b, [f"All 3 × 330 = 990 preserved outputs: success {st['success']}, abstained {st['abstained']}, partial "
              f"{st['partial']}. Scoring reads each output's `extracted_fields` regardless of status."])
    question(b, 3, qs[2])
    plain(b, "No. The MARS entry rules were not applied, and duplicates were not merged (merging duplicates would not "
             "change whether some answer was right anyway).")
    facts(b, ["`oracle_ceiling()` in `tests/eval/extraction/mars_evaluation.py` iterates over the model list directly; it "
              "does not call `build_candidates()`."])
    question(b, 4, qs[3])
    plain(b, "Yes. A segment counts if at least one model got all four fields right, where \"right\" includes leaving a "
             "field empty when the reference is empty. The free-text behaviour field is not part of it.")
    facts(b, ["`is_fully_correct()`: every one of domain, modality, prompt_level, outcome judged `correct` or "
              "`correctly_null`. Values compared after lower-casing and collapsing spaces (`agreement_report.normalize`)."])
    question(b, 5, qs[4])
    plain(b, f"{base['numerator']} of {base['denominator']} segments = {pct(base['rate'])}.")
    facts(b, ["Computed by `oracle_ceiling()` (`mars_evaluation.py`), recorded as `oracle_ceiling` "
              f"{o['recorded_mars_evaluation']:.4f} in `results/mars_evaluation.json`; input `results/scored_against_gold.json`."])
    question(b, 6, qs[5])
    plain(b, "Here are the other ways it could be defined, with their numbers. The 6.1% in the paper is the first row.")
    b.table("Oracle ceiling under each candidate definition. Source: `round9_analysis.json` (r9_8_oracle).",
            ["Definition", "Numerator / denominator", "Rate"], [
                ["3 generators, all outputs, all 330 segments (the reported 6.1%)", f"{base['numerator']}/{base['denominator']}", pct(base["rate"])],
                ["3 generators, MARS-reviewable outputs only (success/partial; behavior or domain present)",
                 f"{o['three_generators_reviewable_only_330']['numerator']}/330", pct(o["three_generators_reviewable_only_330"]["rate"])],
                ["3 generators, all outputs, only segments MARS returned a record for",
                 f"{o['three_generators_all_outputs_mars_returned']['numerator']}/{o['three_generators_all_outputs_mars_returned']['denominator']}",
                 pct(o["three_generators_all_outputs_mars_returned"]["rate"])],
                ["3 generators, all outputs, only rerank-required segments (>1 distinct candidate)",
                 f"{o['three_generators_all_outputs_rerank_required']['numerator']}/{o['three_generators_all_outputs_rerank_required']['denominator']}",
                 pct(o["three_generators_all_outputs_rerank_required"]["rate"])],
                ["4 generators incl. Qwen 3 14B, all outputs, all 330", f"{o['four_generators_all_outputs_330']['numerator']}/330",
                 pct(o["four_generators_all_outputs_330"]["rate"])],
            ], [4.2, 1.4, 0.9])


def r9_6(b: Doc, _qs, _ctx) -> None:
    plain(b, "No action taken, as requested. We found nothing that contradicts the stated protocol. For completeness: "
             "the annotation sheet showed four identifier columns next to each transcript.")
    facts(b, ["`tests/eval/extraction/export_gold_sheet.py`: `CONTEXT_COLUMNS = [\"segment_id\", \"child_id\", \"corpus\", "
              "\"session_id\", \"transcript\"]`. The session ID contains the source file name. The models received only "
              "the rendered transcript turns (`render_segment_to_chat_text`)."])


ANSWERS: dict[str, Callable] = {"R9-1": r9_1, "R9-2": r9_2, "R9-3": r9_3, "R9-4": r9_4, "R9-5": r9_5,
                                "R9-7": r9_7, "R9-8": r9_8, "R9-6": r9_6}


def build() -> Path:
    FIG.mkdir(parents=True, exist_ok=True)
    ctx = {"r9": load("round9_analysis.json"), "rg": json.loads((ROOT / "tests/eval/reporting/results/reporting_groundedness.json").read_text())}
    items = read_items()
    ids = [i["title"].split()[0] for i in items]
    assert ids == ["R9-1", "R9-2", "R9-3", "R9-4", "R9-5", "R9-7", "R9-8", "R9-6"], ids
    expected_q = {"R9-1": 5, "R9-2": 6, "R9-3": 7, "R9-4": 6, "R9-5": 6, "R9-7": 8, "R9-8": 6, "R9-6": 0}
    assert all(len(i["questions"]) == expected_q[i["title"].split()[0]] for i in items)

    pk = ctx["r9"]["r9_2_3_4_packets"]
    b = Doc()
    grey = RGBColor(0x52, 0x51, 0x4E)
    b.doc.add_heading("Round 9 — Answers to the Technical Information Request", level=0)
    b.p("Every question in 'R9_review_Questions.docx', in its original order and wording, answered twice: first "
        "**in plain words**, then as **the facts**, with the file, function, commit or artefact behind each one. "
        "Facts only; no manuscript wording. Branch model-finetuning, built 2026-09-27. Internal working document.",
        colour=grey)
    b.p("No model was run. Numbers come from `round9_analysis.json` (computed from saved outputs, the deterministic "
        "pipeline and git history) and earlier recorded artefacts, and are checked by assertion when this document is "
        "built. Where something cannot be established from what was saved, the answer says so.", colour=grey)

    b.h("At a glance", 1)
    b.table("One line per item.", ["Item", "Plain-language answer"], [
        ["R9-1", "Models ran on 7 Sep (Qwen 3 14B on 16–17 Sep); codebook written 15 Sep; the prompt never contained the "
                 "codebook rules and has not changed since 28 Aug."],
        ["R9-2", "The reporter was told to preserve abstention messages, not to list each withheld analysis; the answer "
                 "form had no required place for it; the checker was a word search in section text."],
        ["R9-3", f"Inputs were {pk['token_summary']['min']:,}–{pk['token_summary']['max']:,} tokens in an 8,192 window; "
                 "input fitted, but report length was not saved, so the fit of input + report cannot be shown."],
        ["R9-4", "No goals at all; pretend weekly dates from 5 Jan 2026; only the two known checks failed."],
        ["R9-5", "Unsupported-claim chances are almost all in modality and outcome; per-field tables given."],
        ["R9-7", "Equal quota per kind (rounded up), earliest first, cut the excess, no redistribution; exact pseudocode "
                 "and reweighting formula given."],
        ["R9-8", "6.1% = 20/330: at least one of the three models fully right, over all saved answers, no filtering."],
        ["R9-6", "No action; nothing contradicts the protocol."],
    ], [0.6, 6.1])
    b.label("Things you should know before writing (facts that may affect the manuscript)")
    b.bullets([
        "**Prototype model:** the configured default extraction model is Llama 3.2, not Qwen 2.5 7B (R9-1 Q4).",
        "**Reporter window:** only 183–776 tokens were left for each report and no output limit was set; report length "
        "and text were not saved (R9-3 Q3–Q4).",
        "**Report text not preserved:** the 11 reports cannot be re-read; only their scores exist (R9-2 Q5).",
        "**Validator scope:** it searched only section body text, not the abstention-note field (R9-2 Q4).",
        "**121 set:** domain and prompt level have zero reference-null opportunities (R9-5 Q1).",
    ])

    for it in items:
        b.page_break()
        b.h(it["title"], 1)
        if it["why"]:
            b.label("Why they asked")
            for t in it["why"]:
                b.p(t, italic=True)
        if it["title"].startswith("R9-6"):
            for _s, t in it["other"]:
                b.p(t, italic=True)
        ANSWERS[it["title"].split()[0]](b, it["questions"], ctx)

    b.page_break()
    b.h("Appendix A — Reporter prompt and schema, verbatim (reporting_client.py @ 9ecc330)", 1)
    from src.backend.ml.reporting_client import REPORTING_SYSTEM_PROMPT, ReportingClient
    b.code(REPORTING_SYSTEM_PROMPT.splitlines())
    b.code(json.dumps(ReportingClient.REPORT_SCHEMA, indent=2).splitlines())
    b.h("How this document was produced", 1)
    b.code(["uv run --with tokenizers --with gguf python -m tests.eval.extraction.round9_analysis   # no model calls",
            "uv run python docs/paper/make_round9_response.py"])
    b.doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    sys.path.insert(0, str(ROOT))
    print(build())
