"""Build the response to the MDPI round-4 review comments.

Reads the nine comments from ``MDPI_Round4_Review_Comments.docx`` in document
order (P-1, P-2, P-3, R2-1 ... R2-6) and answers each underneath, verbatim
question first. Every number is read from a recorded artefact, chiefly
``round4_analysis.json`` (added this round, no model calls), and each table
names its source. Statements about the data are asserted before they are
printed, so a rebuild fails rather than prints a claim the data contradict.

    uv run python -m tests.eval.extraction.round4_analysis
    uv run python docs/paper/make_round4_response.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

PAPER = Path(__file__).resolve().parent
sys.path.insert(0, str(PAPER))

import docx  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from make_round2_response import (  # noqa: E402
    AQUA, BLUE, COLOUR, GRID, INK, INK2, MODELS, ORANGE, QUOTE, YELLOW, Doc, pct, short,
)
from make_round3_response import GREY, _arrow, _box  # noqa: E402

ROOT = PAPER.parents[1]
RESULTS = ROOT / "tests" / "eval" / "extraction" / "results"
QUESTIONS = PAPER / "MDPI_Round4_Review_Comments.docx"
FIG = PAPER / "figures" / "round4"
OUT = PAPER / "Round4_Review_Response.docx"
FIELDS = ["domain", "modality", "prompt_level", "outcome"]
GEN3 = ["llama3.2:latest", "qwen2.5:3b", "qwen2.5:7b"]


def load(name: str) -> Any:
    return json.loads((RESULTS / name).read_text())


def pp(x: float) -> str:
    return f"{100 * x:+.1f}"


def ci_pp(d: dict) -> str:
    return f"{pp(d['value'])} [{pp(d['lower'])}, {pp(d['upper'])}]" + (" *" if d.get("excludes_zero") else "")


def val_ci(d: dict) -> str:
    return f"{pct(d['value'])} [{pct(d['ci95']['lower'])}, {pct(d['ci95']['upper'])}]"


# ---------------------------------------------------------------------------
# Comments, verbatim and in order
# ---------------------------------------------------------------------------

def read_comments() -> list[dict[str, Any]]:
    src = docx.Document(str(QUESTIONS))
    comments: list[dict[str, Any]] = []
    for el in src.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "tbl":
            text = Table(el, src).rows[0].cells[0].text.strip()
            parts = [p.strip() for p in text.split("|")]
            if len(parts) == 3 and (parts[0].startswith("P-") or parts[0].startswith("R2-")):
                comments.append({"id": parts[0], "severity": parts[1], "title": parts[2], "items": []})
        elif tag == "p" and comments:
            t = Paragraph(el, src).text.strip()
            if t and any(t.startswith(k) for k in ("Location:", "Reviewer concern:", "Required action:", "Status:")):
                comments[-1]["items"].append(t)
    return comments


def render_comment(b: Doc, c: dict[str, Any]) -> None:
    b.label("The comment, as written")
    for t in c["items"]:
        head, _, rest = t.partition(":")
        para = b.doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(12)
        para.paragraph_format.space_after = Pt(3)
        r = para.add_run(head + ": ")
        r.bold = True
        r.font.color.rgb = QUOTE
        b._runs(para, rest.strip(), colour=QUOTE)


# ---------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------

def fig_frame(fa: dict) -> Path:
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10.2, 3.8), gridspec_kw={"width_ratios": [1.05, 1]})
    a1.set_xlim(0, 10); a1.set_ylim(0, 6); a1.axis("off")
    _box(a1, 0.1, 2.3, 2.3, 1.5, f"{fa['n_universe_3371']:,} segments\n(cap 455 / child)", INK2)
    _box(a1, 3.2, 3.8, 3.0, 1.6, f"ID rule (used)\nexclude 1,320 pool IDs\n→ {fa['n_used_frame_id_rule']:,}", ORANGE)
    _box(a1, 3.2, 0.6, 3.0, 1.6, f"Identity rule (intended)\nexclude 1,320 pool episodes\n→ {fa['n_intended_frame_identity_rule']:,}", AQUA)
    _box(a1, 6.8, 2.1, 3.2, 1.9, f"{fa['n_in_both']:,} in both\n{fa['n_wrongly_retained_pool']} pool episodes\nkept by mistake\n{fa['n_wrongly_excluded_non_pool']} non-pool dropped", BLUE, fill=BLUE, tc="white", size=7.2, weight="bold")
    _arrow(a1, (2.4, 3.3), (3.2, 4.5)); _arrow(a1, (2.4, 2.8), (3.2, 1.5))
    _arrow(a1, (6.2, 4.5), (6.8, 3.5)); _arrow(a1, (6.2, 1.5), (6.8, 2.6))
    a1.text(0.1, 5.7, "A. Two exclusion rules on the same universe", fontsize=9.5, weight="bold")
    pc = fa["per_child"]
    kids = sorted(pc, key=lambda c: pc[c]["coverage_of_intended"])
    vals = [100 * (1 - pc[c]["coverage_of_intended"]) for c in kids]
    a2.barh(range(len(kids)), vals, color=ORANGE, height=0.62, zorder=3)
    for i, (c, v) in enumerate(zip(kids, vals)):
        a2.text(v + 0.8, i, f"{v:.1f}%  ({pc[c]['wrongly_excluded']}/{pc[c]['intended_frame']})", va="center", fontsize=7.2)
    a2.set_yticks(range(len(kids)), kids, fontsize=7.6)
    a2.set_xlabel("% of the child's intended frame swapped out")
    a2.set_xlim(0, 48); a2.grid(axis="x", color=GRID, lw=0.6); a2.set_axisbelow(True)
    a2.set_title("B. Frame distortion by child", loc="left", fontsize=9.5, weight="bold")
    fig.tight_layout()
    p = FIG / "r4_fig1_frame_audit.png"; fig.savefig(p); plt.close(fig); return p


def fig_frame_types(fa: dict) -> Path:
    prof = fa["profile"]
    sets = [("analysed frame\n(in both)", "analysed_frame_in_both"), ("non-pool episodes\ndropped by ID rule", "wrongly_excluded"),
            ("pool episodes\nkept by ID rule", "wrongly_retained")]
    types = ["other", "play", "request", "response", "joint_attention", "protest", "initiation"]
    cols = [GREY, BLUE, ORANGE, AQUA, YELLOW, "#7b5ea7", "#c9c7c0"]
    fig, ax = plt.subplots(figsize=(9.6, 2.9))
    for i, (lab, k) in enumerate(sets):
        et, n = prof[k]["episode_types"], prof[k]["n"]
        left = 0
        for t, col in zip(types, cols):
            v = 100 * et.get(t, 0) / n
            ax.barh(i, v, left=left, color=col, edgecolor="white", lw=2, height=0.6)
            if v >= 7:
                ax.text(left + v / 2, i, f"{v:.0f}%", ha="center", va="center", fontsize=7.2,
                        color="white" if col in (BLUE, ORANGE, "#7b5ea7") else INK)
            left += v
        ax.text(101, i, f"n = {n:,}", va="center", fontsize=7.6)
    ax.set_yticks(range(3), [s[0] for s in sets], fontsize=8); ax.invert_yaxis()
    ax.set_xlim(0, 112); ax.set_xlabel("% of episodes, by episode type")
    ax.legend(handles=[Patch(color=col, label=t.replace("_", " ")) for t, col in zip(types, cols)],
              frameon=False, fontsize=7.4, ncol=7, loc="upper center", bbox_to_anchor=(0.45, 1.2))
    ax.spines["left"].set_visible(False)
    fig.tight_layout()
    p = FIG / "r4_fig2_frame_types.png"; fig.savefig(p); plt.close(fig); return p


def fig_universe(su: dict) -> Path:
    g = su["groups"]
    fig, ax = plt.subplots(figsize=(9.8, 3.9))
    ax.set_xlim(0, 12); ax.set_ylim(0, 6); ax.axis("off")
    _box(ax, 0.1, 3.9, 3.2, 1.6, "2,442-segment universe\n(cap 273 / child)", BLUE)
    _box(ax, 0.1, 0.9, 3.2, 1.6, "3,371-segment universe\n(cap 455 / child)\n⊃ the 2,442", AQUA)
    _box(ax, 4.0, 4.2, 2.6, 1.3, "Pool 1,320\n→ 330 judgeable\n(main batch)", BLUE, fill=BLUE, tc="white", weight="bold")
    _box(ax, 4.0, 0.9, 2.6, 1.6, f"Evaluation set\n132 frozen\n→ {g['overlap_excluded_121']['n']} overlap-excluded", AQUA, fill=AQUA, tc="white", weight="bold")
    _box(ax, 7.4, 3.2, 2.2, 1.3, f"same universe\nunconditioned\nn = {g['same_universe_unconditioned']['n']}", BLUE, size=7.8)
    _box(ax, 10.0, 3.2, 1.9, 1.3, f"…and\njudgeable\nn = {g['same_universe_judgeable']['n']}", BLUE, size=7.8)
    _box(ax, 7.4, 0.9, 2.2, 1.3, f"higher-cap\nonly\nn = {g['higher_cap_only']['n']}", AQUA, size=7.8)
    _arrow(ax, (3.3, 4.7), (4.0, 4.8)); _arrow(ax, (3.3, 1.7), (4.0, 1.7))
    _arrow(ax, (6.6, 2.0), (7.4, 3.6)); _arrow(ax, (6.6, 1.5), (7.4, 1.5)); _arrow(ax, (9.6, 3.85), (10.0, 3.85))
    ax.text(7.4, 5.2, "Split of the 121 by stable episode identity", fontsize=8.5, weight="bold")
    p = FIG / "r4_fig3_universe_design.png"; fig.savefig(p); plt.close(fig); return p


def fig_universe_forest(su: dict) -> Path:
    order = [("main_batch_330", "Main batch (330)\n2,442 univ., judgeable"),
             ("same_universe_judgeable", "Same universe, judgeable\n(eval set)"),
             ("same_universe_unconditioned", "Same universe,\nunconditioned (eval set)"),
             ("higher_cap_only", "Higher-cap only\n(eval set)"),
             ("overlap_excluded_121", "All 121")]
    fig, axes = plt.subplots(1, 2, figsize=(10.4, 3.9), sharey=True)
    for ax, key, title in zip(axes, ("merged_agreement", "unsupported_rate"), ("Merged agreement", "Unsupported assertion")):
        for i, (g, _) in enumerate(order):
            for j, m in enumerate(MODELS):
                v = su["groups"][g]["per_model"][m][key]
                y = i + (j - 1.5) * 0.17
                ax.plot([100 * v["ci95"]["lower"], 100 * v["ci95"]["upper"]], [y, y], color=COLOUR[m], lw=2)
                ax.scatter([100 * v["value"]], [y], s=26, color=COLOUR[m], zorder=3, edgecolor="white", linewidth=0.8)
        ax.set_title(title, loc="left", fontsize=9.5, weight="bold")
        ax.set_xlabel("% (95% child-cluster CI)")
        ax.grid(axis="x", color=GRID, lw=0.6); ax.set_axisbelow(True)
    axes[0].set_yticks(range(len(order)), [o[1] for o in order], fontsize=7.6); axes[0].invert_yaxis()
    for m in MODELS:
        axes[0].scatter([], [], color=COLOUR[m], label=short(m))
    fig.legend(frameon=False, fontsize=8, loc="lower center", ncol=4, bbox_to_anchor=(0.55, -0.02))
    fig.tight_layout(rect=(0, 0.07, 1, 1))
    p = FIG / "r4_fig4_universe_forest.png"; fig.savefig(p); plt.close(fig); return p


def fig_mars(mc: dict) -> Path:
    comps = [("vs_random_reviewable", "vs random candidate\n(reviewable, primary)", BLUE),
             ("vs_random_all_generators", "vs random candidate\n(published definition)", GREY),
             ("vs_best_reviewable_fallback", "vs strongest generator\n(reviewable fallback, primary)", ORANGE),
             ("vs_best_fixed", "vs strongest generator\n(fixed qwen2.5:7b)", "#c9c7c0")]
    fig, ax = plt.subplots(figsize=(9.8, 4.2))
    for i, f in enumerate(FIELDS):
        for j, (k, lab, col) in enumerate(comps):
            d = mc["per_field"][f][k]
            y = i + (j - 1.5) * 0.19
            ax.plot([100 * d["lower"], 100 * d["upper"]], [y, y], color=col, lw=2.2)
            ax.scatter([100 * d["value"]], [y], s=30, color=col, zorder=3, edgecolor="white", linewidth=0.8,
                       label=lab if i == 0 else None)
    ax.axvline(0, color=INK2, lw=1)
    ax.set_yticks(range(4), [f.replace("_", " ") for f in FIELDS]); ax.invert_yaxis()
    ax.set_xlabel("MARS minus comparator, percentage points (95% child-cluster paired CI)")
    ax.grid(axis="x", color=GRID, lw=0.6); ax.set_axisbelow(True)
    ax.legend(frameon=False, fontsize=7.4, loc="upper center", ncol=4, bbox_to_anchor=(0.5, -0.14))
    fig.tight_layout()
    p = FIG / "r4_fig5_mars_clustered.png"; fig.savefig(p); plt.close(fig); return p


def fig_abstention(ab: dict) -> Path:
    d = ab["main_batch_330"]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10.2, 3.4), gridspec_kw={"width_ratios": [1.25, 1]})
    parts = [("correct", AQUA), ("incorrect", GREY), ("missing", ORANGE), ("correctly_null", BLUE), ("hallucinated", YELLOW)]
    for i, m in enumerate(MODELS):
        left = 0
        for k, col in parts:
            v = 100 * d[m]["counts"].get(k, 0) / d[m]["n_field_slots"]
            a1.barh(i, v, left=left, color=col, edgecolor="white", lw=2, height=0.6)
            if v >= 6:
                a1.text(left + v / 2, i, f"{v:.0f}", ha="center", va="center", fontsize=7,
                        color="white" if col in (BLUE, ORANGE) else INK)
            left += v
    a1.set_yticks(range(4), [short(m) for m in MODELS]); a1.invert_yaxis()
    a1.set_xlim(0, 100); a1.set_xlabel("% of 1,320 field slots (main batch)")
    a1.legend(handles=[Patch(color=col, label=k.replace("_", " ")) for k, col in parts], frameon=False, fontsize=7, ncol=5, loc="upper center", bbox_to_anchor=(0.5, 1.2))
    rates = [("reported_missing_over_all_slots", "reported (missing ÷ all slots)", GREY),
             ("null_on_gold_populated", "null on gold-populated fields", ORANGE),
             ("overall_model_null_rate", "overall model-null rate", BLUE)]
    for j, (k, lab, col) in enumerate(rates):
        a2.barh([i + (j - 1) * 0.26 for i in range(4)], [100 * d[m][k] for m in MODELS], height=0.24, color=col, label=lab)
        for i, m in enumerate(MODELS):
            a2.text(100 * d[m][k] + 0.8, i + (j - 1) * 0.26, f"{100 * d[m][k]:.1f}", va="center", fontsize=6.6)
    a2.set_yticks(range(4), [short(m) for m in MODELS]); a2.invert_yaxis(); a2.set_xlim(0, 72)
    a2.set_xlabel("%"); a2.legend(frameon=False, fontsize=7, loc="lower right")
    a2.grid(axis="x", color=GRID, lw=0.6); a2.set_axisbelow(True)
    fig.tight_layout()
    p = FIG / "r4_fig6_abstention.png"; fig.savefig(p); plt.close(fig); return p


def fig_cap(ec: dict) -> Path:
    rs = ec["resegmentation"]
    keys = ["8", "12", "16", "24", "none"]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10.0, 3.3))
    a1.bar(range(5), [100 * rs[k]["share_at_cap"] for k in keys], color=BLUE, width=0.6, zorder=3)
    for i, k in enumerate(keys):
        a1.text(i, 100 * rs[k]["share_at_cap"] + 1.2, f"{100 * rs[k]['share_at_cap']:.0f}%\n({rs[k]['n_episodes']:,} ep.)",
                ha="center", fontsize=7.2)
    a1.set_xticks(range(5), ["8\n(used)", "12", "16", "24", "no cap"]); a1.set_xlabel("maximum turns per episode")
    a1.set_ylabel("% of episodes reaching the cap"); a1.set_ylim(0, 68)
    a1.grid(axis="y", color=GRID, lw=0.6); a1.set_axisbelow(True)
    a1.set_title("A. Re-segmentation under larger caps (no model)", loc="left", fontsize=9, weight="bold")
    ab = ec["annotated_by_cap"]
    labels, vals, cols = [], [], []
    for side, nm in (("main_batch_330", "main batch"), ("frozen_132", "132 set")):
        for g, lab, col in (("capped_8", "8 turns", BLUE), ("under_8", "< 8 turns", AQUA)):
            labels.append(f"{nm}\n{lab}\n(n={ab[side][g]['n']})"); vals.append(100 * ab[side][g]["gold_null_mean"]); cols.append(col)
    a2.bar(range(4), vals, color=cols, width=0.6, zorder=3)
    for i, v in enumerate(vals):
        a2.text(i, v + 0.4, f"{v:.1f}%", ha="center", fontsize=7.6)
    a2.set_xticks(range(4), labels, fontsize=7.2); a2.set_ylabel("gold null share, mean of 4 fields")
    a2.set_ylim(0, 20); a2.grid(axis="y", color=GRID, lw=0.6); a2.set_axisbelow(True)
    a2.set_title("B. Gold nulls in capped vs shorter episodes", loc="left", fontsize=9, weight="bold")
    fig.tight_layout()
    p = FIG / "r4_fig7_episode_cap.png"; fig.savefig(p); plt.close(fig); return p


def _age_m(a: str) -> float:
    y, m = a.split(";")
    return int(y) * 12 + int(m)


def fig_corpus(rows: list[dict]) -> Path:
    rows = sorted(rows, key=lambda r: _age_m(r["age_first"]))
    fig, ax = plt.subplots(figsize=(9.4, 3.4))
    for i, r in enumerate(rows):
        a, b = _age_m(r["age_first"]) / 12, _age_m(r["age_last"]) / 12
        col = BLUE if r["corpus"] == "rollins" else ORANGE
        ax.plot([a, b], [i, i], color=col, lw=5, solid_capstyle="round")
        ax.text(b + 0.08, i, f"{r['sessions']} sessions", va="center", fontsize=7.2, color=INK2)
    ax.set_yticks(range(len(rows)), [r["child_id"] for r in rows], fontsize=7.8)
    ax.set_xlabel("age at recording (years), from each session's @ID header")
    ax.plot([], [], color=BLUE, lw=5, label="Rollins (child–investigator)")
    ax.plot([], [], color=ORANGE, lw=5, label="Flusberg (child–mother)")
    ax.legend(frameon=False, fontsize=7.8, loc="lower right")
    ax.grid(axis="x", color=GRID, lw=0.6); ax.set_axisbelow(True); ax.set_xlim(1.8, 10.2)
    fig.tight_layout()
    p = FIG / "r4_fig8_corpus_ages.png"; fig.savefig(p); plt.close(fig); return p


def fig_confusion(cf: dict) -> Path:
    fig, axes = plt.subplots(1, 4, figsize=(11.0, 3.2))
    for ax, f in zip(axes, FIELDS):
        cats = cf[f]["categories"]; mat = cf[f]["matrix_a_rows_b_cols"]
        ax.imshow(mat, cmap="Blues")
        mx = max(max(r) for r in mat)
        for i in range(len(cats)):
            for j in range(len(cats)):
                if mat[i][j]:
                    ax.text(j, i, mat[i][j], ha="center", va="center", fontsize=6.5,
                            color="white" if mat[i][j] > mx * 0.55 else INK)
        lab = ["null" if c == "None" else c.replace("_", " ")[:11] for c in cats]
        ax.set_xticks(range(len(cats)), lab, rotation=60, ha="right", fontsize=6.2)
        ax.set_yticks(range(len(cats)), lab, fontsize=6.2)
        ax.set_title(f"{f.replace('_', ' ')}  (agree {pct(cf[f]['observed_agreement'])})", fontsize=8.5, weight="bold", loc="left")
        for s in ax.spines.values():
            s.set_visible(False)
    axes[0].set_ylabel("annotator A", fontsize=8); fig.supxlabel("annotator B", fontsize=8)
    fig.tight_layout()
    p = FIG / "r4_fig9_confusion.png"; fig.savefig(p); plt.close(fig); return p


# ---------------------------------------------------------------------------
# Answers
# ---------------------------------------------------------------------------

def a_p1(b: Doc, ctx) -> None:
    b.label("Short answer")
    b.answer_line(
        "**These three statements contain facts only the authors hold, so they cannot be completed from the "
        "study artefacts, and they have not been invented here.** The manuscript builder still carries the "
        "placeholder (`make_manuscript.py:1984`, \"[TO COMPLETE BEFORE SUBMISSION]\"). Below is the exact "
        "structure MDPI expects, with every factual slot left in brackets for the authors to fill.")
    b.boxed("Template (fill every bracket; delete any sentence that is not true)", [
        "**Funding:** \"This research received [no external funding / funding from [funder name] under grant "
        "[number]]. [The funder had no role in study design, data analysis, interpretation or the decision to "
        "publish.]\"",
        "**Author Contributions (CRediT):** Conceptualization, [initials]; Methodology, [ ]; Software, [ ]; "
        "Validation, [ ]; Formal analysis, [ ]; Investigation, [ ]; Data curation, [ ]; Writing—original draft, "
        "[ ]; Writing—review and editing, [ ]; Visualization, [ ]; Supervision, [ ]; Project administration, [ ]. "
        "\"All authors have read and agreed to the published version of the manuscript.\" Annotators and the "
        "adjudicator either appear here or in Acknowledgments, according to their role.",
        "**Conflicts of Interest:** \"The authors declare no conflicts of interest.\" or a specific declaration.",
        "**Use of generative AI (MDPI policy):** a statement naming the tools used in preparing the manuscript "
        "and code, and confirming author responsibility for the content. This is separate from the LLMs "
        "studied, which are already described in Methods.",
    ], fill="FFF6E5", edge="EDA100")
    b.p("Status after this round: **author action required.** Nothing in the analysis depends on it.")


def a_p2(b: Doc, ctx) -> None:
    rows = ctx["r4"]["corpus_characteristics"]
    tot_u = sum(r["utterances"] for r in rows); tot_s = sum(r["sessions"] for r in rows)
    assert (tot_u, tot_s) == (75003, 85)
    b.label("Short answer")
    b.answer_line(
        "**Added: a per-child table built only from metadata in the released CHAT files.** Every value "
        "below is read from each session's `@ID` header (age, sex), the `@Types` header, or counted from the "
        "transcript. Nothing is inferred. Script: `round4_analysis.py` (corpus_characteristics). "
        f"Totals reconcile with the manuscript: {tot_s} sessions, {tot_u:,} utterances.")
    b.table("Corpus characteristics, per child, as recorded in the CHAT headers. Age = child age at the first "
            "and last recorded session (years;months). Source: `round4_analysis.json`.",
            ["Corpus", "Child", "Sessions", "Age span", "Sex", "@Types", "Utterances", "Child utts", "Partner"],
            [[r["corpus"].title(), r["child_id"], str(r["sessions"]), f"{r['age_first']}–{r['age_last']}",
              r["sex"], r["types_header"], f"{r['utterances']:,}", f"{r['child_utterances']:,}",
              {"MOT": "mother", "INV": "investigator"}[r["main_partner"]]] for r in rows],
            [0.7, 1.05, 0.6, 0.85, 0.5, 1.2, 0.8, 0.7, 0.9])
    b.figure(fig_corpus(rows), "Age span covered by each child's sessions. The two corpora do not overlap much "
             "in age: the Rollins children were recorded between 2 and 4 years, the Flusberg children mostly "
             "between 3½ and 10.")
    b.label("Supporting argument")
    b.bullets([
        "**Not available in the released files, and so not reported:** standardised developmental, cognitive "
        "or language scores, diagnostic instrument results, race/ethnicity and socioeconomic information. "
        "The table should say so explicitly rather than leave the columns out silently.",
        "**Diagnosis is a corpus label.** `@Types: long, toyplay, ASD` is the corpus-level designation used "
        "by ASDBank; the files do not document per-child diagnostic procedures. Describe it as \"children "
        "described in the source corpus as autistic\" rather than as a verified diagnosis.",
        "**All eleven children are recorded as male.** This limits generalisability and belongs in Section 4.9.",
        "**Two children share the pseudonym \"Roger\"** (Rollins 2;06–3;03; Flusberg 3;09–5;06). The "
        "metadata cannot establish whether they are the same child, so they are treated as distinct, as "
        "throughout the study. State this.",
    ])


def a_p3(b: Doc, ctx) -> None:
    cf = ctx["r4"]["annotator_confusion"]
    b.label("Short answer")
    b.answer_line(
        "**Every promised non-repository item has a source in the project, and one gap was closed this "
        "round: the extended reliability results lacked confusion matrices, which are now generated.** The "
        "supplementary documents themselves (S1–S7) are assembled by the authors from these sources; this "
        "round checked that each source exists and agrees with the main text.")
    b.table("Supplementary content the manuscript promises, and where each item comes from.",
            ["Promised content", "Source in the project", "Status"],
            [["Final codebook and schema rules", "`docs/annotation_codebook.md`; `src/backend/schemas/observation_schema.py` (allowed values)", "present"],
             ["Exact prompts and configuration", "`src/backend/ml/extraction/prompts.py`; generator/judge roster, thresholds and seeds in `mars_run_summary.json`", "present"],
             ["Adequacy-router decision logic", "`src/backend/ml/adequacy_router.py`; `docs/lld/Adequacy_Router.md`; Round-3 Figure 10", "present"],
             ["Extended reliability and confusion matrices", "`reliability_current.json`, `reliability_round1.json`, `round_comparison.json`; confusion matrices in `round4_analysis.json` (Figure below)", "**gap closed this round**"],
             ["Extended generator and MARS analyses", "`round3_analysis.json`, `round4_analysis.json`, `mars_evaluation.json`, `mars_selection_analysis.json`", "present; MARS intervals replaced (R2-3)"],
             ["Authorization / access-control verification", "10 test suites, 120 passed (run record 2026-09-26); Round-3 Q9 table", "present"],
             ["Corpus characteristics (S5)", "`round4_analysis.json` (P-2 table)", "**added this round**"]],
            [2.0, 3.6, 1.1])
    b.figure(fig_confusion(cf), "Annotator A × annotator B, 330 segments, final pre-adjudication state (the "
             "state in which κ = 0.801 is computed). Observed agreement on each field matches "
             "`reliability_current.json` (checked by the script).")
    b.p("Status after this round: **sources complete; assembly of the S-files is an author step.** One "
        "consistency rule for that step: each S-file should cite the artefact name used here, so that a "
        "reviewer can trace any supplementary number to the same file as the main text.")


def a_r21(b: Doc, ctx) -> None:
    su = ctx["r4"]["same_universe"]; g = su["groups"]; ts = su["tags_summary"]; cn = su["contrasts"]
    m7 = "qwen2.5:7b"
    mm = lambda grp, m=m7: g[grp]["per_model"][m]["merged_agreement"]  # noqa: E731
    c_j = cn["main_batch_330-minus-same_universe_judgeable"]
    # Checked claims
    assert not c_j[m7]["merged_agreement"]["excludes_zero"]
    assert mm("higher_cap_only")["value"] >= mm("same_universe_unconditioned")["value"]
    assert all(c_j[m]["merged_agreement"]["value"] > 0 for m in MODELS)
    assert [m for m in MODELS if c_j[m]["merged_agreement"]["excludes_zero"]] == ["llama3.2:latest", "qwen2.5:3b"]
    hc = cn["higher_cap_only-minus-same_universe_unconditioned"]
    assert not any(hc[m]["merged_agreement"]["upper"] < 0 for m in MODELS)
    batch_yield = 330 / 506
    b.label("Short answer")
    b.answer_line(
        "**The preferred option was feasible and is done, with no model re-run.** Matching episodes by stable "
        f"identity shows that {g['same_universe_unconditioned']['n']} of the {g['overlap_excluded_121']['n']} "
        "overlap-excluded evaluation segments also belong to the 2,442-segment universe of the main batch. "
        f"Of those, {g['same_universe_judgeable']['n']} would also have passed the batch's judgeability "
        "filter. **The result supports narrowing the wording rather than the causal reading:** "
        f"qwen2.5:7b scores {pct(mm('same_universe_judgeable')['value'])} on the same-universe, judgeable "
        f"segments against {pct(mm('main_batch_330')['value'])} on the main batch. So restoring both the universe "
        "and the filter does not restore main-batch agreement. The higher-cap-only segments score "
        f"{pct(mm('higher_cap_only')['value'])}, so the change of universe does not lower agreement either.")
    b.figure(fig_universe(su), "How the 121 overlap-excluded evaluation segments split by universe and by the "
             "batch's judgeability rule. Pool episodes are excluded from all evaluation groups, so the "
             "same-universe group is the part of the 2,442 universe the main batch never sampled.")
    rows = []
    for key, lab in (("main_batch_330", "Main batch"), ("same_universe_judgeable", "Same universe, judgeable"),
                     ("same_universe_unconditioned", "Same universe, unconditioned"),
                     ("higher_cap_only", "Higher-cap only"), ("overlap_excluded_121", "All 121")):
        rows.append([f"{lab} (n = {g[key]['n']})"] + [val_ci(g[key]["per_model"][m]["merged_agreement"]) for m in MODELS])
    b.table("Merged categorical agreement, point estimate [95% child-cluster CI], 10,000 replicates. Source: "
            "`round4_analysis.json` (same_universe.groups).", ["Group"] + [short(m) for m in MODELS], rows,
            [1.9, 1.2, 1.2, 1.2, 1.2])
    rows = []
    for key, lab in (("main_batch_330-minus-same_universe_judgeable", "Main − same universe, judgeable"),
                     ("main_batch_330-minus-same_universe_unconditioned", "Main − same universe, unconditioned"),
                     ("higher_cap_only-minus-same_universe_unconditioned", "Higher-cap only − same universe"),
                     ("main_batch_330-minus-overlap_excluded_121", "Main − all 121")):
        rows.append([lab] + [ci_pp(cn[key][m]["merged_agreement"]) for m in MODELS])
    b.table("Differences in merged agreement, percentage points [95% CI]; children resampled jointly for both "
            "groups. * marks an interval excluding zero.", ["Contrast"] + [short(m) for m in MODELS], rows,
            [2.1, 1.15, 1.15, 1.15, 1.15])
    b.figure(fig_universe_forest(su), "Agreement and unsupported assertion by group, all four generators. "
             "Intervals for the evaluation-set groups are wide (36–85 segments).")
    b.label("Supporting argument")
    b.bullets([
        "**Why the same-universe group is a fair comparator.** The 1,320-segment pool was a seeded random "
        "draw of 120 segments per child from the 2,442. The remainder of that universe was therefore "
        "exchangeable with the pool within each child. The evaluation segments that fall in it are an "
        "unconditioned sample of the same universe, drawn without reference to model output.",
        f"**The judgeability reconstruction is consistent with the original filter.** The batch rule (at "
        f"least one populated field from each of the three original generators) passes "
        f"{ts['judgeable_all3_among_121']} of 121 evaluation segments ({pct(ts['judgeable_all3_among_121'] / 121)}), "
        f"close to the main batch's own yield of 330 of 506 ({pct(batch_yield)}).",
        "**What remains of the gap is not attributable to either design feature.** For qwen2.5:7b and "
        "qwen3:14b the main-batch advantage over the same-universe judgeable group is not distinguishable "
        "from zero. For llama3.2 and qwen2.5:3b it is (see the table). Other differences between the arms "
        "cannot be separated in this design: the main-batch reference labels came from two annotation "
        "rounds with a codebook introduced between them, while the evaluation set was annotated under the "
        "final codebook from the start, at a different time.",
        "**Caveats.** Judgeability was computed from the evaluation run's generator outputs, not by "
        "re-running the batch procedure, and generation is not guaranteed to be identical across runs. Both "
        "runs recorded five of the eight extracted fields. The subgroups are small, so the absence of a "
        "difference is not evidence of equality.",
    ])
    b.boxed("Suggested manuscript wording (Methods 2.13 / Results 3.3 / Discussion 4.4)", [
        "\"Agreement was lower in the independently sampled, higher-cap evaluation branch, which did not use "
        "model tractability for inclusion. The two branches differ in candidate universe, inclusion "
        "condition and annotation round, so the difference should not be attributed to any single factor. A "
        f"same-universe sensitivity analysis ({g['same_universe_unconditioned']['n']} evaluation segments from "
        f"the main batch's 2,442-segment universe, {g['same_universe_judgeable']['n']} of which met the batch's "
        "judgeability condition) did not reproduce main-batch agreement. Neither removal of tractability "
        "conditioning nor the change of candidate universe accounted for the difference on its own.\"",
        "Remove any sentence that says the decline is caused by removing model-tractability conditioning.",
    ], fill="EAF7F1", edge="1BAF7A")


def a_r22(b: Doc, ctx) -> None:
    fa = ctx["r4"]["frame_audit"]; ps = ctx["r4"]["same_universe"]["poststratified_121"]["episode_type_7"]
    fe = ps["frame_effect"]
    max_m = max(abs(fe[m]["merged_agreement"]) for m in MODELS)
    max_u = max(abs(fe[m]["unsupported_rate"]) for m in MODELS)
    worst = min(fa["per_child"], key=lambda c: fa["per_child"][c]["coverage_of_intended"])
    unacc = ps["to_intended_frame"]["unaccounted_intended_share_by_child"]
    worst_unacc = max(unacc, key=unacc.get)
    assert fa["frozen_draw_reproduced"] and fa["frozen_segments_in_pool"] == 11
    assert fa["n_wrongly_retained_pool"] == fa["n_wrongly_excluded_non_pool"]
    b.label("Short answer")
    b.answer_line(
        "**Done, diagnostically, without redrawing.** Episodes were matched by stable identity (session plus "
        "source start line), and the intended frame was rebuilt. The reconstructed ID-based frame reproduces "
        "the frozen 132-segment draw exactly, which confirms that the audited frame is the one that was "
        f"used. **The ID rule wrongly retained {fa['n_wrongly_retained_pool']} true pool episodes and wrongly "
        f"excluded {fa['n_wrongly_excluded_non_pool']} true non-pool episodes, "
        f"{pct(fa['share_of_intended_frame_excluded'])} of the intended {fa['n_intended_frame_identity_rule']:,}.** "
        f"The distortion is uneven across children ({worst}: "
        f"{pct(1 - fa['per_child'][worst]['coverage_of_intended'])} of its intended frame swapped), so it is "
        "not negligible at child level. The 121 should be described as the reviewer proposes.")
    b.figure(fig_frame(fa), "A: the used and intended frames have the same size because the ID rule swaps "
             "one-for-one. Each pool ID it removes names some episode of the same child, and each pool episode "
             "it keeps has lost its ID to another. B: share of each child's intended frame that was swapped.")
    b.table("Frame audit per child. Source: `round4_analysis.json` (frame_audit.per_child).",
            ["Child", "Intended frame", "In both", "Pool kept (wrong)", "Non-pool dropped (wrong)", "Coverage"],
            [[c, str(v["intended_frame"]), str(v["in_both"]), str(v["wrongly_retained"]), str(v["wrongly_excluded"]),
              pct(v["coverage_of_intended"])] for c, v in fa["per_child"].items()]
            + [["**Total**", f"{fa['n_intended_frame_identity_rule']:,}", f"{fa['n_in_both']:,}",
                str(fa["n_wrongly_retained_pool"]), str(fa["n_wrongly_excluded_non_pool"]),
                pct(1 - fa["share_of_intended_frame_excluded"])]],
            [1.4, 1.1, 0.9, 1.2, 1.4, 0.9])
    b.figure(fig_frame_types(fa), "Composition by episode type. The dropped non-pool episodes are richer in "
             "response and protest episodes; the wrongly kept pool episodes are richer in play, joint-attention "
             "and request episodes. The analysed frame differs little from the intended frame overall.")
    b.label("How much could the distortion move the estimates?")
    b.p("The 121 were re-weighted to each frame's episode-type mix within child, once to the analysed frame "
        "and once to the intended frame. The difference between the two isolates the frame distortion from "
        "ordinary sampling noise.")
    b.table("Frame effect on the 121-segment estimates (intended-frame weighting minus analysed-frame "
            "weighting), percentage points. Source: `round4_analysis.json` (poststratified_121.episode_type_7).",
            ["Generator", "Merged agreement", "Unsupported assertion", "Whole-record"],
            [[short(m)] + [pp(fe[m][k]) for k in ("merged_agreement", "unsupported_rate", "whole_record")] for m in MODELS],
            [1.6, 1.6, 1.8, 1.6])
    b.p(f"On observable episode type, the distortion moves merged agreement by at most "
        f"{100 * max_m:.1f} points and unsupported assertion by at most {100 * max_u:.1f} points. That is "
        "small next to the intervals. **Two limits keep this from proving the distortion negligible:** "
        "re-weighting corrects only for what is observed, and some intended cells have no sampled segment to "
        f"weight ({worst_unacc}: {pct(unacc[worst_unacc])} of its intended frame is in episode types absent "
        "from its sample).")
    b.label("Supporting argument")
    b.bullets([
        "**Why no redraw.** Redrawing from the corrected frame after results were seen would select a sample "
        "with knowledge of the outcome. The audit quantifies the defect instead, as the reviewer asks.",
        f"**Consistency with Round 3.** Stable identity finds the same {fa['frozen_segments_in_pool']} "
        "pool episodes among the 132 that text matching found in Round 3.",
        "**Revision of Round-3 advice.** Round 3 recommended the 121 as the primary unseen-episode estimate. "
        "Given this audit, it should be presented as the reviewer proposes, with the qualification in the "
        "wording below.",
        "**Why text alone is not an identity.** Two episodes of one session can render to identical text "
        "(e.g. runs of unintelligible `xxx .` turns), so the audit keys on source line numbers.",
    ])
    b.boxed("Suggested manuscript wording (Section 2.13, Figure 3, Sections 3.2, 4.4, 4.9)", [
        f"\"A post-hoc audit using stable episode identity found that the identifier-based exclusion retained "
        f"{fa['n_wrongly_retained_pool']} pool episodes and excluded {fa['n_wrongly_excluded_non_pool']} "
        f"non-pool episodes ({pct(fa['share_of_intended_frame_excluded'])} of the intended frame; up to "
        f"{pct(1 - fa['per_child'][worst]['coverage_of_intended'])} for one child). Eleven of the 132 frozen "
        "evaluation segments were pool episodes. We therefore report a post-audit overlap-excluded analysis "
        "(n = 121) from the originally frozen sampling frame, alongside the original 132. The sample was not "
        "redrawn. Re-weighting to the intended frame's episode-type composition changed estimates by less "
        f"than {max(100 * max_m, 100 * max_u):.1f} percentage points, but differences on unobserved "
        "characteristics cannot be excluded.\"",
    ], fill="EAF7F1", edge="1BAF7A")


def a_r23(b: Doc, ctx) -> None:
    mc = ctx["r4"]["mars_clustered"]; msa = ctx["msa"]; pf = mc["per_field"]
    ex = lambda k: [f for f in FIELDS if pf[f][k]["excludes_zero"]]  # noqa: E731
    max_shift = max(abs(100 * (msa["per_field"][f]["vs_best_generator"]["difference"] - pf[f]["vs_best_fixed"]["value"]))
                    for f in FIELDS)
    assert max_shift <= 0.1
    seg_w = lambda d: d["upper"] - d["lower"]  # noqa: E731
    assert any(seg_w(pf[f]["vs_random_all_generators"]) < seg_w(msa["per_field"][f]["vs_random_candidate"]) for f in FIELDS)
    assert not any(pf[f]["vs_best_reviewable_fallback"]["lower"] > 0 for f in FIELDS)
    b.label("Short answer")
    b.answer_line(
        "**Recomputed with a child-cluster paired bootstrap on all four fields, against both comparators.** "
        f"The bootstrap uses 10,000 replicates. It resamples the {mc['n_children']} children and keeps all of "
        f"each child's reranked segments ({min(mc['segments_per_child'].values())}–"
        f"{max(mc['segments_per_child'].values())} per child, {mc['n_judged']} in total). "
        "**One conclusion changes in size, not direction.** MARS still beats a random candidate and still "
        "does not beat the strongest generator. Its margin over random roughly halves once the random "
        f"comparator is limited to candidates MARS could select (mean {pp(mc['mean']['mars'] - mc['mean']['random_expected_reviewable'])} "
        f"points instead of {pp(mc['mean']['mars'] - mc['mean']['random_expected_all_generators'])}; R2-6).")
    rows = []
    for f in FIELDS:
        rows.append([f.replace("_", " "), pct(pf[f]["mars"]), pct(pf[f]["random_expected_reviewable"]),
                     ci_pp(pf[f]["vs_random_reviewable"]), pct(pf[f]["best_reviewable_fallback"]),
                     ci_pp(pf[f]["vs_best_reviewable_fallback"])])
    b.table("Primary comparisons (definitions in R2-6), 142 reranked segments. Differences are observed point "
            "estimates with 95% child-cluster paired intervals, percentage points. * excludes zero. Source: "
            "`round4_analysis.json` (mars_clustered).",
            ["Field", "MARS", "Random (expected)", "MARS − random", "Strongest (fallback)", "MARS − strongest"],
            rows, [1.0, 0.7, 1.0, 1.45, 1.1, 1.45])
    rows = []
    for f in FIELDS:
        old_r, old_b = msa["per_field"][f]["vs_random_candidate"], msa["per_field"][f]["vs_best_generator"]
        rows.append([f.replace("_", " "),
                     f"[{pp(old_r['lower'])}, {pp(old_r['upper'])}]" + (" *" if old_r["excludes_zero"] else ""),
                     ci_pp(pf[f]["vs_random_all_generators"]),
                     f"[{pp(old_b['lower'])}, {pp(old_b['upper'])}]" + (" *" if old_b["excludes_zero"] else ""),
                     ci_pp(pf[f]["vs_best_fixed"])])
    b.table("Sensitivity: the published comparator definitions, before (segment-level bootstrap, as printed in "
            "the manuscript) and after (child-cluster). Intervals only for the published segment-level column.",
            ["Field", "vs random: segment-level", "vs random: child-cluster", "vs strongest: segment-level",
             "vs strongest: child-cluster"], rows, [1.0, 1.4, 1.5, 1.4, 1.5])
    b.figure(fig_mars(mc), "MARS minus each comparator, per field, with child-cluster paired intervals. Primary "
             "definitions in colour; published definitions in grey.")
    b.label("What changes in the interpretation")
    b.bullets([
        f"**Against a random candidate (primary definition):** the interval excludes zero for "
        f"{', '.join(ex('vs_random_reviewable'))}, and includes zero for "
        f"{', '.join(f for f in FIELDS if f not in ex('vs_random_reviewable')) or 'none'}. Under the published "
        f"definition it excluded zero for {', '.join(ex('vs_random_all_generators'))}.",
        f"**Against the strongest generator:** MARS is lower on domain "
        f"({ci_pp(pf['domain']['vs_best_reviewable_fallback'])}). On the other fields the interval includes "
        "zero. No field shows MARS above the strongest generator.",
        "**Clustering did not widen every interval.** The child-level intervals are narrower than the "
        "segment-level ones for some fields. That happens when between-child variation in the difference is "
        "small relative to within-child noise. It is a property of these data, not an error, and it is why "
        "the reviewer's request was worth doing rather than assuming the direction.",
        "**A second correction.** The earlier script reported the mean of the bootstrap replicates as the "
        "difference. The tables now report the observed difference (for the strongest-generator comparison "
        f"the change is at most {max_shift:.1f} points).",
    ])
    b.boxed("Suggested manuscript wording (Sections 2.16, 3.5)", [
        "\"Paired differences between MARS and each comparator were estimated with a child-cluster paired "
        "bootstrap (10,000 replicates): children were resampled with replacement, all of a sampled child's "
        "reranked segments were retained, and both policies were scored on the same resample.\"",
        f"\"On the {mc['n_judged']} reranked segments, MARS agreed with the reference standard more often than "
        f"the expected random reviewable candidate on {len(ex('vs_random_reviewable'))} of 4 fields, and did not "
        "exceed the strongest-generator policy on any field (domain: "
        f"{ci_pp(pf['domain']['vs_best_reviewable_fallback']).replace(' *', '')} percentage points).\"",
    ], fill="EAF7F1", edge="1BAF7A")


def a_r24(b: Doc, ctx) -> None:
    ab = ctx["r4"]["abstention"]; d = ab["main_batch_330"]; h = ab["frozen_132"]
    prov = ab["null_provenance_in_recorded_outputs"]
    q = d["qwen2.5:7b"]
    assert abs(q["reported_missing_over_all_slots"] - ctx["hm"]["batch"]["per_model"]["qwen2.5:7b"]["mean_abstention_rate"]) < 1e-3
    b.label("Short answer")
    b.answer_line(
        "**Agreed, and the definition is slightly different from the one in the comment.** The reported "
        "metric divides by **all** scored field slots, not by gold-populated ones: per field, missing ÷ N, "
        "then averaged over the four fields (`paper_metrics.summarise_scored`). For qwen2.5:7b on the main "
        f"batch that is {q['counts']['missing']}/{q['n_field_slots']} = {pct(q['reported_missing_over_all_slots'])}. "
        "It counts only cases where the reference records a value and the model gives none. **Rename it and "
        "report two companion rates**, so that each name matches its denominator.")
    b.table("Three null-output rates per generator, main batch (1,320 field slots) and 132 set (528). Source: "
            "`round4_analysis.json` (abstention).",
            ["Generator", "Missed-value rate (reported)", "Null on gold-populated fields", "Overall model-null rate",
             "Same three, 132 set"],
            [[short(m), pct(d[m]["reported_missing_over_all_slots"]), pct(d[m]["null_on_gold_populated"]),
              pct(d[m]["overall_model_null_rate"]),
              f"{pct(h[m]['reported_missing_over_all_slots'])} / {pct(h[m]['null_on_gold_populated'])} / {pct(h[m]['overall_model_null_rate'])}"]
             for m in MODELS],
            [1.2, 1.35, 1.35, 1.3, 1.6])
    b.figure(fig_abstention(ab), "A: where each generator's field slots fall. Model null is the orange plus "
             "blue segments. B: the three rates side by side. The reported metric is the smallest of the "
             "three by construction.")
    b.label("Definitions to use")
    b.bullets([
        "**Missed-value rate** (renamed from \"mean field abstention\"): share of all field slots where the "
        "reference standard records a value and the model returns null or omits the field. Denominator: all "
        "scored slots.",
        "**Null rate on answerable fields**: missing ÷ (slots where the reference records a value). This is "
        "the reviewer's \"abstention on answerable opportunities\".",
        "**Overall model-null rate**: (missing + correctly null) ÷ all slots. This is the model's frequency "
        "of declining to fill a field whatever the reference says, and the only one of the three that "
        "describes willingness to answer.",
        f"**Omission vs explicit null.** The scorer merges them, as the reviewer says. The recorded outputs "
        f"show how rare omission is: {prov['field_key_absent']} absent field keys against "
        f"{prov['explicit_null']:,} explicit nulls across all recorded candidates. Say that omission is scored "
        "as null and that it is uncommon.",
    ])
    b.boxed("Suggested manuscript wording (Sections 2.14, 2.17; Table 7; Section 3.4)", [
        "\"Missed-value rate is the proportion of all scored field slots in which the reference standard "
        "recorded a value and the model returned null or omitted the field; omitted fields are scored as null. "
        "We also report the null rate on answerable fields (denominator: slots with a reference value) and "
        "the overall model-null rate (all slots, regardless of the reference).\"",
        "Replace any sentence that uses the old metric to describe overall willingness to answer with the "
        "overall model-null rate.",
    ], fill="EAF7F1", edge="1BAF7A")


def a_r25(b: Doc, ctx) -> None:
    ec = ctx["r4"]["episode_cap"]; rs = ec["resegmentation"]; ab = ec["annotated_by_cap"]
    lost = rs["none"]["utterances_covered"] - rs["8"]["utterances_covered"]
    for side in ("main_batch_330", "frozen_132"):
        assert all(ab[side]["capped_8"]["gold_null_share"][f] == 0 for f in ("domain", "prompt_level"))
        assert ab[side]["capped_8"]["qwen2.5:7b_unsupported"] > ab[side]["under_8"]["qwen2.5:7b_unsupported"]
    assert ab["main_batch_330"]["capped_8"]["gold_null_mean"] < ab["main_batch_330"]["under_8"]["gold_null_mean"]
    assert ab["frozen_132"]["capped_8"]["gold_null_mean"] < ab["frozen_132"]["under_8"]["gold_null_mean"]
    b.label("Short answer")
    b.answer_line(
        "**How it was chosen: the eight-turn maximum is a design parameter from the system's technical "
        "specification, fixed before the corpora were segmented, and never tuned on the data.** Section 5.7.3 "
        "(\"Phase 2: Deidentification & Segment Selection\") of `Technical_Specification_Autism_AI_System.docx` "
        "states the target as \"1,000-3,000 segments, each consisting of 2-8 consecutive utterance turns with "
        "context\". The specification was committed on 2026-08-24, a day before the segmenter that sets "
        "`MAX_EPISODE_TURNS = 8` (`data/pipeline/segment_episodes.py`, commit 6767d65). The specification "
        "gives no rationale for the value 8, and no requirement, analysis or note chooses it. The spec file "
        "was generated programmatically (python-docx default metadata), and neither the research team's notes "
        "(`exp-data/data repos.docx`, which give only the 1,000–3,000 segment target) nor any exploratory "
        "analysis precedes it. **It is an undocumented implementation default, not a justified choice, and the "
        "manuscript should say so.** A model-free sensitivity analysis "
        "was feasible with existing artefacts and is reported below. So is the association between the cap "
        "and gold nulls in the annotated sets.")
    b.table("Re-segmentation of the same 85 sessions under larger caps. No model involved. Source: "
            "`round4_analysis.json` (episode_cap.resegmentation).",
            ["Maximum", "Episodes", "At cap", "Mean turns", "90th pct", "Longest", "Utterances covered"],
            [[k if k != "none" else "no cap", f"{v['n_episodes']:,}", pct(v["share_at_cap"]), f"{v['mean_turns']:.1f}",
              str(v["p90_turns"]), str(v["max_turns"]), f"{v['utterances_covered']:,} / {v['utterances_total']:,}"]
             for k, v in rs.items()],
            [0.8, 0.85, 0.75, 0.9, 0.8, 0.8, 1.6])
    rows = []
    for side, nm in (("main_batch_330", "Main batch"), ("frozen_132", "132 set")):
        for gname, lab in (("capped_8", "8 turns (capped)"), ("under_8", "< 8 turns")):
            v = ab[side][gname]
            rows.append([nm, lab, str(v["n"])] + [pct(v["gold_null_share"][f]) for f in FIELDS]
                        + [pct(v["qwen2.5:7b_unsupported"])])
    b.table("Reference-standard null share by field, and qwen2.5:7b unsupported-assertion rate, for capped and "
            "shorter episodes. Descriptive; not adjusted for other differences between the groups.",
            ["Set", "Episodes", "n", "domain", "modality", "prompt", "outcome", "qwen2.5:7b unsupported"],
            rows, [0.9, 1.1, 0.4, 0.7, 0.75, 0.65, 0.7, 1.2])
    b.figure(fig_cap(ec), "A: raising the cap reduces, but does not remove, cap-hitting; long child–adult "
             "exchanges continue for dozens of turns. B: capped episodes carry fewer reference nulls than "
             "shorter ones in both annotated sets.")
    b.label("Supporting argument")
    b.bullets([
        f"**The cap splits long exchanges; it does not discard them.** With no cap, episodes cover "
        f"{rs['none']['utterances_covered']:,} utterances; with the cap of 8, {rs['8']['utterances_covered']:,}. "
        f"Only {lost:,} utterances ({pct(lost / rs['8']['utterances_total'])} of the corpus) leave "
        "coverage. The rest of a long exchange becomes the next episode. What the cap removes is context "
        "across the split, not transcript content.",
        "**The direction of the risk is visible in the data.** Capped episodes have fewer reference nulls "
        "(domain and prompt level are never null in them in either set). This is consistent with longer context "
        "making fields more often supported. It is also consistent with capped episodes being "
        "systematically different exchanges, and the design cannot separate the two. qwen2.5:7b's "
        "unsupported-assertion rate is higher on capped episodes, on small denominators.",
        "**Why not re-annotate.** Re-annotation under a longer window would be a different task definition "
        "and a new reference standard. The reviewer does not require it for this point, and nothing above "
        "shows the cap reversing a conclusion.",
    ])
    b.boxed("Suggested manuscript wording (Section 2.12 and Limitations 4.9)", [
        "\"The eight-turn maximum was an implementation default fixed in the system specification before "
        f"segmentation; it was not derived from the transcripts or tuned on the data. It was reached by "
        f"{pct(rs['8']['share_at_cap'])} of episodes, so it actively truncates long exchanges. The truncated "
        "turns form subsequent episodes rather than being discarded "
        f"({pct(1 - rs['8']['utterances_covered'] / rs['none']['utterances_covered'])} fewer utterances covered "
        "than without a cap), but context across those boundaries is lost. In the annotated sets, capped "
        "episodes had fewer reference nulls than shorter episodes, so both extraction behaviour and "
        "reference-null frequencies may differ under longer context windows.\"",
    ], fill="EAF7F1", edge="1BAF7A")


def a_r26(b: Doc, ctx) -> None:
    mc = ctx["r4"]["mars_clustered"]; msa = ctx["msa"]
    b.label("Short answer")
    b.answer_line(
        "**The published 32.5% was a Monte Carlo estimate: the mean of 2,000 seeded uniform draws (seed "
        "20260916) over all three generators' outputs for each segment.** It equals the analytical "
        f"expectation over those three outputs ({pct(mc['mean']['random_expected_all_generators'])}). But the "
        "paired intervals were taken against **one** further draw (seed 20260917), and the comparator "
        "included outputs that MARS could never select. Both are fixed by the definitions below.")
    b.table("Comparator definitions, and their mean accuracy over the four fields on the 142 reranked "
            "segments. Source: `round4_analysis.json` (mars_clustered.mean).",
            ["Comparator", "Definition", "Mean"],
            [["**Random candidate (primary)**", "Analytical expectation: for each segment, the share of MARS's "
              "distinct reviewable candidates (SUCCESS/PARTIAL, after MARS deduplication) that match the "
              "reference; averaged over segments. No random draw, so no seed is needed.",
              pct(mc["mean"]["random_expected_reviewable"])],
             ["Random candidate (published)", "Same expectation, over all three generators' outputs, including "
              "outputs MARS's admission gate rejects and counting agreeing generators separately.",
              f"{pct(mc['mean']['random_expected_all_generators'])} (published {pct(msa['mean_accuracy']['random_candidate'])})"],
             ["**Strongest generator (primary)**", "qwen2.5:7b if it has a reviewable candidate; otherwise the "
              "next generator with one, in the fixed order qwen2.5:7b > llama3.2 > qwen2.5:3b (the order of "
              "main-batch merged agreement). A policy realisable at runtime.",
              pct(mc["mean"]["best_reviewable_fallback"])],
             ["Strongest generator (published)", "qwen2.5:7b's output always, even where MARS's gate rejected it "
              f"({mc['n_segments_best_generator_not_reviewable']} of {mc['n_judged']} segments).",
              pct(mc["mean"]["best_fixed"])],
             ["MARS", "The candidate MARS selected.", pct(mc["mean"]["mars"])]],
            [1.6, 4.2, 0.9])
    b.label("Supporting argument")
    b.bullets([
        "**Why restrict to reviewable candidates.** MARS chooses among the candidates that pass its admission "
        "gate. A comparator that may pick an output MARS was never offered is not measuring selection. It "
        "mixes selection with admission, and it lowers the random baseline, which flatters MARS.",
        "**Why an expectation, not draws.** The expectation is exact and reproducible without a seed. Pairing "
        "per segment is still possible, because each segment has an expected hit rate. The child-cluster "
        "intervals in R2-3 pair MARS's hit with that expectation.",
        f"**The strongest-generator edge case is small but real.** qwen2.5:7b lacked a reviewable candidate on "
        f"{mc['n_segments_best_generator_not_reviewable']} segments. The two definitions differ by "
        f"{abs(100 * (mc['mean']['best_fixed'] - mc['mean']['best_reviewable_fallback'])):.1f} points on the mean.",
        "**One disclosure.** The generator order used for the fallback, like the choice of qwen2.5:7b as "
        "strongest, comes from the same main-batch data. State this.",
    ])
    b.boxed("Suggested manuscript wording (Sections 2.18, 3.5)", [
        "\"The random-candidate comparator is the analytical expected agreement of a uniform choice among the "
        "distinct reviewable candidates MARS received for each segment. The strongest-generator comparator "
        "uses qwen2.5:7b, the generator with the highest main-batch merged agreement; where it had no "
        f"reviewable candidate ({mc['n_segments_best_generator_not_reviewable']} of {mc['n_judged']} segments), the next generator in that order with a reviewable "
        "candidate is used.\"",
        f"Replace the 32.5% baseline with {pct(mc['mean']['random_expected_reviewable'])}, and report the "
        "published definition as a sensitivity analysis.",
    ], fill="EAF7F1", edge="1BAF7A")


ANSWERS = {"P-1": a_p1, "P-2": a_p2, "P-3": a_p3, "R2-1": a_r21, "R2-2": a_r22, "R2-3": a_r23,
           "R2-4": a_r24, "R2-5": a_r25, "R2-6": a_r26}


# ---------------------------------------------------------------------------

def build() -> Path:
    FIG.mkdir(parents=True, exist_ok=True)
    ctx = {"r4": load("round4_analysis.json"), "msa": load("mars_selection_analysis.json"),
           "hm": load("holdout_metrics.json")}
    comments = read_comments()
    assert [c["id"] for c in comments] == list(ANSWERS), [c["id"] for c in comments]
    assert all(len(c["items"]) == 4 for c in comments)
    r4 = ctx["r4"]; fa = r4["frame_audit"]; su = r4["same_universe"]; mc = r4["mars_clustered"]

    b = Doc()
    grey = RGBColor(0x52, 0x51, 0x4E)
    b.doc.add_heading("Response to MDPI Round-4 Review Comments", level=0)
    b.p("Answers the nine active comments in 'MDPI_Round4_Review_Comments.docx', in their original order and "
        "wording. Branch model-finetuning, built 2026-09-26. Internal working document; not for submission.", colour=grey)
    b.p("Every number is read from a recorded artefact named in the table that shows it. This round adds "
        "`round4_analysis.json`, computed from preserved outputs and the deterministic CHAT pipeline with **no "
        "model call**. Claims about the data are checked by assertion when this document is built.", colour=grey)

    b.h("Summary", 1)
    b.table("Each comment and the answer in one line.", ["ID", "Severity", "Answer"], [
        ["P-1", "Submission", "Template supplied; facts are the authors' to fill. Not invented"],
        ["P-2", "Moderate", "Per-child table from CHAT headers added: age, sex, sessions, partner. What is unavailable is stated"],
        ["P-3", "Moderate", "Every promised item mapped to its source; confusion matrices were missing and are now generated"],
        ["R2-1", "Major", f"Same-universe analysis done ({su['groups']['same_universe_unconditioned']['n']} segments; "
                          f"{su['groups']['same_universe_judgeable']['n']} judgeable). Does not reproduce main-batch agreement → narrow the causal wording"],
        ["R2-2", "Major", f"Frame rebuilt by episode identity: {fa['n_wrongly_retained_pool']} pool kept, "
                          f"{fa['n_wrongly_excluded_non_pool']} non-pool dropped ({pct(fa['share_of_intended_frame_excluded'])}); "
                          "uneven by child → describe the 121 as post-audit, from the original frame"],
        ["R2-3", "Major", "Child-cluster paired bootstrap, 4 fields × 2 comparators. MARS > random on "
                          f"{sum(mc['per_field'][f]['vs_random_reviewable']['excludes_zero'] for f in FIELDS)}/4 fields; "
                          "never > strongest; lower on domain"],
        ["R2-4", "Moderate", "Metric divides by all slots; renamed 'missed-value rate'; answerable-field and overall model-null rates added"],
        ["R2-5", "Moderate", f"Cap came from the specification, not tuned. Model-free sensitivity: the cap splits rather than drops turns; limitation wording supplied"],
        ["R2-6", "Minor", f"Random baseline = analytical expectation over reviewable candidates ({pct(mc['mean']['random_expected_reviewable'])}, "
                          f"not 32.5%); strongest-generator fallback defined"],
    ], [0.55, 0.85, 5.3])
    b.label("Two results that change what the manuscript says")
    b.bullets([
        "**R2-6 / R2-3:** the published random baseline counted outputs MARS could not select. Corrected, "
        f"MARS's mean margin over random falls from {pp(mc['mean']['mars'] - mc['mean']['random_expected_all_generators'])} "
        f"to {pp(mc['mean']['mars'] - mc['mean']['random_expected_reviewable'])} points. The direction is unchanged.",
        "**R2-2:** Round 3's advice to present the 121 as the unqualified primary unseen-episode estimate is "
        "withdrawn in favour of the reviewer's wording.",
    ])

    for c in comments:
        b.page_break()
        b.h(f"{c['id']}  ·  {c['severity']}  ·  {c['title']}", 1)
        render_comment(b, c)
        b.h("Response", 2)
        ANSWERS[c["id"]](b, ctx)

    b.page_break()
    b.h("Recommended order of resolution, with status", 1)
    b.table("The reviewer's order, unchanged, with what this round did.", ["#", "Comment", "Next action (reviewer)", "Status"], [
        ["1", "R2-2", "Audit the true evaluation sampling frame using stable episode identity.", "Done; wording supplied"],
        ["2", "R2-1", "Correct or narrow the main-batch vs evaluation-branch interpretation.", "Done; same-universe analysis + wording"],
        ["3", "R2-3", "Recompute MARS uncertainty with child-cluster paired bootstrap.", "Done; all 4 fields, both comparators"],
        ["4", "R2-4", "Rename/clarify the abstention metric; add overall model-null frequency.", "Done"],
        ["5", "R2-5", "Justify the eight-turn cap or add the limitation.", "Done; sensitivity + limitation"],
        ["6", "R2-6", "Define the random and strongest-generator policies precisely.", "Done; baseline corrected"],
        ["7", "P-2 / P-3", "Finish corpus characteristics and supplementary materials.", "Table added; sources complete; S-file assembly by authors"],
        ["8", "P-1", "Complete the manuscript declarations.", "Author action"],
    ], [0.3, 0.8, 3.6, 2.0])
    b.h("How this document was produced", 1)
    b.code(["uv run python -m tests.eval.extraction.round4_analysis   # no model calls; needs exp-data/ locally",
            "uv run python docs/paper/make_round4_response.py"])
    b.doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    print(build())
