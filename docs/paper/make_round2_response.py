"""Build the response to the round-2 manuscript-finalization questions.

Reads the questions verbatim from
``Round 2- Clarification Questions for Manuscript Finalization.docx`` (same
order, same headings, same bullet hierarchy) and answers each one underneath.
Every metric is read from the recorded artefact named beside it; every date is
either an artefact's own ``generated_at`` or a commit timestamp, and the table
of evidence says which.

    uv run python docs/paper/make_round2_response.py
"""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch  # noqa: E402

import docx  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
PAPER = ROOT / "docs" / "paper"
RESULTS = ROOT / "tests" / "eval" / "extraction" / "results"
QUESTIONS = PAPER / "Round 2- Clarification Questions for Manuscript Finalization.docx"
FIG_DIR = PAPER / "figures" / "round2"
OUT = PAPER / "Round2_Clarification_Response.docx"

# Categorical slots 1-4 of the validated reference palette (CVD-checked);
# every mark is also direct-labelled, so identity is never colour-alone.
BLUE, ORANGE, AQUA, YELLOW = "#2a78d6", "#eb6834", "#1baf7a", "#eda100"
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"
QUOTE = RGBColor(0x2A, 0x4E, 0x7A)
HEAD_FILL = "E8EEF6"

plt.rcParams.update({
    "font.size": 9, "axes.edgecolor": INK2, "axes.labelcolor": INK,
    "text.color": INK, "xtick.color": INK2, "ytick.color": INK2,
    "axes.spines.top": False, "axes.spines.right": False,
    "savefig.dpi": 200, "savefig.bbox": "tight", "figure.facecolor": "white",
})

MODELS = ["llama3.2:latest", "qwen2.5:3b", "qwen2.5:7b", "qwen3:14b"]
COLOUR = dict(zip(MODELS, [BLUE, ORANGE, AQUA, YELLOW]))


def short(m: str) -> str:
    return m.replace(":latest", "")


def pct(x: float, d: int = 1) -> str:
    return f"{100 * x:.{d}f}%"


# ---------------------------------------------------------------------------
# Facts from artefacts
# ---------------------------------------------------------------------------

def load_facts() -> dict[str, Any]:
    hm = json.loads((RESULTS / "holdout_metrics.json").read_text())
    hs = json.loads((RESULTS / "holdout_set.json").read_text())
    facts: dict[str, Any] = {"hm": hm, "holdout_set": hs}
    for side in ("batch", "holdout"):
        rows = {}
        for m in MODELS:
            v = hm[side]["per_model"][m]
            rows[m] = {
                "merged": v["mean_field_accuracy"],
                "abstain": v["mean_abstention_rate"],
                "uns_n": v["unsupported_n"],
                "uns_d": v["unsupported_opportunities"],
                "uns": v["unsupported_rate"],
                "whole": hm[side]["whole_record"]["per_generator"][m],
            }
        facts[side] = rows
        facts[f"n_{side}"] = hm[side]["n_segments"]

    def ts(name: str, key: str = "generated_at") -> str:
        return json.loads((RESULTS / name).read_text())[key]

    facts["ts"] = {
        "batch_frozen": ts("annotation_batch.json"),
        "batch_scored": ts("scored_against_gold.json"),
        "mars_summary": ts("mars_run_summary.json"),
        "qwen3_added": ts("annotation_batch_augmented.json", "augmented_at"),
        "qwen3_scored": ts("scored_against_gold_augmented.json"),
        "holdout_frozen": hs["generated_at"],
        "holdout_generated": ts("holdout_batch.json"),
        "holdout_scored": ts("scored_holdout.json"),
        "holdout_metrics": hm["generated_at"],
    }

    def gold_quotes(name: str) -> tuple[int, int]:
        gold = json.loads((RESULTS / name).read_text())["gold"]
        beh = [str(v.get("field_values", {}).get("behavior") or "") for v in gold.values()]
        return sum(1 for b in beh if "'" in b or '"' in b), len(gold)

    facts["quotes_batch"] = gold_quotes("gold_labels.json")
    facts["quotes_holdout"] = gold_quotes("gold_holdout_labels.json")
    return facts


def utc(iso: str) -> datetime:
    return datetime.fromisoformat(iso).astimezone(timezone.utc)


def fmt_utc(iso: str) -> str:
    return utc(iso).strftime("%Y-%m-%d %H:%M UTC")


# Commit timestamps (git log --date=iso), converted to UTC. A commit time is an
# upper bound on when its content existed, never a lower bound.
COMMITS = {
    "fa11861": ("2026-08-28T09:54:06-04:00", "Application default extractor set to llama3.2:latest (src/backend/config.py)"),
    "f9a53c6": ("2026-08-26T13:17:00-03:00", "Commit carrying the only existing tag, research-prototype-v1.0"),
    "04a2b36": ("2026-09-07T12:11:10-04:00", "Best-single-generator rule written: argmax over generators (mars_evaluation.py)"),
    "7a4b725": ("2026-09-16T11:41:38-04:00", "qwen2.5:7b named 'strongest generator'; DEFAULT_MODEL of the reporting harness"),
    "f404ca0": ("2026-09-18T11:57:50-03:00", "Held-out double annotation completed"),
    "63a73a6": ("2026-09-21T17:55:27-04:00", "VERIFIED status, analytics gate, deterministic reporter"),
    "7d2e744": ("2026-09-21T23:57:02-03:00", "Held-out adjudication completed"),
}


def commit_utc(h: str) -> str:
    return COMMITS[h][0]


# ---------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------

def fig_chronology(f: dict[str, Any]) -> Path:
    t = f["ts"]
    SEL, MAIN, HOLD, GOV = ORANGE, BLUE, AQUA, INK2
    events = [
        (commit_utc("04a2b36"), SEL, "Selection rule written: 'best single generator' = argmax agreement with gold",
         "commit 04a2b36 · before any gold label existed"),
        (t["batch_frozen"], MAIN, "Main batch frozen: 330 segments, 3 generators",
         "annotation_batch.json"),
        (t["batch_scored"], MAIN, "Main batch scored against adjudicated gold → qwen2.5:7b ranks first",
         "scored_against_gold.json"),
        (t["mars_summary"], MAIN, "MARS run summarised (qwen2.5:7b used as the comparator)",
         "mars_run_summary.json"),
        (commit_utc("7a4b725"), SEL, "qwen2.5:7b named 'strongest generator' in manuscript code and reporting harness",
         "commit 7a4b725"),
        (t["holdout_frozen"], HOLD, "Held-out set frozen: 132 segments, seeded, no model output consulted",
         "holdout_set.json"),
        (t["qwen3_scored"], MAIN, "qwen3:14b added and scored on main batch: ranks second, order unchanged",
         "scored_against_gold_augmented.json"),
        (commit_utc("f404ca0"), HOLD, "Held-out double annotation completed", "commit f404ca0"),
        (t["holdout_generated"], HOLD, "Four generators run over held-out set (outputs not yet scorable)",
         "holdout_batch.json"),
        (commit_utc("63a73a6"), GOV, "Verification gate and deterministic reporter implemented", "commit 63a73a6"),
        (commit_utc("7d2e744"), HOLD, "Held-out adjudication completed: reference standard exists",
         "commit 7d2e744"),
        (t["holdout_scored"], HOLD, "First held-out generator performance number",
         "scored_holdout.json"),
        (t["holdout_metrics"], HOLD, "Batch vs held-out comparison", "holdout_metrics.json"),
    ]
    events.sort(key=lambda e: utc(e[0]))
    n = len(events)
    fig, ax = plt.subplots(figsize=(8.6, 0.5 * n + 1.0))
    ax.set_xlim(0, 10)
    ax.set_ylim(n - 0.3, -1.4)
    ax.axis("off")
    ax.plot([2.55, 2.55], [-0.2, n - 0.6], color=GRID, lw=2, zorder=1)

    first_holdout_number = next(i for i, e in enumerate(events) if "First held-out" in e[2])
    ax.axhspan(-0.5, first_holdout_number - 0.5, xmin=0, xmax=1, color="#f4f7fb", zorder=0)
    ax.axhline(first_holdout_number - 0.5, color=INK2, lw=0.9, ls=(0, (4, 3)))
    ax.text(9.95, first_holdout_number - 0.62,
            "No held-out performance result exists above this line",
            ha="right", va="bottom", fontsize=7.8, color=INK2, style="italic")

    for i, (iso, col, title, src) in enumerate(events):
        d = utc(iso)
        ax.text(2.3, i, d.strftime("%d %b  %H:%M"), ha="right", va="center", fontsize=8.2, color=INK2)
        ax.scatter([2.55], [i], s=64, color=col, edgecolor="white", linewidth=1.8, zorder=3)
        weight = "bold" if col == SEL else "normal"
        ax.text(2.8, i - 0.13, title, ha="left", va="center", fontsize=8.4, weight=weight)
        ax.text(2.8, i + 0.25, src, ha="left", va="center", fontsize=7, color=INK2, family="monospace")

    for k, (lab, col) in enumerate([("Selection of qwen2.5:7b", SEL), ("Main batch (330)", MAIN),
                                    ("Held-out set (132)", HOLD), ("Governance change", GOV)]):
        x = 0.05 + k * 2.5
        ax.scatter([x], [-1.0], s=42, color=col, edgecolor="white", linewidth=1.2)
        ax.text(x + 0.18, -1.0, lab, va="center", fontsize=8)
    ax.text(0.0, n - 0.35, "All times UTC, September 2026 unless shown. Artefact rows use the file's own generated_at; "
            "commit rows use the commit time, an upper bound on when the content existed.",
            fontsize=7, color=INK2, va="top")
    out = FIG_DIR / "r2_fig1_chronology.png"
    fig.savefig(out)
    plt.close(fig)
    return out


def fig_tradeoff(f: dict[str, Any]) -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(9.2, 3.9), gridspec_kw={"width_ratios": [1.35, 1]})
    ax = axes[0]
    for m in MODELS:
        b, h = f["batch"][m], f["holdout"][m]
        x0, y0, x1, y1 = 100 * b["merged"], 100 * b["uns"], 100 * h["merged"], 100 * h["uns"]
        ax.add_patch(FancyArrowPatch((x0, y0), (x1, y1), arrowstyle="-|>", mutation_scale=10,
                                     color=COLOUR[m], lw=1.6, zorder=2, shrinkA=5, shrinkB=5))
        ax.scatter([x0], [y0], s=70, facecolor="white", edgecolor=COLOUR[m], linewidth=2, zorder=3)
        ax.scatter([x1], [y1], s=70, color=COLOUR[m], edgecolor="white", linewidth=1.5, zorder=4)
        dx, dy, ha = {"qwen3:14b": (6, -7, "left"), "llama3.2:latest": (6, 6, "left"),
                      "qwen2.5:7b": (0, -16, "center"), "qwen2.5:3b": (4, -16, "left")}[m]
        ax.annotate(short(m), (x1, y1), xytext=(dx, dy), textcoords="offset points",
                    fontsize=8.5, weight="bold", color=INK, ha=ha)
    ax.set_xlim(8, 50)
    ax.set_ylim(0, 105)
    ax.set_xlabel("Merged mean agreement, 4 categorical fields (%)  →  better")
    ax.set_ylabel("Unsupported assertion on gold-null opportunities (%)\n↓ safer")
    ax.grid(color=GRID, lw=0.6)
    ax.set_axisbelow(True)
    ax.scatter([], [], s=50, facecolor="white", edgecolor=INK2, linewidth=1.6, label=f"main batch (n={f['n_batch']})")
    ax.scatter([], [], s=50, color=INK2, label=f"held-out (n={f['n_holdout']})")
    ax.legend(loc="lower left", fontsize=7.8, frameon=False)
    ax.set_title("A. Agreement against unsupported assertion", loc="left", fontsize=9.5, weight="bold")

    ax = axes[1]
    import numpy as np
    y = np.arange(len(MODELS))
    hgt = 0.36
    for k, (side, face) in enumerate([("batch", False), ("holdout", True)]):
        vals = [100 * f[side][m]["whole"] for m in MODELS]
        for i, m in enumerate(MODELS):
            yy = y[i] + (k - 0.5) * (hgt + 0.04)
            ax.barh(yy, vals[i], height=hgt, color=COLOUR[m] if face else "white",
                    edgecolor=COLOUR[m], linewidth=1.6)
            ax.text(vals[i] + 0.12, yy, f"{vals[i]:.1f}", va="center", fontsize=7.6, color=INK2)
    ax.set_yticks(y, [short(m) for m in MODELS])
    ax.invert_yaxis()
    ax.set_xlim(0, 6.2)
    ax.set_xlabel("Whole-record agreement (%)")
    ax.grid(axis="x", color=GRID, lw=0.6)
    ax.set_axisbelow(True)
    ax.set_title("B. Whole-record agreement\n(outline = batch, filled = held-out)", loc="left",
                 fontsize=9.5, weight="bold")
    fig.tight_layout()
    out = FIG_DIR / "r2_fig2_tradeoff.png"
    fig.savefig(out)
    plt.close(fig)
    return out


def _box(ax, x, y, w, h, title, lines, col, title_col="white"):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                                linewidth=1.4, edgecolor=col, facecolor="white"))
    ax.add_patch(FancyBboxPatch((x, y + h - 0.55), w, 0.55, boxstyle="round,pad=0.02,rounding_size=0.08",
                                linewidth=0, facecolor=col))
    ax.text(x + 0.15, y + h - 0.28, title, color=title_col, weight="bold", fontsize=9, va="center")
    for i, ln in enumerate(lines):
        ax.text(x + 0.15, y + h - 0.85 - i * 0.36, ln, fontsize=7.7, va="center", color=INK)


def fig_release_tiers() -> Path:
    fig, ax = plt.subplots(figsize=(9.6, 5.0))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 6.6)
    ax.axis("off")
    _box(ax, 0.1, 0.35, 3.8, 6.1, "Tier A · Public (DOI)", [
        "Source code (after fixture audit)",
        "Extraction schema + prompts (v1.0.0)",
        "Codebook (examples paraphrased)",
        "Adequacy-router rules + thresholds",
        "Deterministic analytics definitions",
        "Reporting schema + det. reporter",
        "Evaluation + figure scripts",
        "MARS code and run configuration",
        "Manifest + SHA-256 of every file",
        "Aggregate metrics (reliability, agreement,",
        "   unsupported assertion, MARS, held-out)",
        "Categorical labels, opaque segment IDs",
        "Per-segment scoring verdicts",
        "Synthetic example segments",
    ], AQUA)
    _box(ax, 4.1, 0.35, 3.8, 6.1, "Tier B · On request only", [
        "Real segment IDs",
        "Free-text behavior labels",
        f"   (quote the child's words)",
        "Adjudication rationales",
        "Frozen generator outputs",
        "Returned annotation sheets",
        "",
        "Condition: requester holds current",
        "ASDBank password access.",
        "TalkBank rule permits sharing only with",
        "people who already have password access.",
        "The authors hold ASDBank credentials.",
    ], YELLOW, title_col=INK)
    _box(ax, 8.1, 0.35, 3.8, 6.1, "Tier C · Never posted", [
        "Raw CHAT transcripts",
        "De-identified segment text",
        "pool / candidates / batch JSONs",
        "holdout_set.json, holdout_batch.json",
        "Any file with *CHI: / *MOT: lines",
        "",
        "Reason: ASDBank is password-protected;",
        "protected data cannot be posted on",
        "other servers. De-identification",
        "does not change an access rule.",
        "Authorized users regenerate it",
        "from their own ASDBank copy.",
    ], ORANGE)
    out = FIG_DIR / "r2_fig3_release_tiers.png"
    fig.savefig(out)
    plt.close(fig)
    return out


def fig_repro_path() -> Path:
    fig, ax = plt.subplots(figsize=(9.6, 3.6))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 4.4)
    ax.axis("off")

    def node(x, y, w, h, text, col, fill=False):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.1",
                                    linewidth=1.4, edgecolor=col, facecolor=col if fill else "white"))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=7.8,
                color="white" if fill else INK, wrap=True)

    def arrow(x0, y0, x1, y1, col=INK2, ls="-"):
        ax.add_patch(FancyArrowPatch((x0, y0), (x1, y1), arrowstyle="-|>", mutation_scale=11,
                                     color=col, lw=1.3, linestyle=ls))

    ax.text(0.05, 4.2, "Public path (anyone)", fontsize=8.5, weight="bold", color=AQUA)
    node(0.05, 2.85, 2.2, 1.0, "Clean public repo\n(no Tier B/C files\nin any commit)", AQUA)
    node(3.0, 2.85, 2.2, 1.0, "Annotated release tag\ne.g. manuscript-v1.0", AQUA)
    node(5.95, 2.85, 2.2, 1.0, "Zenodo archive\nversion DOI +\nconcept DOI", AQUA, fill=True)
    node(8.9, 2.85, 3.0, 1.0, "Readers re-run every table\nand figure from aggregate\nartefacts + verdicts", AQUA)
    arrow(2.25, 3.35, 3.0, 3.35)
    arrow(5.2, 3.35, 5.95, 3.35)
    arrow(8.15, 3.35, 8.9, 3.35)

    ax.text(0.05, 2.15, "Authorized path (ASDBank account holders)", fontsize=8.5, weight="bold", color=ORANGE)
    node(0.05, 0.6, 2.2, 1.1, "Own ASDBank\ndownload\n(Rollins, Flusberg)", ORANGE)
    node(3.0, 0.6, 2.2, 1.1, "Tier A pipeline:\nparse → de-identify →\nsegment → select", ORANGE)
    node(5.95, 0.6, 2.2, 1.1, "Check SHA-256 against\npublished manifest", ORANGE, fill=True)
    node(8.9, 0.6, 3.0, 1.1, "Match: identical 330 + 132\nsegments. Tier B on request\nfor frozen outputs + text labels", ORANGE)
    arrow(2.25, 1.15, 3.0, 1.15)
    arrow(5.2, 1.15, 5.95, 1.15)
    arrow(8.15, 1.15, 8.9, 1.15)
    arrow(7.05, 2.85, 7.05, 1.7, col=INK2, ls="--")
    ax.text(7.15, 2.27, "hashes published\nin the archive", fontsize=7, color=INK2, va="center")
    out = FIG_DIR / "r2_fig4_repro_path.png"
    fig.savefig(out)
    plt.close(fig)
    return out


# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------

class Doc:
    def __init__(self) -> None:
        self.doc = docx.Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Inches(0.9))
        sec.top_margin = sec.bottom_margin = Inches(0.85)
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10.5)
        self.tables = 0
        self.figures = 0

    def _runs(self, para, text: str, *, italic=False, colour=None, bold=False, size=None) -> None:
        # `backticks` render as code; **double stars** render bold.
        import re
        for chunk in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
            if not chunk:
                continue
            if chunk.startswith("`"):
                run = para.add_run(chunk[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9.2)
            elif chunk.startswith("**"):
                run = para.add_run(chunk[2:-2])
                run.bold = True
            else:
                run = para.add_run(chunk)
                run.bold = bold or None
            run.italic = italic or None
            if size:
                run.font.size = Pt(size)
            if colour is not None:
                run.font.color.rgb = colour

    def h(self, text: str, level: int) -> None:
        self.doc.add_heading(text, level=level)

    def p(self, text: str, **kw) -> None:
        self._runs(self.doc.add_paragraph(), text, **kw)

    def label(self, text: str) -> None:
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(8)
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x33)

    def sub(self, question: str) -> None:
        """Heading for one sub-question: its own text, verbatim, in the question colour."""
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.keep_with_next = True
        _shade_para(para, "EEF3F9")
        r = para.add_run("Q  ")
        r.bold = True
        r.font.color.rgb = QUOTE
        r2 = para.add_run(question.strip())
        r2.bold = True
        r2.font.color.rgb = QUOTE

    def bullets(self, items: list[str], level: int = 1) -> None:
        style = "List Bullet" if level == 1 else "List Bullet 2"
        for it in items:
            self._runs(self.doc.add_paragraph(style=style), it)

    def answer_line(self, text: str) -> None:
        para = self.doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(10)
        _left_border(para, "1BAF7A")
        self._runs(para, text)

    def boxed(self, title: str, lines: list[str], fill: str = "FFF6E5", edge: str = "EDA100") -> None:
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = Inches(6.7)
        _cell_fill(cell, fill)
        _cell_border(cell, edge)
        cell.paragraphs[0].text = ""
        self._runs(cell.paragraphs[0], title, bold=True)
        for ln in lines:
            self._runs(cell.add_paragraph(), ln)
        self.doc.add_paragraph()

    def code(self, lines: list[str]) -> None:
        para = self.doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(14)
        for i, ln in enumerate(lines):
            run = para.add_run(ln)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            if i < len(lines) - 1:
                run.add_break()

    def table(self, caption: str, header: list[str], rows: list[list[str]], widths: list[float]) -> None:
        self.tables += 1
        cap = self.doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(6)
        r = cap.add_run(f"Table {self.tables}. ")
        r.bold = True
        r.font.size = Pt(9.5)
        self._runs(cap, caption, size=9.5)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htext in enumerate(header):
            c = t.rows[0].cells[i]
            c.paragraphs[0].text = ""
            self._runs(c.paragraphs[0], htext, bold=True, size=9)
            _cell_fill(c, HEAD_FILL)
        _repeat_header(t.rows[0])
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].paragraphs[0].text = ""
                self._runs(cells[i].paragraphs[0], val, size=9)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def figure(self, path: Path, caption: str, width: float = 6.7) -> None:
        self.figures += 1
        self.doc.add_picture(str(path), width=Inches(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.doc.add_paragraph()
        r = cap.add_run(f"Figure {self.figures}. ")
        r.bold = True
        r.font.size = Pt(9.5)
        self._runs(cap, caption, size=9.5, italic=False)

    def page_break(self) -> None:
        self.doc.add_page_break()


def _shade_para(para, fill: str) -> None:
    ppr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _left_border(para, colour: str) -> None:
    ppr = para._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "18"), ("space", "8"), ("color", colour)):
        left.set(qn(f"w:{k}"), v)
    bdr.append(left)
    ppr.append(bdr)


def _cell_fill(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _cell_border(cell, colour: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        for k, v in (("val", "single"), ("sz", "10"), ("space", "0"), ("color", colour)):
            el.set(qn(f"w:{k}"), v)
        borders.append(el)
    tcpr.append(borders)


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


# ---------------------------------------------------------------------------
# The questions, verbatim
# ---------------------------------------------------------------------------

def read_questions() -> list[dict[str, Any]]:
    """Split the round-2 document into its sections, keeping every paragraph and its bullet level."""
    src = docx.Document(str(QUESTIONS))
    sections: list[dict[str, Any]] = []
    title = None
    for para in src.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        ppr = para._p.pPr
        num = ppr.numPr if ppr is not None else None
        level = None
        if num is not None and num.ilvl is not None:
            level = int(num.ilvl.val)
        bold = any(r.bold for r in para.runs)
        if title is None:
            title = text
            continue
        is_heading = level is None and bold and (
            text[:2] in ("1.", "2.") or text.startswith("Optional confirmation"))
        if is_heading:
            sections.append({"heading": text, "items": []})
            continue
        sections[-1]["items"].append({"text": para.text, "level": level, "bold": bold})
    return [{"title": title}] + sections


def render_question(b: Doc, items: list[dict[str, Any]]) -> None:
    b.label("The question, as asked")
    for it in items:
        lines = [ln for ln in it["text"].split("\n") if ln.strip()]
        if it["level"] is None:
            for k, ln in enumerate(lines):
                para = b.doc.add_paragraph()
                para.paragraph_format.left_indent = Pt(12)
                para.paragraph_format.space_after = Pt(3)
                is_label = ln.rstrip().endswith(":") and k == 0 and len(ln) < 40
                b._runs(para, ln.strip(), colour=QUOTE, bold=is_label or (it["bold"] and len(lines) == 1))
        else:
            style = "List Bullet" if it["level"] == 0 else "List Bullet 2"
            for ln in lines:
                para = b.doc.add_paragraph(style=style)
                b._runs(para, ln.strip(), colour=QUOTE, bold=it["bold"])


def sub_questions(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Top-level bullets with their children — the unit each answer addresses."""
    out: list[dict[str, Any]] = []
    for it in items:
        if it["level"] == 0:
            out.append({"text": it["text"].strip(), "children": []})
        elif it["level"] == 1 and out:
            out[-1]["children"].append(it["text"].strip())
    return out


# ---------------------------------------------------------------------------
# Section 1
# ---------------------------------------------------------------------------

def answer_q1(b: Doc, f: dict[str, Any], subs: list[dict[str, Any]]) -> None:
    B, H, t = f["batch"], f["holdout"], f["ts"]
    q7, q14, q3, ll = "qwen2.5:7b", "qwen3:14b", "qwen2.5:3b", "llama3.2:latest"
    gap = 100 * (H[q14]["merged"] - H[q7]["merged"])

    b.label("Short answer")
    b.answer_line(
        "**Yes, it was identified before the held-out set existed, and the held-out results played no "
        "part in it.** qwen2.5:7b ranked first on the 330-segment main batch when that batch was scored, "
        f"on {fmt_utc(t['batch_scored'])}. The ranking came from a rule written on 2026-09-07, before any "
        "gold label existed. qwen2.5:7b was named in committed code on "
        f"{fmt_utc(commit_utc('7a4b725'))}. The held-out set was frozen about 1.5 hours after that, and its "
        f"first performance number did not appear until {fmt_utc(t['holdout_scored'])}, six days later.")
    b.answer_line(
        "**One premise needs correcting before this goes into the manuscript: qwen2.5:7b was never made the "
        "application's operational extractor.** The running application's default is `llama3.2:latest` "
        "(`src/backend/config.py:27`, unchanged since commit fa11861 on 2026-08-28). What qwen2.5:7b holds is "
        "three research roles: best single generator on the main batch, the comparator MARS is measured "
        "against, and the model used in the reporting-layer experiment. The manuscript should say that, not "
        "\"operational extractor\". Options are given at the end of this section.")

    # --- bullet 1
    b.sub(subs[0]["text"])
    b.p("Yes. The ordering below is taken from artefact timestamps and commit times, not from memory.")
    b.table(
        "When qwen2.5:7b was identified, compared with every held-out milestone. All times UTC.",
        ["Event", "When", "Source"],
        [
            ["Best-generator rule written (argmax over generators)", fmt_utc(commit_utc("04a2b36")), "commit 04a2b36, `mars_evaluation.py`"],
            ["Main batch scored; **qwen2.5:7b ranks first**", fmt_utc(t["batch_scored"]), "`scored_against_gold.json` generated_at"],
            ["qwen2.5:7b named 'strongest generator' in code", fmt_utc(commit_utc("7a4b725")), "commit 7a4b725"],
            ["Held-out set frozen (did not exist before this)", fmt_utc(t["holdout_frozen"]), "`holdout_set.json` generated_at"],
            ["Held-out double annotation completed", fmt_utc(commit_utc("f404ca0")), "commit f404ca0"],
            ["Held-out generator outputs produced (unscorable: no gold)", fmt_utc(t["holdout_generated"]), "`holdout_batch.json` generated_at"],
            ["Held-out adjudication completed", fmt_utc(commit_utc("7d2e744")), "commit 7d2e744"],
            ["**First held-out performance number**", fmt_utc(t["holdout_scored"]), "`scored_holdout.json` generated_at"],
        ],
        [3.1, 1.6, 2.0],
    )
    b.figure(FIG_DIR / "r2_fig1_chronology.png",
             "The chronology. qwen2.5:7b (orange, bold) was identified and named before the held-out set was "
             "frozen. Everything above the dashed line happened before any held-out performance result existed.")

    # --- bullet 2
    b.sub(subs[1]["text"])
    b.p("\"At the time of selection\" is dated two ways. The first is when the ranking was first computable "
        f"({fmt_utc(t['batch_scored'])}). The second is when qwen2.5:7b was first named in committed code "
        f"({fmt_utc(commit_utc('7a4b725'))}). The table uses both.")
    b.table(
        "Which results existed when qwen2.5:7b was selected.",
        ["Result (as listed in the question)", "Available?", "First timestamp", "Note"],
        [
            [subs[1]["children"][0], "**Yes**", fmt_utc(t["batch_scored"]),
             "Three-generator roster at the time: llama3.2, qwen2.5:3b, qwen2.5:7b"],
            [subs[1]["children"][1], "At naming only", fmt_utc(t["mars_summary"]),
             "After the ranking, before the naming commit. MARS uses the best generator as its comparator. It "
             "cannot change which generator is best"],
            [subs[1]["children"][2], "**No**", fmt_utc(t["qwen3_scored"]),
             f"Added afterwards. Ranked second on the main batch ({pct(B[q14]['merged'])} vs "
             f"{pct(B[q7]['merged'])}), so the ranking did not change"],
            [subs[1]["children"][3], "**No**", fmt_utc(commit_utc("f404ca0")),
             "The held-out set itself did not exist until "
             f"{fmt_utc(t['holdout_frozen'])}"],
            [subs[1]["children"][4], "**No**", fmt_utc(t["holdout_scored"]),
             "Needs held-out gold, which was completed on 2026-09-22"],
        ],
        [1.9, 0.9, 1.35, 2.55],
    )

    # --- bullet 3
    b.sub(subs[2]["text"])
    b.p("The rule is one sentence: **use the single generator with the highest agreement against the "
        "adjudicated main-batch reference standard.** It exists in code in two forms:")
    b.bullets([
        "**Original form, written 2026-09-07 (commit 04a2b36, before any gold label existed).** "
        "`mars_evaluation.build_evaluation()` takes the argmax of each generator's share of fully correct "
        "records, i.e. whole-record agreement.",
        "**Later form, 2026-09-16 (commit 7a4b725).** `mars_selection_analysis.py:224` takes the argmax of "
        "mean agreement over the four categorical fields, i.e. merged mean agreement.",
    ])
    b.p("Both forms pick qwen2.5:7b on the main batch, so the choice does not depend on which form is used. "
        "Each criterion listed in the question is answered below.")
    rows = []
    crit = subs[2]["children"]
    rows.append([crit[0], "**Yes: the criterion**",
                 f"qwen2.5:7b {pct(B[q7]['merged'])}; llama3.2 {pct(B[ll]['merged'])}; qwen2.5:3b "
                 f"{pct(B[q3]['merged'])}. qwen3:14b, added later, {pct(B[q14]['merged'])}"])
    rows.append([crit[1], "**Yes: the original form of the criterion**",
                 f"qwen2.5:7b {pct(B[q7]['whole'])}; llama3.2 {pct(B[ll]['whole'])}; qwen2.5:3b "
                 f"{pct(B[q3]['whole'])}. qwen3:14b {pct(B[q14]['whole'])}"])
    rows.append([crit[2], "**No**",
                 f"Not a criterion, and it could not have been one: qwen2.5:3b has the lowest rate "
                 f"({pct(B[q3]['uns'])} vs {pct(B[q7]['uns'])}). A safety criterion would have chosen qwen2.5:3b"])
    rows.append([crit[3], "**No**",
                 f"Not a criterion. Mean abstention: qwen2.5:7b {pct(B[q7]['abstain'])}, qwen3:14b "
                 f"{pct(B[q14]['abstain'])}, qwen2.5:3b {pct(B[q3]['abstain'])}"])
    rows.append([crit[4], "**No: this is a result, not a criterion**",
                 "The best single generator is the baseline MARS is tested against. 'MARS did not beat always "
                 "using the strongest generator' is a finding reported after the fact. It was not a reason "
                 "for the choice"])
    rows.append([crit[5], "None recorded",
                 "No design note gives any further reason. The reporting harness uses the same model "
                 "(`DEFAULT_MODEL = \"qwen2.5:7b\"`, `tests/eval/reporting/groundedness.py:67`) without a "
                 "written rationale. The authors should confirm it was simply 'reuse the strongest generator'"])
    b.table("Each candidate criterion in the question, and whether it was used. Main batch, n = 330. "
            "Source: `holdout_metrics.json` (batch side).",
            ["Criterion", "Used?", "Evidence"], rows, [1.7, 1.3, 3.7])
    b.boxed("A correction to the question's own framing", [
        "The question's rationale says qwen2.5:7b \"has a substantially safer unsupported-assertion "
        "profile\". That holds against qwen3:14b "
        f"({pct(B[q14]['uns'])} batch, {pct(H[q14]['uns'])} held-out) and llama3.2 "
        f"({pct(B[ll]['uns'])}, {pct(H[ll]['uns'])}). It does not hold against qwen2.5:3b "
        f"({pct(B[q3]['uns'])}, {pct(H[q3]['uns'])}), which is the safest of the four on both sets.",
        "So the manuscript must not say qwen2.5:7b was chosen for safety. The accurate statement is that "
        "it was chosen on agreement, and that among the two highest-agreement generators it asserts "
        "unsupported values far less often.",
    ])

    # --- bullet 4
    b.sub(subs[3]["text"])
    b.p("Yes for the selection itself; see the table below. There is **no** dated record designating "
        "qwen2.5:7b as the operational extractor, because no such designation was ever made.")
    b.table(
        "Dated records that support the chronology. A commit time is an upper bound on when its "
        "content existed: several commits in this history are bulk commits.",
        ["Record", "Type", "When (UTC)", "What it shows"],
        [
            ["`mars_evaluation.py` @ 04a2b36", "Commit", fmt_utc(commit_utc("04a2b36")),
             "The best-generator rule existed eight days before any gold label"],
            ["`scored_against_gold.json`", "Run artefact (SHA-256 pinned)", fmt_utc(t["batch_scored"]),
             "The scoring that makes qwen2.5:7b first"],
            ["`mars_selection_analysis.py:224`, `groundedness.py:67` @ 7a4b725", "Commit",
             fmt_utc(commit_utc("7a4b725")),
             "qwen2.5:7b named as the strongest generator and set as the reporting model"],
            ["`holdout_set.json`", "Run artefact (pinned)", fmt_utc(t["holdout_frozen"]),
             "selection_rule: \"No model output was consulted. Frozen before any generator is run over it.\""],
            ["`src/backend/config.py:27` @ fa11861", "Configuration", fmt_utc(COMMITS["fa11861"][0]),
             "The application default is `llama3.2:latest`, not qwen2.5:7b"],
        ],
        [1.95, 1.15, 1.35, 2.25],
    )
    b.p("The pinned artefacts are covered by `docs/paper/artifact_manifest.json`. Anyone can confirm "
        "the timestamps are unaltered by running `uv run python docs/paper/verify_artifacts.py`.",
        italic=True)

    # --- bullet 5
    b.sub(subs[4]["text"])
    b.p("**Confirmed.** The held-out set was used only for later validation. Four independent facts support "
        "this:")
    b.bullets([
        "**Order of events.** The selection rule (2026-09-07) and its outcome (2026-09-15/16) both predate "
        "the held-out set (frozen 2026-09-16 17:11 UTC) and its first result (2026-09-22).",
        "**How the held-out set was drawn.** It was stratified by child, seeded (20260916) and drawn from "
        "segments outside the annotation pool. No model output was consulted, so the held-out data cannot "
        "encode a preference for any generator.",
        "**Nothing moved after the held-out results.** `holdout_metrics.py` recomputes a held-out 'best "
        "generator' figure and reports it. Nothing reads that figure back: no configuration, prompt or "
        "manuscript role changed after 2026-09-22.",
        "**The held-out result does not uniformly favour qwen2.5:7b.** On merged agreement qwen3:14b is "
        f"ahead ({pct(H[q14]['merged'])} vs {pct(H[q7]['merged'])}). A post-hoc choice would have picked "
        "whichever model looked best on held-out data. The pre-specified choice was kept even though it "
        "no longer leads on that metric.",
    ])

    b.label("What the held-out set says about the choice")
    b.table(
        f"Each criterion on both sets. Winner in bold. Batch n = {f['n_batch']}, held-out n = "
        f"{f['n_holdout']}, same scoring policy (4 categorical fields, behavior excluded). Source: "
        "`holdout_metrics.json`.",
        ["Generator", "Merged, batch", "Merged, held-out", "Whole-record, batch", "Whole-record, held-out",
         "Unsupported, batch", "Unsupported, held-out"],
        [[short(m)] + [
            _win(pct(B[m]["merged"]), m == max(MODELS, key=lambda x: B[x]["merged"])),
            _win(pct(H[m]["merged"]), m == max(MODELS, key=lambda x: H[x]["merged"])),
            _win(pct(B[m]["whole"]), m == max(MODELS, key=lambda x: B[x]["whole"])),
            _win(pct(H[m]["whole"]), m == max(MODELS, key=lambda x: H[x]["whole"])),
            _win(f"{B[m]['uns_n']}/{B[m]['uns_d']} ({pct(B[m]['uns'])})", m == min(MODELS, key=lambda x: B[x]["uns"])),
            _win(f"{H[m]['uns_n']}/{H[m]['uns_d']} ({pct(H[m]['uns'])})", m == min(MODELS, key=lambda x: H[x]["uns"])),
        ] for m in MODELS],
        [1.05, 0.8, 0.9, 0.95, 1.0, 1.0, 1.0],
    )
    b.figure(FIG_DIR / "r2_fig2_tradeoff.png",
             "Batch (outline) against held-out (filled) for all four generators. A: arrows run from batch to "
             "held-out; lower-right is better. qwen3:14b sits highest on unsupported assertion on both sets. "
             "B: qwen2.5:7b has the highest whole-record agreement on both sets.")
    b.p(f"On the held-out set, qwen3:14b leads merged agreement by "
        f"{gap:.1f} points over {f['n_holdout']} segments. qwen2.5:7b "
        f"leads whole-record agreement ({pct(H[q7]['whole'])} vs {pct(H[q14]['whole'])}) and asserts "
        f"unsupported values on {H[q7]['uns_n']} of {H[q7]['uns_d']} gold-null opportunities, against "
        f"{H[q14]['uns_n']} of {H[q14]['uns_d']} for qwen3:14b. No paired significance test has been run "
        f"on the {gap:.1f}-point gap. The manuscript should describe it as a small difference in one direction, "
        "not as a reversal of the ranking.")

    b.label("What the manuscript should say")
    b.boxed("Suggested Methods wording (adjust names to the manuscript's conventions)", [
        "\"Before any reference labels existed, we specified that the single-generator baseline would be the "
        "generator with the highest agreement with the adjudicated reference standard on the annotation "
        "batch. When the batch was scored (15 September 2026), this rule selected qwen2.5:7b under both "
        "whole-record and mean categorical-field agreement. qwen3:14b, added to the roster afterwards, ranked "
        "second. The model-independent held-out set was drawn after this choice, by a seeded rule that "
        "consulted no model output. It was adjudicated and scored six days later and was used only to "
        "validate the choice, never to make it.\"",
        f"\"On the held-out set, qwen3:14b's mean agreement was {gap:.1f} points higher than qwen2.5:7b's. qwen2.5:7b "
        "retained the higher whole-record agreement and a substantially lower unsupported-assertion rate "
        "(32/63 vs 58/63 gold-null opportunities).\"",
        "Do not write \"operational extractor\", and do not write \"selected for safety\".",
    ], fill="EAF7F1", edge="1BAF7A")

    b.label("Decision needed from the authors: what to call qwen2.5:7b")
    b.table("Options for the 'operational extractor' wording.",
            ["Option", "What changes", "Chronology consequence"],
            [
                ["**A (recommended).** Describe qwen2.5:7b as the best-performing generator and the "
                 "evaluated configuration. State that the application's extractor is configurable and "
                 "currently defaults to llama3.2",
                 "Wording only", "None. Every statement is already true and dated"],
                ["B. Make qwen2.5:7b the application default now",
                 "One line in `config.py` plus a dated commit",
                 "The change would be made **after** the held-out results were known. It must be disclosed "
                 "as a post-held-out deployment decision, justified by whole-record agreement and "
                 "unsupported assertion. It cannot be described as the original selection"],
            ],
            [2.6, 1.6, 2.5])


def _win(text: str, win: bool) -> str:
    return f"**{text}**" if win else text


# ---------------------------------------------------------------------------
# Section 2
# ---------------------------------------------------------------------------

TIER_ROWS = [
    ("source code", "A · public",
     "No licence file exists yet. Before release, audit the 13 test and source files that contain CHAT "
     "speaker lines, and the files that cite real segment IDs (e.g. `tests/phase4/test_mars_envelope.py`, "
     "`src/backend/ml/extraction/mars/provider.py`). Keep synthetic fixtures; replace anything that is corpus text"),
    ("schemas", "A · public", "`src/backend/schemas/observation_schema.py`, `EXTRACTION_TOOL_SCHEMA`"),
    ("prompts", "A · public", "`src/backend/ml/extraction/prompts.py` (PROMPT_VERSION 1.0.0). The scan found no transcript excerpt in it"),
    ("annotation codebook", "A · public after edit",
     "`docs/annotation_codebook.md` has three worked examples that cite real segment IDs and quote the "
     "child's words (lines ~100-110). Paraphrase them or replace them with synthetic examples"),
    ("adequacy-router rules", "A · public", "`adequacy_router.py`, `config.py`, `docs/lld/Adequacy_Router.md`"),
    ("deterministic analytics definitions", "A · public", "`src/backend/ml/analytics_engine.py`"),
    ("reporting schema", "A · public", "`src/backend/ml/schemas.py`, deterministic reporter in `reporting_client.py`"),
    ("evaluation scripts", "A · public", "`tests/eval/**`. Runnable end to end only with the corpus, which is the point of Tier B/C"),
    ("MARS code/configuration", "A · public after audit",
     "`src/backend/ml/extraction/mars/`, `mars_run_summary.json`. `docs/lld/Mars_Reranking_Specs.md` cites real segment IDs"),
    ("artifact manifest", "A · public", "Keep an entry for every file, including restricted ones: path, size and hash"),
    ("SHA-256 hashes", "A · public",
     "Safe to publish for restricted files too. A hash reveals nothing about the text, and it lets authorized "
     "users prove their regeneration is identical"),
    ("aggregated metrics/results", "A · public",
     "`paper_metrics.json`, `holdout_metrics.json`, `reliability_*.json`, `round_comparison.json`, "
     "`mars_evaluation.json`, `mars_selection_analysis.json`. The scan found no transcript lines"),
    ("adjudicated annotation labels", "Split: A / B",
     "Categorical fields (domain, modality, prompt_level, outcome) under opaque IDs: public. Free-text "
     "`behavior`: Tier B. {qb} of the main-batch behavior labels quote the child's words. Adjudication "
     "rationales: Tier B"),
    ("held-out labels", "Split: A / B", "As above. {qh} of the held-out behavior labels quote the child's words"),
    ("preserved model outputs", "Split: A / B / C",
     "Per-segment verdicts (`scored_*.json`: correct / incorrect / missing): public. Raw outputs are stored "
     "inside `annotation_batch*.json` and `holdout_batch.json` beside the transcript text: Tier C as files. "
     "Their content without transcript text could go to Tier B"),
    ("de-identified ASDBank-derived segment text", "C · not public",
     "ASDBank is password-protected. De-identification removes identifiers but does not change the "
     "access rule"),
]


def answer_q2(b: Doc, f: dict[str, Any], subs: list[dict[str, Any]]) -> None:
    qb = f"{f['quotes_batch'][0]} of {f['quotes_batch'][1]}"
    qh = f"{f['quotes_holdout'][0]} of {f['quotes_holdout'][1]}"

    b.label("Short answer")
    b.answer_line(
        "**Settled by the authors' decision, which follows TalkBank's terms.** The authors hold ASDBank "
        "credentials, will cite ASDBank, and will not keep corpus data in the repository. ASDBank is password-protected. "
        "TalkBank's rules state that \"password-protected data cannot be posted on other websites or servers or "
        "shared with anyone who does not already have password access\". So no transcript-derived text can go "
        "into a public release, de-identified or not. Code, schemas, prompts, rules, hashes and aggregate "
        "metrics can all be public, and so can categorical labels under opaque IDs. The release can be "
        "frozen with a Zenodo DOI.")
    b.answer_line(
        "**Data handling is confirmed (authors, 2026-09-24).** Corpus data will not be kept in the repository. "
        "One step makes that true for the release: the working history currently tracks the batch JSONs and "
        "the returned annotation and adjudication sheets. Build the release from a copy whose history does "
        "not contain those files. Deleting them in a new commit is not enough, because git keeps them in "
        "earlier commits. The files to exclude are the Tier B and Tier C items in Table 9.")
    b.answer_line(
        "**Please confirm the target journal.** This round names \"AI (MDPI)\". Every earlier document, the "
        "reviewer assessment and the manuscript title all say JMIR AI. The two journals have different "
        "data-availability and AI-disclosure templates.")

    # --- repository
    b.sub(subs[0]["text"] + " — " + subs[0]["children"][0])
    b.p("**What exists:** one GitHub remote, "
        "`adnimbuslab/agentic-autism-intervention-progress-tracking`, which holds the working history. Corpus-derived files are to be "
        "removed from the repository (authors' decision). No Zenodo or OSF record exists yet.")
    b.p("**Recommendation: GitHub and Zenodo together.** Use a new, clean public GitHub repository, or a "
        "public mirror of the current one with rewritten history, for the code and Tier A artefacts. Archive "
        "the tagged release on Zenodo through its GitHub integration. That produces a version DOI and a "
        "concept DOI, and the integration requires a public repository, which is one more reason the release "
        "must be the clean copy. OSF adds nothing here unless the authors want a preregistration-style "
        "project page.")
    b.p("Decision: the authors.", italic=True)

    # --- tag
    b.sub(subs[1]["text"])
    b.p("**What exists:** exactly one tag, `research-prototype-v1.0`, created on commit f9a53c6 on "
        "2026-08-26. That was before any annotation, gold label, generator scoring or held-out work, so it "
        "**must not** be cited as the study release.")
    b.p("**Recommendation:** a new annotated tag, e.g. `manuscript-v1.0.0`, on the commit that builds the "
        "submitted manuscript. Tag revisions after review as `manuscript-v1.1.0` and so on. The Zenodo "
        "concept DOI then resolves to the latest version, and each version DOI stays fixed. Current HEAD is "
        "`2d83b77`. It is not yet the release commit, because the Section 2 release audit and the "
        "[TO COMPLETE] markers are still open.")

    # --- licence
    b.sub(subs[2]["text"])
    b.p("**What exists:** no `LICENSE` file anywhere in the repository.")
    b.table("Recommended licences by artefact type. The choice of code licence depends on institutional IP "
            "policy; check it before choosing.",
            ["Artefact type", "Recommended licence", "Why"],
            [
                ["Source code, scripts", "Apache-2.0 (or MIT)",
                 "Permissive; Apache-2.0 adds an explicit patent grant. Code contains no corpus data once audited"],
                ["Codebook, schemas, prompts, documentation, aggregate metrics, figures", "CC BY 4.0",
                 "Standard for research outputs; allows reuse with attribution"],
                ["Anything derived from ASDBank text (Tier B)", "Not relicensable",
                 "Inherits TalkBank's terms: CC BY-NC-SA, and shared only with password holders. The authors "
                 "cannot grant broader rights than they hold"],
            ],
            [2.2, 1.5, 3.0])

    # --- what can be released
    b.sub(subs[3]["text"])
    b.p("Each item is placed in one of three tiers. **Tier A** is public and DOI-archived. **Tier B** is shared "
        "on request, only with people who already have ASDBank password access. **Tier C** is never posted.")
    b.table("The 16 items in the question, in the question's order.",
            ["Item", "Tier", "Condition and location"],
            [[name, tier, cond.format(qb=qb, qh=qh)] for name, tier, cond in TIER_ROWS],
            [1.6, 1.05, 4.05])
    b.figure(FIG_DIR / "r2_fig3_release_tiers.png",
             "The proposed release boundary. Tier A is the public DOI package. Tier B goes on request only to "
             "people with ASDBank password access. Tier C is never posted, and "
             "authorized users regenerate it from their own corpus copy.")

    # --- TalkBank restrictions
    b.sub(subs[4]["text"])
    b.p("**Yes.** Checked on 2026-09-24 against TalkBank's live pages:")
    b.table("TalkBank / ASDBank terms that bind this release.",
            ["Term (quoted)", "Source", "Consequence here"],
            [
                ["ASDBank requires login: \"For access to ASD materials, click the green Login button… to set up "
                 "your TalkBank account.\"", "talkbank.org/asd/",
                 "ASDBank data is password-protected, not open access"],
                ["\"Password-protected data cannot be posted on other websites or servers or shared with anyone "
                 "who does not already have password access.\"", "talkbank.org/0share/rules.html",
                 "No public release of segment text, annotation sheets that contain transcripts, or raw model "
                 "outputs stored with transcripts. Sharing with existing password holders is allowed (Tier B)"],
                ["\"The use of TalkBank data is governed by the Creative Commons CC BY-NC-SA 3.0 copyright "
                 "license.\" (The ASDBank page says CC BY-NC-SA 4.0.)", "rules.html; asd/",
                 "Non-commercial and share-alike for anything derived. The version number differs between "
                 "the two pages; ask TalkBank which applies"],
                ["\"Users cannot upload TalkBank data to web-based systems unless those systems include "
                 "assurance that they will not keep the data.\"", "rules.html",
                 "Covers the GitHub remote holding transcript text, and any future frontier-API arm"],
                ["\"…the data themselves cannot be included in models.\" / NC excludes \"large language "
                 "models\" in commercial products", "rules.html",
                 "Not relevant to this paper (no training). It **is** relevant to the deferred fine-tuning "
                 "phase and must be resolved before that phase starts"],
                ["\"When using TalkBank corpora in published articles, you need to cite your sources by "
                 "following these rules.\"", "rules.html",
                 "Cite both corpora by DOI (Rollins 10.21415/T5KK53; Flusberg 10.21415/T5459F) plus the "
                 "TalkBank references the rules page lists"],
            ],
            [2.6, 1.2, 2.9])
    b.p("**Model outputs that reproduce source text** fall under the same rule as the text itself. A "
        "behavior string that quotes the child is transcript content wherever it is stored. Categorical "
        "values (e.g. `modality = verbal`) and correct/incorrect verdicts reproduce no source text and "
        "can go in Tier A.")
    b.boxed("Authors' position (confirmed 2026-09-24)", [
        "The authors hold ASDBank credentials and will cite ASDBank and both corpora. No corpus data will be "
        "kept in the repository or the public release. That matches the TalkBank rule quoted above, so no "
        "further permission is needed for the Tier A release. One small open point: the rules page says "
        "CC BY-NC-SA 3.0 and the ASDBank page says 4.0. Cite the version on the page current at submission.",
    ], fill="EAF7F1", edge="1BAF7A")

    # --- substitutes
    b.sub(subs[5]["text"])
    b.p("All of the substitutes listed should be released. Together they let an authorized reader rebuild "
        "the exact dataset and let any reader re-derive every reported number.")
    sub_rows = [
        ["hashes", "Yes", "SHA-256 for every artefact, restricted ones included (the manifest already pins 38 of 61)"],
        ["segment identifiers", "Yes, as opaque IDs",
         "Real IDs embed corpus filenames and child pseudonyms (e.g. `<child>_flusberg_…cha_seg_020`). Publish "
         "`SHA-256(real_id)`. Authorized users compute the same value from their regenerated segments"],
        ["scripts that recreate the dataset for authorized users", "Yes",
         "`data/pipeline/` (parse → de-identify → segment → select), `build_annotation_batch`, "
         "`build_holdout_set` (seed 20260916). Deterministic given the corpus"],
        ["schemas/codebooks", "Yes", "Codebook after the example edit above"],
        ["aggregate metrics", "Yes", "All paper tables and figures regenerate from these plus the verdicts"],
        ["synthetic examples", "Yes, recommended",
         "A small synthetic CHAT set that runs the whole pipeline without the corpus. The test fixtures "
         "are a starting point"],
        ["other permitted artifacts", "Yes",
         "Per-segment verdicts, categorical labels, figures, and the manuscript-consistency tests that tie "
         "each reported number to its artefact"],
    ]
    b.table("Reproducibility substitutes, in the question's order.",
            ["Substitute", "Release?", "Detail"], sub_rows, [1.9, 1.1, 3.7])
    b.p("**One limit, which must be stated rather than engineered around.** Generator outputs cannot be "
        "regenerated bit-for-bit, because no generation seed was set. An authorized user can rebuild the "
        "identical segments and verify them by hash. To get the identical model outputs, they need the "
        "frozen outputs from Tier B. Everything downstream of those outputs is deterministic.")
    b.figure(FIG_DIR / "r2_fig4_repro_path.png",
             "Two reproduction paths. Public: from the DOI archive to every table and figure. Authorized: "
             "from the reader's own ASDBank copy to the identical segment set, checked against the published "
             "hashes.")

    # --- DOI
    b.sub(subs[6]["text"])
    b.p("**Yes.** A Zenodo record is immutable once published, and it receives a version DOI plus a concept "
        "DOI for all versions. Cite the version DOI in the manuscript. Minting it needs three things first: "
        "(1) the clean public repository, (2) the release tag, and (3) the Tier A audit. A DOI can be "
        "reserved on Zenodo before publication, so it can go into the manuscript before the record goes "
        "live.")

    b.label("Draft Data Availability statement")
    b.boxed("Draft", [
        "\"The source transcripts are available to registered users from ASDBank (TalkBank) at "
        "https://doi.org/10.21415/T5KK53 (Rollins) and https://doi.org/10.21415/T5459F (Flusberg). Under "
        "TalkBank's terms they are not redistributed. The code, extraction schema and prompts, annotation "
        "codebook, adequacy-router rules, analytics and reporting definitions, evaluation scripts, "
        "categorical reference labels and per-segment scoring verdicts (keyed by hashed segment "
        "identifiers), aggregate results, and a manifest with SHA-256 hashes of every artefact are "
        "archived at [Zenodo DOI] (release [tag]) under Apache-2.0 (code) and CC BY 4.0 (other materials). "
        "Researchers with ASDBank access can regenerate the exact segment sets with the released scripts "
        "and verify them against the published hashes. Free-text labels and frozen model outputs are "
        "available from the corresponding author to researchers who hold ASDBank access.\"",
    ], fill="EAF7F1", edge="1BAF7A")


# ---------------------------------------------------------------------------
# Optional confirmation
# ---------------------------------------------------------------------------

def answer_optional(b: Doc, subs: list[dict[str, Any]]) -> None:
    b.label("Short answer")
    b.answer_line(
        "**All three changes arrived together in one commit: 63a73a6**, "
        f"{fmt_utc(commit_utc('63a73a6'))} (2026-09-21 17:55:27 −04:00). Full hash "
        "`63a73a61fd96c7c6b5738f682f8c0c6ae86ddb2b`. No later commit changed any of the three. Their tests "
        "pass at current HEAD `2d83b77`: 55 passed in `test_analytics_gate.py`, `test_reporting_client.py` "
        "and `test_structured_records.py`, run on 2026-09-24.")
    rows = [
        [subs[0]["text"], "`src/backend/models/observation.py` (RecordStatus.VERIFIED, verified_by, verified_at); "
         "`alembic/versions/004_add_verified_record_status.py`; `POST /records/{record_id}/verify`",
         "`tests/phase3/test_structured_records.py`"],
        [subs[1]["text"], "`src/backend/services/analytics_service.py`: "
         "`status.in_([VALIDATED, CORRECTED])` → `status.in_([VERIFIED, CORRECTED])`",
         "`tests/phase4/test_analytics_gate.py` (compiles the query; asserts the admitted set)"],
        [subs[2]["text"], "`src/backend/config.py`: `reporting_mode = \"deterministic\"`; "
         "`src/backend/ml/reporting_client.py`: `_build_deterministic_sections`",
         "`tests/phase4/test_reporting_client.py::TestOperationalReportingMode`"],
    ]
    b.table("Where each change was made. All three are in commit 63a73a6.",
            ["Change (as asked)", "Code", "Test that pins it"], rows, [1.8, 3.0, 1.9])
    b.p("For the manuscript: cite the release tag and DOI from Section 2, which will contain `63a73a6`. The "
        "commit hash can go in a footnote or the supplementary artefact list. Note one ordering detail: "
        "the gate was implemented at 21:55 UTC on 2026-09-21, about five hours before held-out adjudication "
        "was completed. It governs analytics, not extraction, so it has no bearing on the Section 1 "
        "chronology.")


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------

def build() -> Path:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    facts = load_facts()
    fig_chronology(facts)
    fig_tradeoff(facts)
    fig_release_tiers()
    fig_repro_path()

    parsed = read_questions()
    title = parsed[0]["title"]
    sections = parsed[1:]

    b = Doc()
    b.doc.add_heading(f"Response to {title}", level=0)
    b.p("Answers the questions in 'Round 2- Clarification Questions for Manuscript Finalization.docx'. That "
        "round followed the strict-reviewer assessment ('JMIR_Strict_Reviewer_Critical_Issues_Prioritized.docx') "
        "and the point-by-point response to it ('Reviewer_Response.docx'). Branch model-finetuning, HEAD "
        "`2d83b77`, built 2026-09-24. Internal working document; not for submission.", colour=RGBColor(0x52, 0x51, 0x4E))
    b.p("Each question is reproduced verbatim, with its original numbering, order and bullet structure, and "
        "the answer follows it. Every metric is read from the recorded artefact named beside it. Every date "
        "is either an artefact's own timestamp or a commit time, and each table says which.",
        colour=RGBColor(0x52, 0x51, 0x4E))

    b.h("Summary", 1)
    b.table("What each question asked, and what the answer changes in the manuscript.",
            ["#", "Question", "Answer", "Manuscript consequence"],
            [
                ["1", "Operational Extractor Selection — Qwen 2.5 7B",
                 "Selected before the held-out set existed, by a rule written before any gold label. The "
                 "held-out set played no part. **But it was never the application's operational extractor**",
                 "Replace \"operational extractor\" with \"best-performing / evaluated generator\". Do not "
                 "claim it was selected for safety"],
                ["2", "Final Reproducibility and Public-Release Plan",
                 "ASDBank is password-protected, so no transcript-derived text can be public. A three-tier "
                 "release with a Zenodo DOI is feasible. Authors hold ASDBank credentials, will cite ASDBank, "
                 "and will keep no corpus data in the repository",
                 "Data Availability statement drafted below. Needs a release built without the corpus "
                 "files in its history, a tag, licences and a DOI"],
                ["Opt.", "Optional confirmation (verification gate)",
                 "All three changes landed in commit 63a73a6 on 2026-09-21; tests pass at HEAD",
                 "Cite via the release tag; commit hash in the supplement"],
            ],
            [0.4, 1.6, 2.5, 2.2])
    b.label("Decisions only the authors can make")
    b.bullets([
        "Remove the corpus-derived files (Table 9, Tiers B and C) and build the release from a history "
        "that never contained them (Section 2).",
        "Choose option A or B for how qwen2.5:7b is described (Section 1). A is recommended.",
        "Confirm the target journal: JMIR AI or AI (MDPI).",
        "Approve the repository and archive plan, the tag name, and the licences (Section 2).",
    ])

    answerers = [answer_q1, answer_q2]
    for i, sec in enumerate(sections):
        b.page_break()
        b.h(sec["heading"], 1)
        render_question(b, sec["items"])
        subs = sub_questions(sec["items"])
        b.h("Response", 2)
        if i < len(answerers):
            answerers[i](b, facts, subs)
        else:
            answer_optional(b, subs)

    b.page_break()
    b.h("How this document was produced", 1)
    b.p("Rebuild it with the command below. The questions are re-read from the round-2 document, and "
        "every number is re-read from its artefact, so a changed artefact changes this response instead of "
        "leaving it out of date.")
    b.code(["uv run python docs/paper/make_round2_response.py",
            "uv run python docs/paper/verify_artifacts.py   # confirms the pinned artefacts are unaltered"])
    b.p("External sources consulted on 2026-09-24: TalkBank data-usage rules "
        "(https://talkbank.org/0share/rules.html), data access levels "
        "(https://talkbank.org/0share/access.html), and the ASDBank home page (https://talkbank.org/asd/). "
        "Re-check them on the submission date, since these pages can change.", italic=True)

    b.doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    print(build())
