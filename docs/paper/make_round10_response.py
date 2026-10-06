"""Build the response to the round-10 technical request (R10-1 audit preparation).

Reads ``R10-Technical_Request.docx`` and answers its "What we need back" list in
order, each in plain words and then as facts. Numbers are read from the audit
package written by ``tests/eval/extraction/codebook_audit.py build``. No label
is proposed or changed here; no model is called.

    uv run python -m tests.eval.extraction.codebook_audit build
    uv run python docs/paper/make_round10_response.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

PAPER = Path(__file__).resolve().parent
sys.path.insert(0, str(PAPER))

import docx  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from make_round2_response import QUOTE, Doc  # noqa: E402

ROOT = PAPER.parents[1]
AUDIT = ROOT / "tests" / "eval" / "extraction" / "results" / "codebook_audit_r10"
QUESTIONS = PAPER / "R10-Technical_Request.docx"
OUT = PAPER / "Round10_Technical_Response.docx"
PLAIN_FILL, PLAIN_EDGE = "EAF7F1", "1BAF7A"


def read_needs() -> list[str]:
    src = docx.Document(str(QUESTIONS))
    items, on = [], False
    for el in src.element.body.iterchildren():
        if el.tag.split("}")[1] != "p":
            continue
        p = Paragraph(el, src)
        if p.style.name == "Heading 1":
            on = p.text.strip() == "What we need back from Claude"
        elif on and p.text.strip():
            items.append(p.text.strip())
    return items


def ask(b: Doc, n: int, text: str) -> None:
    para = b.doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.keep_with_next = True
    r = para.add_run(f"{n}. ")
    r.bold = True
    r.font.color.rgb = QUOTE
    b._runs(para, text, colour=QUOTE)


def plain(b: Doc, text: str) -> None:
    b.boxed("In plain words", [text], fill=PLAIN_FILL, edge=PLAIN_EDGE)


def audit_results(b: Doc, key: dict) -> None:
    import filecmp
    from collections import Counter
    res = json.loads((AUDIT / "audit_result.json").read_text())
    rc = json.loads((AUDIT / "recompute_pre_vs_post.json").read_text())
    same = filecmp.cmp(AUDIT / "audit_worksheet_specialist_1_completed.csv", AUDIT / "audit_worksheet_specialist_2_completed.csv", shallow=False)
    assert res["segments_audited"] == 152 and res["labels_audited"] == 456
    n_changes = len(res["changes"])
    assert n_changes == sum(v["changed"] for v in res["per_field"].values())

    b.page_break()
    b.h("Audit results (returned worksheets, commit 60c4704)", 1)
    plain(b, f"The specialists checked all 456 labels on the 152 segments and required {n_changes} changes in "
             f"{len(res['segments_changed'])} segments; most changes are in prompt level (the largest group is verbal → "
             "independent). Scored against the corrected reference, the models' saved answers give lower whole-record "
             "and oracle figures. MARS's comparisons keep the same direction and significance on domain, modality and "
             "outcome; on prompt level both of its (non-significant) differences change sign.")
    if same:
        b.boxed("Record-keeping note", ["The two returned worksheets are byte-identical and were committed together "
                                         "(60c4704). The record therefore shows one joint audit, not two independent "
                                         "passes; no disagreement could arise, and none was resolved by consensus."])
    b.table("Labels checked and changed. Source: `codebook_audit_r10/audit_result.json`.",
            ["Field", "Checked", "Changed", "Changes (old → new: count)"],
            [[f, str(v["checked"]), str(v["changed"]),
              "; ".join(f"{a} → {c}: {n}" for (a, c), n in Counter(
                  ("none" if x["old"] is None else x["old"], "none" if x["new"] is None else x["new"])
                  for x in res["changes"] if x["field"] == f).most_common())]
             for f, v in res["per_field"].items()], [1.0, 0.7, 0.7, 4.3])
    src = Counter(x["old_source"] for x in res["changes"])
    b.p(f"Of the {n_changes} changed labels, {src.get('round1_agreed', 0)} had been agreed by both annotators in Round 1 "
        f"and {src.get('adjudicated_round1_conflict', 0)} had been adjudicated. Every change (segment, field, old, new) is "
        "listed in `audit_result.json`. The post-audit reference is `codebook_audit_r10/gold_labels_postaudit_r10.json`; "
        "`results/gold_labels.json` is unchanged.")

    four = rc["four_generators"]
    pct1 = lambda x: f"{100 * x:.1f}"  # noqa: E731
    b.table("Main batch (330), preserved outputs re-scored against the pre- and post-audit reference, % (pre → post). "
            "Source: `codebook_audit_r10/recompute_pre_vs_post.json`.",
            ["Generator", "Merged agreement", "Unsupported assertion", "Whole record"],
            [[m, *[f"{pct1(four['pre']['per_model'][m][k])} → {pct1(four['post']['per_model'][m][k])}"
                   for k in ("merged_agreement", "unsupported_rate", "whole_record")]] for m in four["pre"]["per_model"]],
            [1.4, 1.8, 1.8, 1.7])
    t3 = rc["three_generators"]
    b.table("Oracle ceiling and MARS, % (pre → post).", ["Quantity", "Pre → post"], [
        ["Oracle ceiling, 3 generators (all outputs, 330)", f"{pct1(t3['pre']['oracle_ceiling'])} → {pct1(t3['post']['oracle_ceiling'])}"],
        ["Oracle ceiling, 4 generators", f"{pct1(four['pre']['oracle_ceiling'])} → {pct1(four['post']['oracle_ceiling'])}"],
        ["MARS coverage (unchanged: selection does not use the reference)", f"{pct1(t3['pre']['mars']['coverage'])} → {pct1(t3['post']['mars']['coverage'])}"],
        ["MARS whole-record accuracy over returned", f"{pct1(t3['pre']['mars']['accuracy_over_returned'])} → {pct1(t3['post']['mars']['accuracy_over_returned'])}"],
        ["MARS whole-record accuracy over all 330", f"{pct1(t3['pre']['mars']['accuracy_over_all'])} → {pct1(t3['post']['mars']['accuracy_over_all'])}"],
    ], [4.2, 2.5])
    mp = rc["mars_per_field"]
    assert mp["pre"]["strength_order"][0] == mp["post"]["strength_order"][0]

    def d(x):
        return f"{100 * x['value']:+.1f} [{100 * x['lower']:+.1f}, {100 * x['upper']:+.1f}]" + (" *" if x["excludes_zero"] else "")
    b.table(f"MARS per field on the {mp['post']['n_judged']} reranked segments: MARS minus random reviewable candidate and "
            "minus strongest reviewable generator, percentage points [95% child-cluster CI]; pre-audit reproduces "
            "round4_analysis.json exactly (asserted). * interval excludes zero.",
            ["Field", "MARS % pre → post", "vs random: pre", "vs random: post", "vs strongest: pre", "vs strongest: post"],
            [[f, f"{pct1(mp['pre']['per_field'][f]['mars'])} → {pct1(mp['post']['per_field'][f]['mars'])}",
              d(mp["pre"]["per_field"][f]["vs_random_reviewable"]), d(mp["post"]["per_field"][f]["vs_random_reviewable"]),
              d(mp["pre"]["per_field"][f]["vs_best_reviewable_fallback"]), d(mp["post"]["per_field"][f]["vs_best_reviewable_fallback"])]
             for f in mp["pre"]["per_field"]],
            [0.9, 1.1, 1.2, 1.2, 1.2, 1.2])
    b.bullets([
        f"Strongest generator on merged agreement unchanged: {mp['post']['strength_order'][0]}.",
        "Domain is not audited, so its rows are identical before and after.",
        "Not yet rebuilt: round3/round4/round8/round9 main-batch tables, paper_metrics.json, mars_selection_analysis.json, "
        "figures and the manuscript. The procedure is in section 5; it has not been run.",
    ])


def build() -> Path:
    man: dict[str, Any] = json.loads((AUDIT / "audit_set_152.json").read_text())
    key = json.loads((AUDIT / "audit_key.json").read_text())
    c = man["counts"]
    assert c["n_unreissued"] == 152 and c["n_reissued"] == 178 and c["n_gold"] == 330
    assert len(man["unreissued_segment_ids"]) == 152 and set(man["unreissued_segment_ids"]) == set(key)
    prov = c["label_provenance_counts"]
    assert prov["round1_agreed"] + prov["adjudicated_round1_conflict"] == 456
    needs = read_needs()
    assert len(needs) == 6, needs
    retained_dist = {f: {} for f in ("modality", "outcome", "prompt_level")}
    for v in key.values():
        for f, val in v["retained"].items():
            k = "none" if val is None else val
            retained_dist[f][k] = retained_dist[f].get(k, 0) + 1

    b = Doc()
    grey = RGBColor(0x52, 0x51, 0x4E)
    b.doc.add_heading("Round 10 — R10-1 Audit Preparation: Response", level=0)
    b.p("Answers the six items under 'What we need back from Claude' in 'R10-Technical_Request.docx', in order, each "
        "in plain words and then as facts. Branch model-finetuning, built 2026-09-27. Internal working document.", colour=grey)
    b.p("Boundary kept: no label was judged, proposed, changed or adjudicated by software or any model. The package "
        "only reconstructs the set, carries the human-authored values to the specialists, and counts what they return. "
        "No generator, MARS judge or reporter was run.", colour=grey)

    b.h("At a glance", 1)
    b.bullets([
        f"**The 152 are reconstructed exactly** (330 − 178, disjoint), with file hashes and commits.",
        "**The audit package is ready**: two identical model-blind worksheets (one per specialist), an instruction sheet "
        "with the full final codebook, and a key file kept away from the auditors.",
        "**The tooling is ready and tested**: a counter that reports old → audited changes and sends specialist "
        "disagreements back to humans, and a recompute path that runs only if a label changes.",
        f"**One correction to the request's premise:** not all 152 had agreement in Round 1. {c['segments_with_round1_categorical_conflict']} "
        f"had a Round-1 disagreement on a categorical field that the codebook does not settle, so it went to adjudication "
        f"instead of re-issue. For the three audited fields, {prov['adjudicated_round1_conflict']} of the 456 labels "
        f"({c['segments_with_adjudicated_audit_field']} segments) are adjudicated values; the other {prov['round1_agreed']} "
        "are values both annotators chose in Round 1.",
    ])

    b.page_break()
    b.h("1. " + needs[0].split(",")[0], 1)
    ask(b, 1, needs[0])
    plain(b, "Yes. Of the 330 main-batch segments, 178 were sent back in Round 2. The other 152 are exactly the "
             "ones never sent back, and none of them appears in the 178. Every source file is identified by its "
             "fingerprint (hash) and commit, so anyone can check the list is the same one.")
    b.table("Reconstruction. Source: `results/codebook_audit_r10/audit_set_152.json`.", ["Quantity", "Value"], [
        ["Rule", man["rule"]],
        ["Main-batch segments with a reference standard", str(c["n_gold"])],
        ["Re-issued in Round 2", str(c["n_reissued"])],
        ["Unreissued (audit population)", str(c["n_unreissued"])],
        ["Overlap between the two", "0 (asserted)"],
        ["Labels to audit (3 fields × 152)", str(man["labels_to_audit"])],
        ["Annotator records of the 152 replaced by Round 2?", "No — identical to the Round-1 snapshots (asserted)"],
    ], [3.0, 3.7])
    b.table("Round-1 disagreements among the 152 (categorical fields), and where the audited labels come from.",
            ["Item", "Count"],
            [["Segments with any categorical Round-1 disagreement", str(c["segments_with_round1_categorical_conflict"])]]
            + [[f"  … on {f}", str(n)] for f, n in c["round1_conflicts_by_field"].items()]
            + [["Audited labels = both annotators' Round-1 value", str(prov["round1_agreed"])],
               ["Audited labels = adjudicated from a Round-1 disagreement", str(prov["adjudicated_round1_conflict"])]],
            [4.5, 2.2])
    b.table("Source artefacts (sha256 prefix and last commit).", ["File", "sha256", "Last commit"],
            [[n, v["sha256"][:16] + "…", v["last_commit"]] for n, v in man["sources"].items()], [2.6, 1.8, 2.3])
    b.p("Full segment-ID list: `audit_set_152.json` → `unreissued_segment_ids`.")

    b.page_break()
    b.h("2. The model-blind worksheet", 1)
    ask(b, 2, needs[1])
    plain(b, "Each specialist gets a spreadsheet with one row per segment: an anonymous code (A001…A152), the "
             "transcript exactly as the annotators saw it, and the current value of the three fields. Next to each "
             "value are two empty boxes: a verdict (compliant / change required / not applicable) and, only if a change "
             "is required, the corrected value. Nothing about any AI model appears anywhere in it.")
    b.table("Worksheet columns (`audit_worksheet_specialist_1.csv`, `…_specialist_2.csv`, identical).",
            ["Column", "Content"], [
                ["audit_id", "A001–A152, opaque; the segment ID is kept in `audit_key.json`, not on the sheet"],
                ["transcript", "de-identified transcript text, identical to the Round-1 annotation sheet (`chat_text` of annotation_batch.json)"],
                ["<field>_retained", "current reference value; 'none' = the transcript supports no value"],
                ["<field>_status", "blank — compliant | change_required | not_applicable"],
                ["<field>_corrected", "blank — filled only when change_required"],
                ["notes", "blank, optional"],
            ], [1.5, 5.2])
    b.bullets([
        "Model-blind: the build asserts that no model name appears anywhere in the sheet and that it has no columns "
        "beyond those listed. No generator output, status, confidence, score, MARS or selection field is read into it.",
        "Rows are in a fixed shuffled order (seed 20261010), so row order does not reveal child or session sequence.",
        "Not shown to auditors: segment/child/session IDs, and whether a label was agreed or adjudicated in Round 1 "
        "(both in `audit_key.json`).",
    ])
    b.table("Values the auditors will see (retained labels on the 152).", ["Field", "Distribution"],
            [[f, ", ".join(f"{k} {n}" for k, n in sorted(d.items(), key=lambda kv: -kv[1]))] for f, d in retained_dist.items()],
            [1.3, 5.4])

    b.h("3. The instruction sheet", 1)
    ask(b, 3, needs[2])
    plain(b, "A text file for the specialists: what to check, what not to look at, how to fill each box, the allowed "
             "values, the three boundaries in one line each, and then the full final codebook word for word.")
    b.bullets(["`results/codebook_audit_r10/AUDIT_INSTRUCTIONS.txt`; the codebook is embedded verbatim with its commit.",
               "The protocol from the request is followed: judge compliance with the clarified boundary only; do not "
               "re-annotate; stay blind; keep the retained value untouched; corrected value only when a change is required."])

    b.page_break()
    b.h("4. Counting old → audited changes", 1)
    ask(b, 4, needs[3])
    plain(b, "When both sheets come back, one command checks them and counts. If the two specialists agree that a "
             "label is fine, it stays. If both require the same new value, that is a change. Anything else is a "
             "disagreement: the program writes those rows to a consensus sheet for the humans and stops — it never picks "
             "a side. After the consensus sheet is returned, it produces the final counts and the list of every change.")
    b.code(["# return the sheets as audit_worksheet_specialist_1_completed.csv / ..._2_completed.csv",
            "uv run python -m tests.eval.extraction.codebook_audit compare",
            "# if it reports disagreements: humans fill consensus_sheet.csv, then",
            "uv run python -m tests.eval.extraction.codebook_audit compare --consensus <completed consensus csv>"])
    b.bullets([
        "Validation before counting: every row present, retained column unedited, status from the allowed three, a "
        "corrected value only with change_required, from the allowed list and different from the retained value. Any "
        "problem stops the run and nothing is counted.",
        "'compliant' and 'not_applicable' both leave the label unchanged; the status counts are still reported per specialist.",
        "Output `audit_result.json`: segments audited; per field, labels checked and changed; changed segment IDs; every "
        "change as old → new with its Round-1 source; number of disputes resolved by consensus.",
        "Tested on synthetic sheets (`tests/eval/extraction/test_codebook_audit.py`, 6 tests): zero-change path, "
        "disagreement goes to consensus unresolved, invalid entries rejected.",
    ])

    b.h("5. Recomputation, only if a label changes", 1)
    ask(b, 5, needs[4])
    plain(b, "If nothing changes, nothing is recomputed. If something changes, the program writes a new copy of the "
             "reference standard next to the audit files — the original is never touched — and re-scores the answers "
             "the models already gave. No model is run again.")
    b.code(["uv run python -m tests.eval.extraction.codebook_audit apply       # refuses if there are no changes",
            "uv run python -m tests.eval.extraction.codebook_audit recompute   # old vs post-audit, preserved outputs only"])
    b.bullets([
        "`apply` writes `codebook_audit_r10/gold_labels_postaudit_r10.json`, carrying the list of changes and the "
        "sha256 of the pre-audit file; it refuses to overwrite an existing post-audit file.",
        "`recompute` scores the preserved outputs (3- and 4-generator batches) against both versions and reports, side "
        "by side: per-generator merged agreement, unsupported assertion and whole-record agreement; the oracle ceiling; "
        "MARS coverage and accuracy. It first checks that the pre-audit side reproduces the recorded results.",
        "Full rebuild of every downstream artefact (below) uses the existing scripts in order on a separate branch where "
        "the post-audit file replaces `gold_labels.json`, so the pre-audit state stays on the main line: "
        "score_against_gold (default, then --batch annotation_batch_augmented.json) → mars_evaluation → "
        "mars_selection_analysis → paper_metrics → holdout_metrics → round3/4/8/9 analyses → figures and manuscript "
        "builders; then compare each JSON with the main-line version.",
    ])

    b.page_break()
    b.h("6. What would be affected if a label changed", 1)
    ask(b, 6, needs[5])
    plain(b, "Everything scored against the 330 main-batch reference would move; nothing about the held-out sets, the "
             "annotator agreement figures or the models' own answers would.")
    b.table("Outputs that depend on the main-batch modality/outcome/prompt_level reference.", ["Output (results/…)", "What moves", "Script"], [
        ["scored_against_gold.json, scored_against_gold_augmented.json", "every per-field judgement", "score_against_gold.py"],
        ["mars_evaluation.json", "per-generator whole-record accuracy, oracle ceiling (6.1%), MARS accuracy", "mars_evaluation.py"],
        ["mars_selection_analysis.json", "MARS vs random / strongest generator, per field", "mars_selection_analysis.py"],
        ["paper_metrics.json", "headline extraction metrics", "paper_metrics.py"],
        ["holdout_metrics.json (batch side)", "main-batch comparison figures", "holdout_metrics.py"],
        ["round3_analysis.json (main batch)", "merged / unsupported / whole-record, qwen3:14b − qwen2.5:7b paired CIs", "round3_analysis.py"],
        ["round4_analysis.json", "same-universe main batch and contrasts; MARS clustered bootstrap; null-output rates; capped vs shorter reference nulls", "round4_analysis.py"],
        ["round5 (reference nulls by cap), round8 (performance by position), round9 (R9-5 330 tables, R9-8 oracle)", "the main-batch parts", "round5/8/9_analysis.py"],
        ["Model-selection statement (strongest generator = qwen2.5:7b on merged agreement)", "could change if the ordering changes", "round4 mars_clustered asserts the order"],
        ["Reporter experiment (11 EvidencePackets)", "packet inputs are built from this reference; the recorded 0/11 describes the pre-audit packets and cannot be re-scored without a new reporter run (out of scope)", "tests/eval/reporting/groundedness.py"],
        ["Figures and response/manuscript builders", "any figure drawn from the above", "docs/paper/make_*.py"],
    ], [2.4, 2.9, 1.4])
    b.bullets([
        "Not affected: Round-1/Round-2 inter-annotator kappa (computed from the annotators' own files); the 132/121 "
        "evaluation sets (separate reference); preserved generator outputs; MARS selections in mars_outcomes.json (made "
        "without the reference).",
    ])

    if (AUDIT / "audit_result.json").exists() and (AUDIT / "recompute_pre_vs_post.json").exists():
        audit_results(b, key)

    b.h("Choices made in preparing the package (please confirm or change)", 1)
    b.bullets([
        "Opaque audit IDs and shuffled row order instead of segment IDs (the request allowed either).",
        "Round-1 provenance (agreed vs adjudicated) kept off the worksheet to avoid anchoring; it is in the key.",
        "Two separate identical worksheets, one per specialist, so the first pass is independent.",
        "The post-audit reference is written inside `codebook_audit_r10/`, never beside or over `gold_labels.json`.",
    ])
    b.h("How this document was produced", 1)
    b.code(["uv run python -m tests.eval.extraction.codebook_audit build",
            "uv run python docs/paper/make_round10_response.py"])
    b.doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    print(build())
