"""Build the response to the round-6 technical questions (R6-1, R6-3).

Reads ``Round6_review_feedback.docx`` in document order and answers each
technical item underneath its request, quoted verbatim; R6-2 and R6-4 need no
technical action and are acknowledged only. Every number is read from
``round6_analysis.json`` (added this round, no model calls) or an earlier
recorded artefact, and each table names its source. Statements about the data
are asserted before they are printed.

    uv run python -m tests.eval.extraction.round6_analysis
    uv run python docs/paper/make_round6_response.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

PAPER = Path(__file__).resolve().parent
sys.path.insert(0, str(PAPER))

import docx  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from make_round2_response import AQUA, BLUE, COLOUR, GRID, MODELS, ORANGE, QUOTE, Doc, short  # noqa: E402
from make_round3_response import GREY  # noqa: E402

ROOT = PAPER.parents[1]
RESULTS = ROOT / "tests" / "eval" / "extraction" / "results"
QUESTIONS = PAPER / "Round6_review_feedback.docx"
FIG = PAPER / "figures" / "round6"
OUT = PAPER / "Round6_Review_Response.docx"
FIELDS = ["domain", "modality", "prompt_level", "outcome"]
METRIC_NAMES = {"merged_agreement": "Merged agreement", "unsupported_rate": "Unsupported assertion",
                "whole_record": "Whole-record agreement"}


def load(name: str) -> Any:
    return json.loads((RESULTS / name).read_text())


def pp(x: float) -> str:
    return f"{100 * x:+.1f}"


def v_ci(d: dict) -> str:
    return f"{100 * d['value']:.1f} [{100 * d['ci95']['lower']:.1f}, {100 * d['ci95']['upper']:.1f}]"


def k_ci(d: dict) -> str:
    c = d["cluster_bootstrap_ci"]
    return f"[{c['lower']:.3f}, {c['upper']:.3f}]"


# ---------------------------------------------------------------------------
# Requests, verbatim and in order
# ---------------------------------------------------------------------------

def read_sections() -> list[dict[str, Any]]:
    """Heading-1 sections of the feedback document, with their paragraphs and tables in order."""
    src = docx.Document(str(QUESTIONS))
    sections: list[dict[str, Any]] = []
    for el in src.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, src)
            t = p.text.strip()
            if not t:
                continue
            if p.style.name == "Heading 1":
                sections.append({"title": t, "items": []})
            elif sections:
                sections[-1]["items"].append((p.style.name, t))
        elif tag == "tbl" and sections:
            rows = [[c.text for c in r.cells] for r in Table(el, src).rows]
            sections[-1]["items"].append(("table", rows))
    return sections


def render_request(b: Doc, sec: dict[str, Any]) -> None:
    b.label("The request, as written")
    for style, t in sec["items"]:
        if style == "table":
            continue
        para = b.doc.add_paragraph()
        indent = {"List Bullet": 24, "List Bullet 2": 36}.get(style, 12)
        para.paragraph_format.left_indent = Pt(indent)
        para.paragraph_format.space_after = Pt(2)
        if style == "Heading 2":
            r = para.add_run(t)
            r.bold = True
            r.font.color.rgb = QUOTE
            continue
        prefix = "• " if style == "List Bullet" else "– " if style == "List Bullet 2" else ""
        b._runs(para, prefix + t, colour=QUOTE)


# ---------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------

def fig_turn_overlap(ta: dict) -> Path:
    aff = ta["affected"]
    fig, ax = plt.subplots(figsize=(10, 0.75 * len(aff) + 1.4))
    for i, a in enumerate(aff):
        m = a["matches"][0]
        n_e, n_m, shared = a["n_turns"], m["main_n_turns"], m["n_shared"]
        # Place the evaluation episode at turns 0..n_e-1 and the main-batch episode so the shared turn aligns.
        if m["eval_position"][0] == "last":
            e0, m0 = 0, n_e - shared
        else:
            m0, e0 = 0, n_m - shared
        y = i
        for t in range(n_e):
            ax.barh(y - 0.17, 1, left=e0 + t, height=0.3, color=AQUA, edgecolor="white", lw=1)
        for t in range(n_m):
            ax.barh(y + 0.17, 1, left=m0 + t, height=0.3, color=BLUE, edgecolor="white", lw=1)
        sx = e0 + n_e - shared if m["eval_position"][0] == "last" else e0
        ax.add_patch(plt.Rectangle((sx, y - 0.34), shared, 0.68, fill=False, edgecolor=ORANGE, lw=2.2, zorder=4))
        ax.text(15.3, y, f"shared: line {m['shared_lines'][0]}", va="center", fontsize=7.5, color=ORANGE)
        ax.text(-0.3, y, f"{a['child_id']}\n{a['session_id'].split('_')[-1]}", ha="right", va="center", fontsize=7.5)
    ax.set_ylim(-0.7, len(aff) - 0.2)
    ax.set_xlim(-3.2, 19)
    ax.set_xticks([])
    ax.invert_yaxis()
    ax.set_yticks([])
    ax.set_xlabel("consecutive source utterances in the session (one block = one CHAT line)")
    for sp in ("left", "right", "top", "bottom"):
        ax.spines[sp].set_visible(False)
    ax.legend(handles=[Patch(color=AQUA, label="evaluation episode (121-set)"),
                       Patch(color=BLUE, label="main-batch episode (330)"),
                       Patch(fill=False, edgecolor=ORANGE, lw=2.2, label="shared source turn")],
              frameon=False, fontsize=7.5, loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=3)
    fig.tight_layout()
    p = FIG / "r6_fig1_turn_overlap.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


def fig_sensitivity(sens: dict) -> Path:
    groups = [("overlap_excluded_121", "121 (post-audit)", "-"), ("turn_overlap_excluded", f"{sens['turn_overlap_excluded']['n']} (turn-overlap excluded)", (0, (2, 1.5)))]
    fig, axes = plt.subplots(1, 3, figsize=(12.5, 3.6), sharey=True)
    for ax, key in zip(axes, METRIC_NAMES):
        for j, m in enumerate(MODELS):
            for g, (gname, _lab, ls) in enumerate(groups):
                v = sens[gname]["per_model"][m][key]
                y = j + (g - 0.5) * 0.3
                ax.plot([100 * v["ci95"]["lower"], 100 * v["ci95"]["upper"]], [y, y], color=COLOUR[m], lw=2, ls=ls)
                ax.scatter([100 * v["value"]], [y], s=26, zorder=3, color=COLOUR[m] if g == 0 else "white",
                           edgecolor=COLOUR[m], lw=1.4)
        ax.set_title(METRIC_NAMES[key], loc="left", fontsize=9.5, weight="bold")
        ax.grid(axis="x", color=GRID, lw=0.6)
        ax.set_axisbelow(True)
        ax.set_xlabel("% (95% child-cluster CI)")
    axes[0].set_yticks(range(len(MODELS)), [short(m) for m in MODELS], fontsize=8)
    axes[0].invert_yaxis()
    axes[0].scatter([], [], color="k", s=24, label=groups[0][1])
    axes[0].scatter([], [], facecolor="white", edgecolor="k", s=24, label=groups[1][1])
    fig.legend(frameon=False, fontsize=8, loc="lower center", ncol=2)
    fig.tight_layout(rect=(0, 0.08, 1, 1))
    p = FIG / "r6_fig2_sensitivity.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


def fig_kappa(rel: dict) -> Path:
    fig, ax = plt.subplots(figsize=(8.6, 3.2))
    for i, f in enumerate(FIELDS + ["mean"]):
        for g, (key, lab, col) in enumerate((("set_132", "132-set", GREY), ("set_121", "121-set", AQUA))):
            y = i + (g - 0.5) * 0.3
            if f == "mean":
                ax.scatter([rel[key]["mean_kappa"]], [y], s=34, color=col, marker="D", zorder=3)
                continue
            d = rel[key]["fields"][f]
            c = d["cluster_bootstrap_ci"]
            ax.plot([c["lower"], c["upper"]], [y, y], color=col, lw=2.2)
            ax.scatter([d["kappa"]], [y], s=30, color=col, zorder=3, label=lab if i == 0 else None)
            ax.text(c["upper"] + 0.004, y, f"{d['kappa']:.3f}", va="center", fontsize=7)
        if f == "mean":
            for g, key in enumerate(("set_132", "set_121")):
                ax.text(rel[key]["mean_kappa"] + 0.006, i + (g - 0.5) * 0.3, f"{rel[key]['mean_kappa']:.3f}", va="center", fontsize=7)
    ax.set_yticks(range(5), [f.replace("_", " ") for f in FIELDS] + ["mean of 4\n(no CI)"], fontsize=8)
    ax.invert_yaxis()
    ax.set_xlim(0.76, 1.04)
    ax.set_xlabel("Cohen's kappa (95% child-cluster CI)")
    ax.grid(axis="x", color=GRID, lw=0.6)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, fontsize=8, loc="lower left")
    fig.tight_layout()
    p = FIG / "r6_fig3_kappa.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


# ---------------------------------------------------------------------------
# Answers
# ---------------------------------------------------------------------------

def a_r61(b: Doc, ctx) -> None:
    ta, sens = ctx["r6"]["turn_audit"], ctx["r6"]["sensitivity"]
    aff = ta["affected"]
    assert ta["n_eval"] == 121 and ta["n_main"] == 330
    assert all(len(a["matches"]) == 1 and a["n_shared_turns"] == 1 for a in aff)
    assert all({a["matches"][0]["eval_position"][0], a["matches"][0]["main_position"][0]} == {"first", "last"} for a in aff)
    assert ta["context_frozen_132"]["identity_overlap_with_330"] == 1
    base, excl = sens["overlap_excluded_121"], sens["turn_overlap_excluded"]
    max_shift = max(abs(excl["per_model"][m][k]["value"] - base["per_model"][m][k]["value"]) for m in MODELS for k in METRIC_NAMES)
    assert max_shift < 0.012
    pb, pe = base["paired_qwen3_14b_minus_qwen2_5_7b"], excl["paired_qwen3_14b_minus_qwen2_5_7b"]
    assert all(pb[k]["ci95"]["excludes_zero"] == pe[k]["ci95"]["excludes_zero"] for k in METRIC_NAMES)
    for m in MODELS:  # the 121 baseline reproduces the recorded round-4 values
        for k in METRIC_NAMES:
            assert abs(base["per_model"][m][k]["value"] - ctx["r4_121"][m][k]["value"]) < 1e-4

    b.label("Short answer")
    b.answer_line(
        f"**Overlap exists but is minimal: {ta['n_eval_affected']} of the 121 evaluation episodes share exactly one "
        f"source turn each with {ta['n_main_involved']} distinct main-batch episodes ({ta['n_distinct_shared_turns']} "
        "distinct shared utterance lines in total).** Every shared turn is the boundary utterance between two adjacent "
        "episodes of the same session: the last line of one episode is the first line of the next, which is what the "
        "episode-construction rule permits. Excluding the five for sensitivity leaves "
        f"{ta['n_eval_remaining']} episodes; no metric for any generator moves by more than {100 * max_shift:.1f} "
        "percentage points, and every qwen3:14b − qwen2.5:7b paired conclusion is unchanged.")

    b.h("Matching key and logic", 3)
    b.bullets([
        "**Key:** (session_id, CHAT source line number) for every utterance used to build an episode. Line numbers are "
        "those the parser assigns to each utterance in the source `.cha` file, carried unchanged through "
        "de-identification and segmentation.",
        "**Sets compared:** the 121 post-audit evaluation episodes (frozen 132 minus the 11 pool-identity overlaps, "
        "resolved through the cap-5000 selection) and the 330 main-batch episodes (resolved through the cap-3000 "
        "selection). Both selections draw from the same single segmentation of the corpus, so line sets are directly "
        "comparable.",
        "**Rule:** within the same session, an evaluation episode is affected if its line set intersects the line set "
        "of any main-batch episode. Rendered text was not used.",
        "**Why identity alone was not enough:** an episode starts one utterance before the opening child turn, so the "
        "last utterance of one episode can be the first of the next while the two have different (session_id, "
        "start_line) identities.",
        f"**Check:** on the full 132 the same audit finds {ta['context_frozen_132']['identity_overlap_with_330']} "
        f"identity overlap with the 330 (the episode the reviewer refers to) and "
        f"{ta['context_frozen_132']['turn_overlap_with_330']} turn overlaps in total, i.e. that one plus the five here.",
    ])

    b.table("Affected evaluation episodes. Episode IDs are the study segment IDs. Position = where the shared line "
            "sits in each episode. Source: `round6_analysis.json` (turn_audit.affected).",
            ["Evaluation episode", "Child / session", "Matched main-batch episode", "Shared line", "Eval pos.",
             "Main pos.", "Shared turns"],
            [[a["segment_id"].split(".cha_")[-1] + " (" + a["segment_id"].split("_")[0] + ")",
              f"{a['child_id']} / {a['session_id']}",
              a["matches"][0]["main_segment_id"].split(".cha_")[-1],
              str(a["matches"][0]["shared_lines"][0]),
              a["matches"][0]["eval_position"][0], a["matches"][0]["main_position"][0],
              f"{a['n_shared_turns']} of {a['n_turns']}"] for a in aff],
            [1.15, 2.1, 1.05, 0.6, 0.55, 0.55, 0.7])
    b.figure(fig_turn_overlap(ta), "Each affected evaluation episode (top bar) and its matched main-batch episode "
             "(bottom bar), aligned on source turns. The only shared turn in each pair (orange) is the boundary "
             "utterance between adjacent episodes.")

    b.table("Summary counts. Source: `round6_analysis.json` (turn_audit).", ["Quantity", "Count"], [
        ["Evaluation episodes sharing ≥ 1 source turn with the 330", f"{ta['n_eval_affected']} of {ta['n_eval']}"],
        ["Distinct main-batch episodes involved", str(ta["n_main_involved"])],
        ["Distinct shared source turns (utterance lines)", str(ta["n_distinct_shared_turns"])],
        ["Shared turns per affected evaluation episode", "1 (each of 5)"],
        ["Evaluation episodes remaining after excluding all turn-overlapping ones", str(ta["n_eval_remaining"])],
        ["Per child (affected episodes / shared turns)",
         "; ".join(f"{c} {v['affected_eval_episodes']}/{v['shared_turns']}" for c, v in ta["per_child"].items())
         + "; all other children 0/0"],
    ], [4.2, 2.5])

    rows = []
    for m in MODELS:
        row = [short(m)]
        for k in METRIC_NAMES:
            row += [v_ci(base["per_model"][m][k]), v_ci(excl["per_model"][m][k])]
        rows.append(row)
    b.table(f"Sensitivity: central metrics on the 121 and on the {excl['n']} after excluding the "
            f"{ta['n_eval_affected']} turn-overlapping episodes, % [95% child-cluster CI, 10,000 replicates]. "
            f"{excl['n_children']} children remain. Nothing redrawn or replaced; reference labels unchanged. "
            "Source: `round6_analysis.json` (sensitivity).",
            ["Model", "Merged: 121", f"Merged: {excl['n']}", "Unsupp.: 121", f"Unsupp.: {excl['n']}",
             "Whole: 121", f"Whole: {excl['n']}"],
            rows, [0.8, 1.0, 1.0, 1.0, 1.0, 0.95, 0.95])
    b.table("Paired difference qwen3:14b − qwen2.5:7b, percentage points [95% child-cluster CI]. "
            "Source: `round6_analysis.json` (sensitivity.*.paired_qwen3_14b_minus_qwen2_5_7b).",
            ["Metric", "121", str(excl["n"])],
            [[METRIC_NAMES[k],
              f"{pp(pb[k]['value'])} [{pp(pb[k]['ci95']['lower'])}, {pp(pb[k]['ci95']['upper'])}]" + (" *" if pb[k]["ci95"]["excludes_zero"] else ""),
              f"{pp(pe[k]['value'])} [{pp(pe[k]['ci95']['lower'])}, {pp(pe[k]['ci95']['upper'])}]" + (" *" if pe[k]["ci95"]["excludes_zero"] else "")]
             for k in METRIC_NAMES],
            [2.0, 2.3, 2.3])
    b.p("* interval excludes zero.")
    b.figure(fig_sensitivity(sens), f"Central metrics on the 121 (filled, solid) and on the {excl['n']} after "
             "excluding turn-overlapping episodes (open, dashed).")

    b.h("Artefacts", 3)
    b.bullets([
        "Script: `tests/eval/extraction/round6_analysis.py` (reuses the round-4 pipeline rebuild and frame audit).",
        "Output: `tests/eval/extraction/results/round6_analysis.json` (turn_audit, sensitivity).",
        "Inputs: `holdout_set.json`, `annotation_pool.json`, `scored_against_gold_augmented.json` (330 IDs), "
        "`scored_holdout.json` (scored outputs), CHAT sources in `exp-data/`.",
    ])


def a_r63(b: Doc, ctx) -> None:
    rel = ctx["r6"]["reliability"]
    s132, s121 = rel["set_132"], rel["set_121"]
    diffs = {f: s121["fields"][f]["kappa"] - s132["fields"][f]["kappa"] for f in FIELDS}
    for f in FIELDS:
        c = s132["fields"][f]["cluster_bootstrap_ci"]
        assert c["lower"] <= s121["fields"][f]["kappa"] <= c["upper"]
    assert all(0 <= d < 0.03 for d in diffs.values())
    dm = s121["mean_kappa"] - s132["mean_kappa"]

    b.label("Short answer")
    b.answer_line(
        f"**Mean kappa on the 121-set is {s121['mean_kappa']:.3f}, against {s132['mean_kappa']:.3f} on the 132-set "
        f"({dm:+.3f}). The result is essentially unchanged:** every field's kappa rises slightly "
        f"({min(diffs.values()):+.3f} to {max(diffs.values()):+.3f}) and every 121-set value lies inside the 132-set "
        "95% child-cluster interval.")
    rows = [[f.replace("_", " "), f"{s132['fields'][f]['kappa']:.3f}", k_ci(s132["fields"][f]),
             f"{s121['fields'][f]['kappa']:.3f}", k_ci(s121["fields"][f]), f"{diffs[f]:+.3f}"] for f in FIELDS]
    rows.append(["Unweighted mean of 4", f"{s132['mean_kappa']:.3f}", "–", f"{s121['mean_kappa']:.3f}", "–", f"{dm:+.3f}"])
    b.table("Cohen's kappa, pre-adjudication holdout annotators A vs B. 95% child-cluster percentile bootstrap, "
            "10,000 replicates, seed 20260915 (as published). The mean is the unweighted mean of the four field "
            "kappas, the published definition; no interval is reported for it, as for the 132. "
            "Source: `round6_analysis.json` (reliability).",
            ["Field", "132-set κ", "132-set 95% CI", "121-set κ", "121-set 95% CI", "Δ (121 − 132)"],
            rows, [1.4, 0.8, 1.2, 0.8, 1.2, 0.9])
    b.figure(fig_kappa(rel), "Field kappas with 95% child-cluster intervals on the 132 (grey) and the 121 (teal); "
             "diamonds are the unweighted means.")
    b.bullets([
        f"**Check:** the same code applied to all 132 reproduces the recorded `reliability_holdout.json` exactly "
        f"(mean {s132['mean_kappa']:.4f}, identical field kappas and intervals) before the 121 subset is taken.",
        f"**Interval notes:** modality's child-cluster interval rests on "
        f"{s121['fields']['modality']['cluster_bootstrap_ci']['n_valid']:,} valid replicates (resamples with a single "
        "category leave kappa undefined and are dropped), as on the 132.",
        "**Labels:** `gold_holdout_annotator_a.json`, `gold_holdout_annotator_b.json` (first-pass, pre-adjudication). "
        "Script: `tests/eval/extraction/round6_analysis.py`, which calls `reliability_stats.build_report`. "
        "Output: `round6_analysis.json` (reliability.set_121).",
    ])


def a_ack(b: Doc, _ctx) -> None:
    b.label("Response")
    b.answer_line("Noted. No code, model or annotation work was done for R6-2 or R6-4. The facts they rely on are "
                  "already recorded in earlier rounds (`round4_analysis.json` mars_clustered; "
                  "`mars_selection_analysis.json`).")


def build() -> Path:
    FIG.mkdir(parents=True, exist_ok=True)
    r6, r4 = load("round6_analysis.json"), load("round4_analysis.json")
    ctx = {"r6": r6, "r4_121": r4["same_universe"]["groups"]["overlap_excluded_121"]["per_model"]}
    sections = read_sections()
    answers = {"R6-1": a_r61, "R6-3": a_r63, "No architect action": a_ack}
    answered = [s for s in sections if any(s["title"].startswith(k) for k in answers)]
    assert [s["title"][:4] for s in answered] == ["R6-1", "R6-3", "No a"], [s["title"] for s in sections]
    ta, rel = r6["turn_audit"], r6["reliability"]

    b = Doc()
    grey = RGBColor(0x52, 0x51, 0x4E)
    b.doc.add_heading("Response to Round-6 Technical Questions", level=0)
    b.p("Answers the technical items in 'Round6_review_feedback.docx' (R6-1, R6-3) in their original order, each "
        "request quoted first. Branch model-finetuning, built 2026-09-26. Internal working document; not for submission.",
        colour=grey)
    b.p("Scope as requested: no model re-run, no new annotation, no re-adjudication, no redraw, no replacement of "
        "frozen records. Every number is read from `round6_analysis.json`, computed from preserved outputs, "
        "annotator labels and the deterministic CHAT pipeline. Claims are checked by assertion when this document "
        "is built.", colour=grey)

    b.h("Results at a glance", 1)
    b.table("The two requested analyses in one line each.", ["Priority", "Comment", "Result"], [
        ["1", "R6-1", f"{ta['n_eval_affected']} of 121 evaluation episodes share 1 boundary turn each with "
                      f"{ta['n_main_involved']} main-batch episodes; excluding them ({ta['n_eval_remaining']} left) "
                      "changes no conclusion"],
        ["2", "R6-3", f"Mean kappa {rel['set_121']['mean_kappa']:.3f} on the 121 vs {rel['set_132']['mean_kappa']:.3f} "
                      "on the 132; essentially unchanged"],
    ], [0.7, 0.8, 5.2])

    for s in answered:
        b.page_break()
        b.h(s["title"], 1)
        render_request(b, s)
        b.h("Response", 2)
        next(fn for k, fn in answers.items() if s["title"].startswith(k))(b, ctx)

    b.h("How this document was produced", 1)
    b.code(["uv run python -m tests.eval.extraction.round6_analysis   # no model calls; needs exp-data/ locally",
            "uv run python docs/paper/make_round6_response.py"])
    b.doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    print(build())
