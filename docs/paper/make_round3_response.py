"""Build the response to the round-3 review comments.

Reads the ten questions verbatim from ``Roun3-Review comments.docx`` (same
order, same bullets) and answers each underneath. Every number is read from a
recorded artefact -- chiefly ``round3_analysis.json`` and
``segmentation_stats.json``, which this round added -- and each table names
its source. Test counts are run records and are labelled as such.

    uv run python -m tests.eval.extraction.round3_analysis
    uv run python -m tests.eval.extraction.segmentation_stats
    uv run python docs/paper/make_round3_response.py
"""

from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

PAPER = Path(__file__).resolve().parent
sys.path.insert(0, str(PAPER))

import docx  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx.shared import RGBColor  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle  # noqa: E402

from make_round2_response import (  # noqa: E402
    AQUA, BLUE, COLOUR, GRID, INK, INK2, MODELS, ORANGE, YELLOW, Doc, pct, render_question, short,
)

ROOT = PAPER.parents[1]
RESULTS = ROOT / "tests" / "eval" / "extraction" / "results"
QUESTIONS = PAPER / "Roun3-Review comments.docx"
FIG = PAPER / "figures" / "round3"
OUT = PAPER / "Round3_Review_Response.docx"
GREY = "#9a9893"


def load(name: str) -> Any:
    return json.loads((RESULTS / name).read_text())


METRICS3 = ("merged_agreement", "unsupported_rate", "whole_record")


def diff_str(d: dict) -> str:
    """A paired difference in points with its interval, as printed everywhere in this document."""
    return f"{100 * d['value']:+.1f} [{100 * d['ci95']['lower']:+.1f}, {100 * d['ci95']['upper']:+.1f}]"


def overlaps(a: dict, b: dict) -> bool:
    return a["ci95"]["lower"] <= b["ci95"]["upper"] and b["ci95"]["lower"] <= a["ci95"]["upper"]


def ranking(per_model: dict, key: str) -> list[str]:
    return sorted(per_model, key=lambda m: -per_model[m][key]["value"])


# ---------------------------------------------------------------------------
# Questions, verbatim
# ---------------------------------------------------------------------------

def read_questions() -> list[dict[str, Any]]:
    src = docx.Document(str(QUESTIONS))
    sections: list[dict[str, Any]] = []
    for para in src.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        ppr = para._p.pPr
        num = ppr.numPr if ppr is not None else None
        level = int(num.ilvl.val) if num is not None and num.ilvl is not None else None
        bold = any(r.bold for r in para.runs)
        head, _, rest = para.text.partition("\n")
        if level is None and re.match(r"^\d+\.\s", head):
            sections.append({"heading": head.strip(), "items": []})
            if rest.strip():
                sections[-1]["items"].append({"text": rest, "level": None, "bold": False})
            continue
        sections[-1]["items"].append({"text": para.text, "level": level, "bold": bold})
    return sections


def bullets_of(items: list[dict[str, Any]]) -> list[str]:
    return [i["text"].strip() for i in items if i["level"] == 0]


# ---------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------

def _box(ax, x, y, w, h, text, edge, fill="white", tc=INK, size=8, weight="normal"):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                                linewidth=1.4, edgecolor=edge, facecolor=fill))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=size, color=tc, weight=weight)


def _arrow(ax, a, b, col=INK2, ls="-"):
    ax.add_patch(FancyArrowPatch(a, b, arrowstyle="-|>", mutation_scale=11, color=col, lw=1.3, linestyle=ls))


def fig_kappa(kc: dict) -> Path:
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(9.6, 3.6), gridspec_kw={"width_ratios": [1.1, 1]})
    a1.set_xlim(0, 10); a1.set_ylim(0, 6); a1.axis("off")
    _box(a1, 0.1, 2.4, 2.0, 1.3, "330 segments\nround 1\n(no codebook)", INK2)
    _box(a1, 3.0, 3.9, 3.3, 1.4, f"152 retained\nround-1 values kept\nκ = {kc['retained_152_round1_values']['mean']:.3f}", BLUE)
    _box(a1, 3.0, 0.5, 3.3, 1.8, f"178 re-issued\n(disagreement-enriched)\nround 1 κ = {kc['reissued_178_round1_values']['mean']:.3f}\n→ round 2 κ = {kc['reissued_178_round2_values']['mean']:.3f}", ORANGE)
    _box(a1, 7.1, 2.3, 2.8, 1.5, f"Final state, 330\nmean κ = {kc['final_state_all_330']['mean']:.3f}\n(pre-adjudication)", AQUA, fill=AQUA, tc="white", weight="bold")
    _arrow(a1, (2.1, 3.3), (3.0, 4.5)); _arrow(a1, (2.1, 2.9), (3.0, 1.5))
    _arrow(a1, (6.3, 4.5), (7.1, 3.3)); _arrow(a1, (6.3, 1.4), (7.1, 2.8))
    a1.text(0.1, 5.7, "A. What 0.801 is computed over", fontsize=9.5, weight="bold")

    labels = ["Round 1,\nall 330", "Retained 152\n(round 1)", "Re-issued 178\nround 1", "Re-issued 178\nround 2", "Final state,\nall 330"]
    keys = ["round1_all_330", "retained_152_round1_values", "reissued_178_round1_values", "reissued_178_round2_values", "final_state_all_330"]
    cols = [INK2, BLUE, ORANGE, ORANGE, AQUA]
    vals = [kc[k]["mean"] for k in keys]
    bars = a2.bar(range(5), vals, color=cols, width=0.62, zorder=3)
    bars[2].set_alpha(0.45)
    for i, v in enumerate(vals):
        a2.text(i, v + 0.02, f"{v:.3f}", ha="center", fontsize=8)
    a2.set_xticks(range(5), labels, fontsize=7.3)
    a2.set_ylim(0, 1.05); a2.set_ylabel("mean Cohen's κ (4 fields)")
    a2.grid(axis="y", color=GRID, lw=0.6); a2.set_axisbelow(True)
    a2.set_title("B. Mean κ by annotation state", loc="left", fontsize=9.5, weight="bold")
    fig.tight_layout()
    p = FIG / "r3_fig1_kappa_composition.png"; fig.savefig(p); plt.close(fig); return p


def fig_pipeline(ss: dict, leak: dict) -> Path:
    fig, ax = plt.subplots(figsize=(9.8, 4.4))
    ax.set_xlim(0, 12); ax.set_ylim(0, 6); ax.axis("off")
    _box(ax, 0.05, 4.2, 2.1, 1.3, f"{ss['n_sessions']} sessions\n{ss['n_utterances']:,} utterances\n11 children", INK2)
    _box(ax, 2.6, 4.2, 2.1, 1.3, f"{ss['n_episodes']:,} episodes\n2–8 turns, ≥1 CHI turn\n(all kept)", INK2)
    _box(ax, 5.15, 4.2, 2.2, 1.3, f"Balanced selection\ncap 273 / child\n→ {ss['n_selected']['3000']:,} segments", BLUE)
    _box(ax, 7.8, 4.2, 1.9, 1.3, "Pool 1,320\n120 / child\nseeded shuffle", BLUE)
    _box(ax, 10.1, 4.2, 1.85, 1.3, "506 examined\n→ 330 batch\n30 / child", BLUE, fill=BLUE, tc="white", weight="bold")
    for a, b in [((2.15, 4.85), (2.6, 4.85)), ((4.7, 4.85), (5.15, 4.85)), ((7.35, 4.85), (7.8, 4.85)), ((9.7, 4.85), (10.1, 4.85))]:
        _arrow(ax, a, b)
    _box(ax, 5.15, 1.6, 2.2, 1.3, f"Balanced selection\ncap 455 / child\n→ {ss['n_selected']['5000']:,} segments", AQUA)
    _box(ax, 7.8, 1.6, 1.9, 1.3, "Exclude pool\nby segment ID\n→ 2,051", AQUA)
    _box(ax, 10.1, 1.6, 1.85, 1.3, "132 frozen eval\n12 / child\nseed 20260916", AQUA, fill=AQUA, tc="white", weight="bold")
    _arrow(ax, (3.65, 4.2), (5.15, 2.3)); _arrow(ax, (7.35, 2.25), (7.8, 2.25)); _arrow(ax, (9.7, 2.25), (10.1, 2.25))
    ax.text(7.8, 0.95, f"ID exclusion missed {leak['n_in_pool']} pool episodes\n(IDs renumber between caps) — see Q3",
            fontsize=7.6, color=ORANGE, weight="bold")
    ax.text(0.05, 5.75, "Main batch", color=BLUE, weight="bold", fontsize=9)
    ax.text(0.05, 3.0, "Evaluation set (original frozen)", color=AQUA, weight="bold", fontsize=9)
    p = FIG / "r3_fig2_pipeline.png"; fig.savefig(p); plt.close(fig); return p


def fig_episode_anatomy() -> Path:
    """Toy sequence; the drawn boundaries come from the real segmenter, not from the figure."""
    from data.pipeline.models import Utterance
    from data.pipeline.segment_episodes import _find_episode_boundaries

    turns = ["MOT", "MOT", "CHI", "MOT", "CHI", "MOT", "MOT", "MOT", "MOT", "CHI", "INV", "CHI", "MOT"]
    bounds = _find_episode_boundaries([Utterance(speaker=t, text="x .", line_number=i) for i, t in enumerate(turns)])
    fig, ax = plt.subplots(figsize=(9.8, 3.0))
    ax.set_xlim(-0.3, len(turns) + 0.2); ax.set_ylim(-1.9, 1.8); ax.axis("off")
    covered = {i for a, e in bounds for i in range(a, e)}
    for i, t in enumerate(turns):
        col = ORANGE if t == "CHI" else BLUE
        ax.add_patch(Rectangle((i + 0.05, 0), 0.9, 0.8, facecolor=col if t == "CHI" else "white", edgecolor=col, lw=1.6,
                               alpha=1 if i in covered else 0.45))
        ax.text(i + 0.5, 0.4, t, ha="center", va="center", fontsize=8, color="white" if t == "CHI" else INK)
        ax.text(i + 0.5, -0.25, str(i), ha="center", fontsize=7, color=INK2)
    for k, (a, e) in enumerate(bounds):
        y = -0.7 - 0.65 * (k % 2)
        col = AQUA if k % 2 == 0 else YELLOW
        ax.plot([a + 0.05, e - 0.05], [y, y], color=col, lw=3)
        ax.text((a + e) / 2, y - 0.33, f"Episode {k + 1}: turns {a}–{e - 1}", ha="center", fontsize=8, weight="bold",
                color=AQUA if col == AQUA else INK)
    (a1, e1) = bounds[0]
    ax.annotate("start: CHI turn after a non-CHI turn;\nwindow opens one turn earlier", xy=(a1 + 1.5, 0.85),
                xytext=(0.0, 1.3), fontsize=7.2, arrowprops=dict(arrowstyle="-", color=INK2))
    ax.annotate("end: none of the next 3 turns is CHI\n(or 8 turns, or session end)", xy=(e1 - 0.5, 0.85),
                xytext=(3.6, 1.3), fontsize=7.2, arrowprops=dict(arrowstyle="-", color=INK2))
    gap = [i for i in range(len(turns)) if i not in covered]
    ax.text(6.5, -0.8, f"faded turns ({', '.join(map(str, gap))})\nbelong to no episode", ha="center", va="center",
            fontsize=7.2, color=INK2)
    p = FIG / "r3_fig3_episode_anatomy.png"; fig.savefig(p); plt.close(fig); return p


def fig_episode_stats(ss: dict) -> Path:
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(9.6, 3.4), gridspec_kw={"width_ratios": [1, 1.25]})
    d = ss["episode_length_distribution"]
    ks = sorted(d, key=int)
    a1.bar([int(k) for k in ks], [d[k] for k in ks], color=BLUE, width=0.7, zorder=3)
    for k in ks:
        a1.text(int(k), d[k] + 60, f"{d[k]:,}", ha="center", fontsize=7.4)
    a1.set_xlabel("turns per episode"); a1.set_ylabel("episodes")
    a1.set_title(f"A. Episode length (n = {ss['n_episodes']:,})", loc="left", fontsize=9.5, weight="bold")
    a1.grid(axis="y", color=GRID, lw=0.6); a1.set_axisbelow(True)
    types = list(ss["episode_types_all"])
    allv = [ss["episode_types_all"][t] / ss["n_episodes"] * 100 for t in types]
    selv = [ss["episode_types_selected_cap3000"].get(t, 0) / ss["n_selected"]["3000"] * 100 for t in types]
    y = range(len(types))
    a2.barh([i - 0.2 for i in y], allv, height=0.38, color=GREY, label=f"all {ss['n_episodes']:,} episodes")
    a2.barh([i + 0.2 for i in y], selv, height=0.38, color=BLUE, label=f"selected {ss['n_selected']['3000']:,}")
    for i, (a, s) in enumerate(zip(allv, selv)):
        a2.text(a + 0.8, i - 0.2, f"{a:.1f}%", va="center", fontsize=7)
        a2.text(s + 0.8, i + 0.2, f"{s:.1f}%", va="center", fontsize=7)
    a2.set_yticks(list(y), [t.replace("_", " ") for t in types]); a2.invert_yaxis()
    a2.set_xlabel("% of episodes"); a2.set_xlim(0, 88)
    a2.legend(frameon=False, fontsize=7.6, loc="lower right")
    a2.set_title("B. Episode types before and after balancing", loc="left", fontsize=9.5, weight="bold")
    fig.tight_layout()
    p = FIG / "r3_fig4_episode_stats.png"; fig.savefig(p); plt.close(fig); return p


def fig_selection(per_child: list[tuple[str, int, int]]) -> Path:
    fig, ax = plt.subplots(figsize=(9.2, 3.2))
    names = [c for c, _, _ in per_child]
    ex = [e for _, e, _ in per_child]; kp = [k for _, _, k in per_child]
    x = range(len(names))
    ax.bar(x, [120] * len(names), color="white", edgecolor=GRID, width=0.7, label="pool (120 per child)")
    ax.bar(x, ex, color=GREY, width=0.7, label="examined, not admissible", zorder=2)
    ax.bar(x, kp, color=BLUE, width=0.7, label="kept (all 3 generators admissible)", zorder=3)
    for i, (e, k) in enumerate(zip(ex, kp)):
        ax.text(i, e + 2, f"{e}", ha="center", fontsize=7.3)
    ax.set_xticks(list(x), names, rotation=35, ha="right", fontsize=8)
    ax.set_ylabel("segments"); ax.set_ylim(0, 130)
    ax.legend(frameon=False, fontsize=7.6, ncol=3, loc="upper center", bbox_to_anchor=(0.5, 1.16))
    ax.grid(axis="y", color=GRID, lw=0.6); ax.set_axisbelow(True)
    fig.tight_layout()
    p = FIG / "r3_fig5_selection_order.png"; fig.savefig(p); plt.close(fig); return p


def fig_state_matrix() -> Path:
    gold = ["value", "unknown", "null"]
    model = ["same value", "different value", "unknown", "null / omitted"]
    cell = {
        ("value", "same value"): ("correct", AQUA), ("value", "different value"): ("incorrect", GREY),
        ("value", "unknown"): ("incorrect", GREY), ("value", "null / omitted"): ("missing", YELLOW),
        ("unknown", "same value"): ("—", "white"), ("unknown", "different value"): ("incorrect", GREY),
        ("unknown", "unknown"): ("correct", AQUA), ("unknown", "null / omitted"): ("missing", YELLOW),
        ("null", "same value"): ("—", "white"), ("null", "different value"): ("hallucinated", ORANGE),
        ("null", "unknown"): ("hallucinated", ORANGE), ("null", "null / omitted"): ("correctly_null", BLUE),
    }
    fig, ax = plt.subplots(figsize=(8.6, 3.0))
    ax.set_xlim(0, 4 * 1.7 + 1.6); ax.set_ylim(0, 3 * 0.8 + 0.8); ax.axis("off")
    for j, m in enumerate(model):
        ax.text(1.6 + j * 1.7 + 0.85, 3 * 0.8 + 0.35, m, ha="center", fontsize=8, weight="bold")
    ax.text(0.8, 3 * 0.8 + 0.35, "gold ↓ / model →", ha="center", fontsize=7.6, color=INK2)
    for i, g in enumerate(gold):
        yy = (2 - i) * 0.8
        ax.text(0.8, yy + 0.4, f"gold {g}", ha="center", va="center", fontsize=8, weight="bold")
        for j, m in enumerate(model):
            lab, col = cell[(g, m)]
            ax.add_patch(Rectangle((1.6 + j * 1.7 + 0.04, yy + 0.04), 1.62, 0.72, facecolor=col,
                                   edgecolor="white", lw=2, alpha=0.9 if col != "white" else 1))
            ax.text(1.6 + j * 1.7 + 0.85, yy + 0.4, lab, ha="center", va="center", fontsize=8,
                    color="white" if col in (BLUE, ORANGE) else INK)
    p = FIG / "r3_fig6_state_matrix.png"; fig.savefig(p); plt.close(fig); return p


def fig_forest(r3: dict) -> Path:
    mets = [("merged_agreement", "Merged agreement"), ("unsupported_rate", "Unsupported assertion"),
            ("whole_record", "Whole-record agreement")]
    fig, axes = plt.subplots(1, 3, figsize=(10.4, 3.5), sharey=True)
    for ax, (key, title) in zip(axes, mets):
        for k, (side, mk) in enumerate([("main_batch", "o"), ("holdout", "s")]):
            perf = r3["model_performance"][side]["per_model"]
            for i, m in enumerate(MODELS):
                v = perf[m][key]
                y = i + (k - 0.5) * 0.32
                ax.plot([100 * v["ci95"]["lower"], 100 * v["ci95"]["upper"]], [y, y], color=COLOUR[m], lw=2)
                ax.scatter([100 * v["value"]], [y], marker=mk, s=38, color=COLOUR[m] if k else "white",
                           edgecolor=COLOUR[m], linewidth=1.6, zorder=3)
        ax.set_title(title, fontsize=9.5, weight="bold", loc="left")
        ax.set_xlabel("% (95% child-cluster CI)")
        ax.grid(axis="x", color=GRID, lw=0.6); ax.set_axisbelow(True)
    axes[0].set_yticks(range(len(MODELS)), [short(m) for m in MODELS]); axes[0].invert_yaxis()
    axes[0].scatter([], [], marker="o", facecolor="white", edgecolor=INK2, label="main batch (330)")
    axes[0].scatter([], [], marker="s", color=INK2, label="original frozen evaluation set (132)")
    fig.legend(frameon=False, fontsize=8, loc="lower center", ncol=2, bbox_to_anchor=(0.5, -0.02))
    fig.tight_layout(rect=(0, 0.07, 1, 1))
    p = FIG / "r3_fig7_cluster_ci.png"; fig.savefig(p); plt.close(fig); return p


def fig_mars(me: dict, n_short: int, n_judged: int, best_on_returned: float) -> Path:
    mars = me["mars"]; ab = mars["abstention_reasons"]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(9.8, 3.1), gridspec_kw={"width_ratios": [1.5, 1]})
    parts = [("reranked", n_judged, BLUE), ("single-candidate\nshort-circuit", n_short, AQUA),
             ("no reviewable\ncandidate", ab.get("no_reviewable_candidate", 0), ORANGE),
             ("judge\ndisagreement", ab.get("judge_disagreement", 0), YELLOW),
             ("insufficient\nvalid judges", ab.get("insufficient_valid_judge_responses", 0), GREY)]
    left = 0
    for k, (lab, n, col) in enumerate(parts):
        a1.barh([0], [n], left=left, color=col, edgecolor="white", lw=2, height=0.5)
        if n >= 15:
            a1.text(left + n / 2, 0, f"{n}", ha="center", va="center", color="white" if col in (BLUE, ORANGE) else INK, fontsize=8.5, weight="bold")
        ypos = -0.5 if k < 2 else [0.62, -0.62, 0.62][k - 2]
        text = lab if n >= 15 else f"{lab} ({n})"
        a1.text(left + n / 2, ypos, text, ha="center", va="center", fontsize=6.8)
        left += n
    a1.set_xlim(0, 335); a1.set_ylim(-0.95, 0.95); a1.axis("off")
    a1.set_title(f"A. MARS outcome on all {mars['n_total']} segments (returned {mars['n_returned']}, abstained {mars['n_abstained']})",
                 loc="left", fontsize=9, weight="bold")
    labs = ["MARS\nall 330\n(abstain\n= miss)", "MARS\n265\nreturned", "qwen2.5\n:7b\nsame 265", "qwen2.5\n:7b\nall 330", "Oracle\nceiling"]
    vals = [mars["accuracy_over_all"], mars["accuracy_over_returned"], best_on_returned, me["best_generator_accuracy"], me["oracle_ceiling"]]
    cols = [BLUE, BLUE, AQUA, AQUA, GREY]
    a2.bar(range(5), [100 * v for v in vals], color=cols, width=0.62, zorder=3)
    for i, v in enumerate(vals):
        a2.text(i, 100 * v + 0.12, f"{100 * v:.1f}", ha="center", fontsize=8)
    a2.set_xticks(range(5), labs, fontsize=6.8); a2.set_ylabel("whole-record agreement (%)")
    a2.grid(axis="y", color=GRID, lw=0.6); a2.set_axisbelow(True); a2.set_ylim(0, 7.5)
    a2.set_title("B. Full-policy whole-record agreement", loc="left", fontsize=9, weight="bold")
    fig.tight_layout()
    p = FIG / "r3_fig8_mars_full_policy.png"; fig.savefig(p); plt.close(fig); return p


def fig_inputs() -> Path:
    fig, ax = plt.subplots(figsize=(9.8, 3.8))
    ax.set_xlim(0, 12); ax.set_ylim(0, 6); ax.axis("off")
    _box(ax, 4.3, 2.1, 3.4, 1.8, "Identical de-identified\nCHAT segment text\n(2–8 turns, all speakers,\n%spa %act %gpx only)", INK2, size=8.2, weight="bold")
    left = ["System prompt: fields, allowed values,", "null rule, no inference / diagnosis", "CHAT format guide", "No child / corpus / session IDs",
            "No timestamps, no goals", "(setting, role, date passed as fixed", " request metadata, never in the prompt)"]
    right = ["Segment ID, child ID, corpus, session", "(visible: needed to key the sheet)", "Blank value columns + reference sheet",
             "Codebook", "No model output, identity, confidence,", "MARS result or selection result", "(checked: no sheet contains any)"]
    _box(ax, 0.1, 0.3, 3.6, 5.3, "", BLUE)
    _box(ax, 8.3, 0.3, 3.6, 5.3, "", AQUA)
    ax.text(1.9, 5.25, "Extraction model saw", ha="center", weight="bold", color=BLUE, fontsize=9)
    ax.text(10.1, 5.25, "Annotators saw", ha="center", weight="bold", color=AQUA, fontsize=9)
    for i, t in enumerate(left):
        ax.text(0.3, 4.7 - i * 0.6, t, fontsize=7.4)
    for i, t in enumerate(right):
        ax.text(8.5, 4.7 - i * 0.6, t, fontsize=7.4)
    _arrow(ax, (4.3, 3.0), (3.7, 3.0)); _arrow(ax, (7.7, 3.0), (8.3, 3.0))
    p = FIG / "r3_fig9_inputs.png"; fig.savefig(p); plt.close(fig); return p


def fig_adequacy() -> Path:
    fig, ax = plt.subplots(figsize=(9.8, 4.2))
    ax.set_xlim(0, 12); ax.set_ylim(0, 6.2); ax.axis("off")
    _box(ax, 0.1, 4.4, 2.6, 1.4, "All VERIFIED /\nCORRECTED records\nfor the child in the\nreview period", INK2, size=7.6)
    checks = ["coverage (critical)", "field completeness (critical)", "temporal distribution",
              "≥ 2 contributor ROLES", "≥ 2 settings", "goal-domain coverage ≥ 0.50"]
    for i, c in enumerate(checks):
        col = ORANGE if "critical" in c else (YELLOW if "≥ 2" in c else BLUE)
        _box(ax, 3.3, 5.3 - i * 0.82, 3.0, 0.62, c, col, size=7.6)
    _arrow(ax, (2.7, 5.1), (3.3, 5.1))
    _box(ax, 7.0, 4.6, 4.9, 1.3, "Level: INSUFFICIENT if 0 records, or a critical check\nfails with < 3 records → all analytics withheld", ORANGE, size=7.5)
    _box(ax, 7.0, 2.9, 4.9, 1.4, "LIMITED otherwise if any check fails\n(or < 3 records): analytics computed;\nfailed representation check withholds the\nmatching cross-comparison", YELLOW, size=7.5)
    _box(ax, 7.0, 1.3, 4.9, 1.3, "ADEQUATE only if ≥ 5 records\nand every check passes", AQUA, size=7.5)
    _arrow(ax, (6.3, 3.5), (7.0, 5.2)); _arrow(ax, (6.3, 3.5), (7.0, 3.6)); _arrow(ax, (6.3, 3.5), (7.0, 1.95))
    ax.text(0.1, 0.35, "No goal-level exemption exists: the Goal model has no single-setting or single-contributor attribute, "
            "and both checks run on the child's whole record set.", fontsize=7.6, color=INK2)
    p = FIG / "r3_fig10_adequacy.png"; fig.savefig(p); plt.close(fig); return p


# ---------------------------------------------------------------------------
# Answers
# ---------------------------------------------------------------------------

def a1(b: Doc, sec, ctx):
    kc = ctx["r3"]["kappa_composition"]
    b.label("Short answer")
    b.answer_line(
        "**Recomputed over all 330, in each segment's final annotation state, before adjudication.** "
        f"For each of the four categorical fields, Cohen's κ was computed over the 330 paired values: "
        "annotator A's value against annotator B's. Each segment contributes its round-2 value if it was one "
        "of the 178 re-issued, and its round-1 value otherwise. 0.801 is the unweighted mean of the four "
        "per-field κ values. It is not a pooled κ. Adjudication plays no part in it: adjudicated values form "
        "the reference standard, not an annotator's labels.")
    b.answer_line(
        "**Your understanding is correct, and the number should not be presented as a before/after effect.** "
        f"The 152 retained segments already agreed at κ = {kc['retained_152_round1_values']['mean']:.3f} in "
        "round 1, because re-issue selection took the disagreement-enriched segments and left these. Part of "
        f"the rise from {kc['round1_all_330']['mean']:.3f} to {kc['final_state_all_330']['mean']:.3f} therefore "
        "comes from which segments were re-annotated, not from the codebook alone.")
    rows = [
        ["Round 1, all 330", "round 1", "330", f"{kc['round1_all_330']['mean']:.3f}"] + [f"{kc['round1_all_330']['per_field'][f]:.3f}" for f in ("domain", "modality", "prompt_level", "outcome")],
        ["Retained 152", "round 1 (unchanged)", "152", f"{kc['retained_152_round1_values']['mean']:.3f}"] + [f"{kc['retained_152_round1_values']['per_field'][f]:.3f}" for f in ("domain", "modality", "prompt_level", "outcome")],
        ["Re-issued 178, before", "round 1", "178", f"{kc['reissued_178_round1_values']['mean']:.3f}"] + [f"{kc['reissued_178_round1_values']['per_field'][f]:.3f}" for f in ("domain", "modality", "prompt_level", "outcome")],
        ["Re-issued 178, after", "round 2", "178", f"{kc['reissued_178_round2_values']['mean']:.3f}"] + [f"{kc['reissued_178_round2_values']['per_field'][f]:.3f}" for f in ("domain", "modality", "prompt_level", "outcome")],
        ["**Final state (reported)**", "mixed", "330", f"**{kc['final_state_all_330']['mean']:.3f}**"] + [f"{kc['final_state_all_330']['per_field'][f]:.3f}" for f in ("domain", "modality", "prompt_level", "outcome")],
    ]
    b.table("κ by annotation state. Source: `round3_analysis.json` (kappa_composition). The script asserts "
            "that its final-state and round-1 means equal `reliability_current.json` and "
            "`reliability_round1.json`, and that the 152 retained segments carry identical values in both files.",
            ["Set", "Values", "n", "Mean κ", "domain", "modality", "prompt_level", "outcome"],
            rows, [1.45, 1.05, 0.4, 0.65, 0.65, 0.7, 0.85, 0.65])
    b.figure(FIG / "r3_fig1_kappa_composition.png",
             "Composition of the reported κ. The 152 retained segments were high-agreement to begin with. The "
             f"re-issued 178 rose from {kc['reissued_178_round1_values']['mean']:.3f} to "
             f"{kc['reissued_178_round2_values']['mean']:.3f}.")
    b.label("How to describe it")
    b.boxed("Suggested wording", [
        "\"After re-annotation of the 178 disagreement-enriched segments under the codebook, mean κ over all 330 "
        "segments, each in its final pre-adjudication state (round-2 values for the 178, round-1 values for the "
        "remaining 152), was 0.801. Because the 152 were not re-annotated and had higher round-1 agreement "
        "(0.884), this figure should not be read as a before/after comparison. The change within the 178 "
        "re-issued segments provides the more direct descriptive comparison of agreement before and after "
        "clarification. Domain, whose coding rules were unchanged, provides a useful internal reference field. "
        "Because the same annotators coded these segments a second time, familiarity and round effects may "
        "also contribute to that change.\"",
        "Also relabel the table row that currently reads \"after codebook re-annotation\" (make_manuscript.py:791). "
        "Only 178 of the 330 were re-annotated.",
    ], fill="EAF7F1", edge="1BAF7A")


def a2(b: Doc, sec, ctx):
    ss = ctx["ss"]
    bl = bullets_of(sec["items"])
    b.label("Short answer")
    b.answer_line(
        f"**One model input is a window of 2–8 consecutive transcript turns (mean {ss['mean_episode_length']}) "
        "built around child turns, including every speaker's turns inside the window.** Segmentation is "
        f"rule-based, runs session by session, and turns {ss['n_utterances']:,} utterances in "
        f"{ss['n_sessions']} sessions into {ss['n_episodes']:,} episodes. Two runs produce identical output. "
        "Code: `data/pipeline/segment_episodes.py`; measured by `segmentation_stats.py` → `segmentation_stats.json`.")
    ans = [
        "A child (`CHI`) turn that either carries a behavioural tier (%spa, %act or %gpx), is the first "
        "utterance of the session, or follows a non-child turn. The window opens **one turn before** the "
        "triggering child turn when a preceding turn exists, so the adult turn that prompted it is included; "
        "if the child turn is the first utterance of the session, the window begins with that child turn.",
        "Checked at each following turn: the window closes when it reaches 8 turns, when the session ends, "
        "or when none of the next 3 utterances is a child turn (once it has at least 2 turns). The closing "
        "turn is included.",
        "One preceding turn before the triggering child turn. After it, turns continue until the end rule "
        "fires, up to 8 in total. Nothing is added after the closing turn. Measured length: "
        + ", ".join(f"{k} turns: {v:,}" for k, v in ss["episode_length_distribution"].items())
        + f". Most episodes ({ss['episode_length_distribution']['8']:,}) hit the 8-turn cap.",
        "All speakers inside the window: "
        + ", ".join(f"{k} {v:,}" for k, v in ss["speakers_in_episodes"].items())
        + ". Of the dependent tiers, only %spa, %act and %gpx are carried into the segment. %mor, %gra, "
        "%com, %tim (timing), %xpho and %gls are parsed but dropped. No @header metadata is included.",
        f"**Yes, by one turn.** Because a new window opens one turn before its triggering child turn, it can "
        f"begin on the previous episode's closing turn. {ss['adjacent_episode_pairs_sharing_an_utterance']:,} "
        "adjacent episode pairs share exactly one utterance. Episodes never overlap by more than that.",
        "Eight types, assigned by the first matching rule in this order: request (%spa REQ/QN/$Q), imitation "
        "(IM), protest (%spa NO/PR, or push/turn away, head shake in %gpx), joint attention (%gpx point/show/"
        "look at/gaze), response (%spa AN/RP/$R), play (play vocabulary in %act/%gpx/text), initiation (the "
        "window opens on a child turn), and other. Imitation never occurs in these corpora. Distribution: see "
        "Figure 4B.",
        f"Two rules exist (windows shorter than 2 turns, and windows with no child turn). Measured on this "
        f"corpus they exclude **{ss['excluded_shorter_than_min_turns']}** and "
        f"**{ss['excluded_no_child_turn']}** windows. Every window is kept, so the balancing step sees all "
        f"{ss['n_episodes']:,}. Balancing then keeps, per child, up to ⌈cap/11⌉ episodes, split equally across "
        "the child's episode types. Within each type it takes the earliest in transcript order, with no "
        "randomness.",
        "**Yes.** There is no random step anywhere in parsing, de-identification, segmentation or selection. "
        "Two consecutive runs gave the identical episode list, SHA-256 "
        f"`{ss['episode_fingerprint_sha256'][:16]}…`. One caveat, which matters for Q3: the segment ID "
        "number is assigned in selection order, so the same ID can name a different episode under a different "
        "cap.",
    ]
    for q, a in zip(bl, ans):
        b.sub(q)
        b.p(a)
    b.figure(FIG / "r3_fig3_episode_anatomy.png",
             "The start and end rules on an illustrative sequence of turns (not corpus text). The episode boundaries "
             "are computed by the pipeline's own `_find_episode_boundaries`, not drawn by hand.")
    b.figure(FIG / "r3_fig4_episode_stats.png",
             "A: most episodes reach the 8-turn cap. B: \"other\" is 78% of all episodes. The per-type quota "
             "reduces it to 39% of the 2,442 selected segments. Source: `segmentation_stats.json`.")
    b.figure(FIG / "r3_fig2_pipeline.png",
             "From utterances to model inputs, for both samples. The original frozen evaluation branch is "
             "where the ID-instability issue surfaces (Q3).")


def a3(b: Doc, sec, ctx):
    bl = bullets_of(sec["items"])
    per_child, leak = ctx["per_child"], ctx["r3"]["holdout_leakage"]
    b.label("Short answer")
    b.answer_line(
        "**Stage 1 shuffled each child's segments with a fixed per-child seed and kept the first 120. "
        "Stage 2 walked those 120 in shuffled order and kept each segment where all three generators gave an "
        "admissible response, stopping at 30 per child.** Script: "
        "`tests/eval/extraction/build_annotation_batch.py` (`build_pool`, `extract_candidates`); seed 42; "
        "artifacts `annotation_pool.json` and `annotation_batch.json`.")
    ans = [
        "Child by child, in alphabetical order of child ID. Within each child, in the pool's shuffled order. "
        "Stage 2 visits every pool entry and skips a child's remaining entries once that child has 30 kept. "
        "The examined set is therefore, for each child, the prefix of its shuffled list up to its 30th kept "
        "segment. Reconstructed this way, it contains exactly 506 segments, matching `segments_attempted` "
        "in the artifact.",
        "Yes. Each child's candidates (from the 2,442) were shuffled before the first 120 were taken, and the "
        "1,320-segment pool keeps that order.",
        "Yes: `random.Random(f\"42:{child_id}\")`, i.e. seed 42 combined with the child ID. Seeding per child "
        "means adding or removing a child does not reshuffle the others.",
        "Yes. For each child, the 30 are the first 30 examined segments where **every** generator returned "
        "an admissible response (`require=\"all\"`). Under the `judgeable` rule, admissible means the "
        "response contains at least one populated extracted field, whatever status the model self-reported.",
        "One upstream filter: segments rendering to 20 characters or fewer are dropped before pooling (none "
        "were). There is no other sampling, weighting or manual choice. Annotator assignment was a separate "
        "step: both annotators received all 330 (100% overlap), each in their own seeded row order.",
    ]
    for q, a in zip(bl, ans):
        b.sub(q)
        b.p(a)
    b.table("Examined and kept, per child. Reconstructed from the artifacts; the total reproduces the recorded 506.",
            ["Child", "Pool", "Examined", "Kept", "Joint yield"],
            [[c, "120", str(e), str(k), pct(k / e)] for c, e, k in per_child]
            + [["**Total**", "1,320", f"**{sum(e for _, e, _ in per_child)}**", f"**{sum(k for _, _, k in per_child)}**",
                pct(sum(k for _, _, k in per_child) / sum(e for _, e, _ in per_child))]],
            [1.9, 0.8, 1.0, 0.8, 1.0])
    b.figure(FIG / "r3_fig5_selection_order.png",
             "How far into each child's shuffled pool the builder had to go to reach 30 segments where all "
             "three generators were admissible.")

    b.label("Evaluation-set overlap identified during review")
    b.boxed("What was found (2026-09-26)", [
        f"`build_holdout_set.py` excluded the 1,320-segment pool by **segment ID**. Segment IDs are numbered "
        "in selection order within each session, so under the original evaluation-set build's larger cap (455 per child, "
        "against 273) the numbering shifts: 338 of the 2,442 IDs name a different episode. Matched by "
        f"transcript text and confirmed by source line numbers, **{leak['n_in_pool']} of the 132 original frozen "
        f"evaluation segments are pool episodes**. Of those, {leak['n_model_output_examined']} had model output examined "
        f"during batch construction, and {leak['n_annotated_in_main_batch']} was annotated in the main batch.",
        "This also corrects my answer in the previous round, which said none of the 132 were in the pool. "
        "That check compared IDs, so it could not see this.",
    ], fill="FDECEA", edge="D03B3B")
    sens = ctx["r3"]["holdout_sensitivity"]
    base = ctx["r3"]["model_performance"]["holdout"]["per_model"]
    rows = []
    for m in MODELS:
        rows.append([short(m)] + [
            f"{pct(base[m][k]['value'])} → {pct(sens['excluding_all_pool_episodes']['per_model'][m][k]['value'])}"
            for k in ("merged_agreement", "unsupported_rate", "whole_record")])
    b.table("Original frozen evaluation set (132) → overlap-excluded evaluation set (121, still 11 children). "
            "Source: `round3_analysis.json` (holdout_sensitivity).",
            ["Generator", "Merged agreement", "Unsupported assertion", "Whole-record"], rows, [1.3, 1.8, 1.8, 1.8])
    main = ctx["r3"]["model_performance"]["main_batch"]["per_model"]
    excl = sens["excluding_all_pool_episodes"]["per_model"]
    # The claims below are checked, not asserted by hand.
    assert all(base[m]["merged_agreement"]["value"] < main[m]["merged_agreement"]["value"]
               and excl[m]["merged_agreement"]["value"] < main[m]["merged_agreement"]["value"] for m in MODELS)
    assert all(ranking(base, k) == ranking(excl, k) for k in METRICS3)
    wr, ur = "whole_record", "unsupported_rate"
    assert all(base[m][wr]["value"] >= main[m][wr]["value"] and excl[m][wr]["value"] >= main[m][wr]["value"]
               for m in ("qwen2.5:7b", "qwen3:14b"))
    assert base["llama3.2:latest"][ur]["value"] > main["llama3.2:latest"][ur]["value"]
    pd = sens["excluding_all_pool_episodes"]["paired_qwen3_14b_minus_qwen2_5_7b"]
    assert not pd["merged_agreement"]["ci95"]["excludes_zero"] and pd["unsupported_rate"]["ci95"]["excludes_zero"]
    b.p("**No conclusion changes.** On both the 132 and the 121, every generator's merged agreement is lower "
        "than on the main batch. (This does not hold for every metric: whole-record agreement for qwen2.5:7b "
        "and qwen3:14b is equal or higher, and on the 132 llama3.2's unsupported-assertion rate is higher, "
        "all within wide intervals.) The order of generators on each of the three metrics is the same in the 132 and the "
        "121. The qwen3:14b minus qwen2.5:7b difference in merged agreement is still not distinguishable from "
        f"zero ({diff_str(pd['merged_agreement'])} points on the 121). The unsupported-assertion difference "
        f"still excludes zero ({diff_str(pd['unsupported_rate'])}).")
    b.boxed("Recommendation for the manuscript", [
        "Use two names consistently: the **original frozen evaluation set (n = 132)** and the "
        "**overlap-excluded evaluation set (n = 121)**. Do not call the 132 simply \"held-out\": 11 of its "
        "segments break the independence criterion it was meant to meet.",
        "Report the 121 as the primary estimate of performance on unseen episodes. Report the 132 as the "
        "originally frozen evaluation, labelled as the set drawn before the identifier defect was found. "
        "Neither set is re-drawn, so no set is selected after its results were seen. Suggested sentence: "
        "\"We report the originally frozen 132-segment evaluation for transparency and a "
        "contamination-excluded analysis of 121 segments. Conclusions were unchanged.\"",
        "The manuscript's claim that \"the entire original stratified pool … is excluded\" "
        "(make_manuscript.py:446) must be removed. Fix the builder to exclude by episode identity before "
        "any future draw.",
    ], fill="EAF7F1", edge="1BAF7A")


def a4(b: Doc, sec, ctx):
    bl = bullets_of(sec["items"])
    c = ctx["counts_q7"]
    b.label("Short answer")
    b.answer_line(
        "**Each field is first given one of five judgements by comparing the model's value with the gold "
        "value (score_against_gold.judge_field). All four metrics are counts of those judgements over the "
        "four categorical fields, with the free-text behavior field excluded.** Code: `score_against_gold.py`, "
        "`paper_metrics.summarise_scored`, `mars_evaluation.is_fully_correct`.")
    b.figure(FIG / "r3_fig6_state_matrix.png",
             "Judgement for every pair of gold state and model state. A dash means the combination cannot "
             "occur: a model value cannot equal a gold value that is unknown or null.")
    defs = [
        ("merged categorical agreement", "For each field: (correct + correctly_null) ÷ all N scored segments. Then the unweighted mean over the four fields.",
         f"qwen2.5:7b, batch: {pct(c['merged'])}"),
        ("whole-record agreement", "Share of segments where all four fields are correct or correctly_null. Denominator: all N scored segments.",
         f"{c['whole_n']}/330 = {pct(c['whole_n'] / 330)}"),
        ("mean field abstention", "For each field: missing ÷ N. Then the mean over the four fields. Counts only cases where gold has a value and the model gave none. A null where gold is also null is correctly_null, not abstention.",
         f"{pct(c['abstain'])}"),
        ("unsupported assertion", "Σ hallucinated ÷ Σ (hallucinated + correctly_null), pooled across the four fields (not averaged per field). Denominator: every field where gold is null.",
         f"{c['hall']}/{c['opp']} = {pct(c['hall'] / c['opp'])}"),
    ]
    for (name, formula, ex), q in zip(defs, bl[:4]):
        b.sub(q)
        b.p(f"**{formula}** Worked value: {ex}.")
    b.table("Judgement counts behind the four metrics: qwen2.5:7b, main batch, 4 fields × 330 = 1,320 field "
            "judgements. Source: `scored_against_gold_augmented.json`.",
            ["correct", "incorrect", "missing", "hallucinated", "correctly_null", "total"],
            [[str(c[k]) for k in ("correct", "incorrect", "missing", "hallucinated", "correctly_null")] + [str(c["total"])]],
            [1.1, 1.1, 1.1, 1.2, 1.2, 0.9])
    states = [
        "An exact match (after case and whitespace are normalised) is correct; otherwise incorrect. If gold is null, any model value is hallucinated.",
        "A null on either side is a distinct state. Gold null with model null is correctly_null, which counts as agreement and as a non-hallucinated opportunity. Gold value with model null is missing. The literal strings \"null\"/\"none\" are converted to JSON null before scoring (`ollama_llm._normalize_null_strings`).",
        "\"unknown\" is an ordinary category, not a null. Gold unknown with model unknown is correct. Model unknown where gold is null is **hallucinated**, i.e. an unsupported assertion. Gold unknown with model null is missing. "
        f"For qwen2.5:7b on the main batch, {c['unknown_hall']} of its {c['hall']} unsupported assertions were the value \"unknown\".",
        "**Scored exactly as JSON null.** The scorer reads a missing field, or a whole missing candidate, as None, so the manuscript must not describe omission as a separate state. On the main batch this cannot occur at candidate level, because inclusion required every generator to respond. On the 132-segment evaluation set, a generator with no admissible response has all four fields judged as null.",
    ]
    for q, a in zip(bl[4:], states):
        b.sub(q)
        b.p(a)


def a5(b: Doc, sec, ctx):
    r3 = ctx["r3"]
    b.label("Short answer")
    b.answer_line(
        "**Yes, and the results are below.** `tests/eval/extraction/round3_analysis.py` computes them from the "
        "saved scored outputs (`scored_against_gold_augmented.json`, `scored_holdout.json`) with no model "
        "call. It uses the same child-cluster percentile bootstrap as the annotation κ: 10,000 replicates, "
        "children drawn with replacement, all segments of each drawn child kept. Seed 20260926. Output: "
        "`round3_analysis.json`. Its point estimates are checked against `holdout_metrics.json`.")
    for side, n in (("main_batch", 330), ("holdout", 132)):
        perf = r3["model_performance"][side]["per_model"]
        rows = [[short(m)] + [
            f"{pct(perf[m][k]['value'])} [{pct(perf[m][k]['ci95']['lower'])}, {pct(perf[m][k]['ci95']['upper'])}]"
            for k in ("merged_agreement", "unsupported_rate", "whole_record")] for m in MODELS]
        b.table(f"{'Main batch' if side == 'main_batch' else 'Original frozen evaluation set'} (n = {n}, 11 children): point "
                "estimate [95% child-cluster CI].",
                ["Generator", "Merged agreement", "Unsupported assertion", "Whole-record"], rows, [1.2, 1.85, 1.85, 1.85])
    mb, ho = r3["model_performance"]["main_batch"]["per_model"], r3["model_performance"]["holdout"]["per_model"]
    u = "unsupported_rate"
    assert not any(overlaps(mb[a][u], mb[b][u]) for i, a in enumerate(MODELS) for b in MODELS[i + 1:])
    assert overlaps(ho["qwen2.5:3b"][u], ho["qwen2.5:7b"][u]) and overlaps(ho["llama3.2:latest"][u], ho["qwen3:14b"][u])
    b.figure(FIG / "r3_fig7_cluster_ci.png",
             "Child-cluster 95% intervals. Merged-agreement intervals for qwen2.5:7b and qwen3:14b overlap "
             "heavily. On the main batch, the unsupported-assertion intervals of the four generators do not "
             "overlap. On the 132-segment set they overlap for qwen2.5:3b and qwen2.5:7b, and for llama3.2 and "
             "qwen3:14b.")
    rows = []
    paired = [("Main batch (330)", r3["model_performance"]["main_batch"]["paired_qwen3_14b_minus_qwen2_5_7b"]),
              ("Original frozen (132)", r3["model_performance"]["holdout"]["paired_qwen3_14b_minus_qwen2_5_7b"]),
              ("Overlap-excluded (121)", r3["holdout_sensitivity"]["excluding_all_pool_episodes"]["paired_qwen3_14b_minus_qwen2_5_7b"])]
    for label, p in paired:
        rows.append([label] + [diff_str(p[k]) + (" *" if p[k]["ci95"]["excludes_zero"] else "") for k in METRICS3])
    b.table("Paired difference, qwen3:14b minus qwen2.5:7b, in percentage points. Both models are scored on "
            "the same resampled children. * marks an interval excluding zero.",
            ["Set", "Merged agreement", "Unsupported assertion", "Whole-record"], rows, [1.2, 1.85, 1.85, 1.85])
    hp = r3["model_performance"]["holdout"]["paired_qwen3_14b_minus_qwen2_5_7b"]
    ep = r3["holdout_sensitivity"]["excluding_all_pool_episodes"]["paired_qwen3_14b_minus_qwen2_5_7b"]
    assert hp["unsupported_rate"]["ci95"]["excludes_zero"] and ep["unsupported_rate"]["ci95"]["excludes_zero"]
    wr = "whole_record"
    assert all(ho[m][wr]["ci95"]["lower"] == 0 for m in MODELS)
    zero_main = [short(m) for m in MODELS if mb[m][wr]["ci95"]["lower"] <= 0]
    n7 = round(ho["qwen2.5:7b"][wr]["value"] * 132)
    b.p("This answers the question left open in Round 2. On the original frozen evaluation set, qwen3:14b's "
        f"{100 * hp['merged_agreement']['value']:.1f}-point lead in merged agreement is not distinguishable "
        f"from zero; on the overlap-excluded 121 it is {100 * ep['merged_agreement']['value']:+.1f} points, "
        "also not distinguishable. The higher unsupported-assertion rate is distinguishable from zero in both "
        f"analyses: {diff_str(hp['unsupported_rate'])} percentage points on the original frozen set and "
        f"{diff_str(ep['unsupported_rate'])} on the overlap-excluded set. On the 132-segment set every whole-record "
        f"interval has a lower bound of zero: qwen2.5:7b's {pct(ho['qwen2.5:7b'][wr]['value'])} is {n7} segments, "
        "and llama3.2 and qwen2.5:3b scored 0. On the main batch, the intervals for "
        f"{' and '.join(zero_main)} reach zero. Report whole-record values as descriptive, not as a ranking.")


def a6(b: Doc, sec, ctx):
    me = ctx["mars_eval"]
    bl = bullets_of(sec["items"])
    b.label("Short answer")
    b.answer_line(
        "**Yes, and it already exists: mars_evaluation.py → mars_evaluation.json.** It scores the complete "
        "MARS policy over all 330 segments at whole-record level and reports it two ways: "
        "`accuracy_over_all`, where **every abstention counts as a miss**, and `accuracy_over_returned`, "
        "conditional on MARS returning a record. Each is reported with `coverage` and the reason for each "
        "abstention. No new definition is needed; the one to cite is `accuracy_over_all`, alongside "
        "coverage.")
    rows = [
        [bl[0], str(ctx["n_short"]), "returned", "Scored on the passed-through candidate"],
        [bl[1], str(me["mars"]["abstention_reasons"].get("no_reviewable_candidate", 0)), "abstained", "Miss in accuracy_over_all"],
        [bl[2], str(me["mars"]["abstention_reasons"].get("judge_disagreement", 0)), "abstained", "Miss in accuracy_over_all"],
        [bl[3], str(me["mars"]["abstention_reasons"].get("insufficient_valid_judge_responses", 0)), "abstained", "Miss in accuracy_over_all"],
        [bl[4], str(ctx["n_judged"]), "returned", "Scored on the selected candidate"],
        ["**Total**", str(me["mars"]["n_total"]), f"{me['mars']['n_returned']} returned / {me['mars']['n_abstained']} abstained",
         f"coverage {pct(me['mars']['coverage'])}"],
    ]
    b.table("How each MARS outcome enters the full-policy score. Source: `mars_evaluation.json`, `mars_selection_analysis.json`.",
            ["Outcome (as listed)", "n", "State", "Scoring"], rows, [2.4, 0.6, 1.9, 1.8])
    b.figure(FIG / "r3_fig8_mars_full_policy.png",
             f"A: the complete policy over {me['mars']['n_total']}. B: MARS at whole-record level "
             f"({pct(me['mars']['accuracy_over_all'])} counting abstentions as misses; "
             f"{pct(me['mars']['accuracy_over_returned'])} over the {me['mars']['n_returned']} it returned) against "
             f"always using qwen2.5:7b ({pct(ctx['best_on_returned'])} on the same {me['mars']['n_returned']}, "
             f"{pct(me['best_generator_accuracy'])} on all {me['mars']['n_total']}) and the oracle ceiling "
             f"({pct(me['oracle_ceiling'])}, 3-generator roster).")
    mm = me["mars"]
    b.boxed("Suggested manuscript wording", [
        f"\"Across all {mm['n_total']} episodes, MARS returned a record for {mm['n_returned']} "
        f"({pct(mm['coverage'])} coverage) and achieved {pct(mm['accuracy_over_all'])} whole-record agreement "
        "when abstentions were counted as failures; among returned records, whole-record agreement was "
        f"{pct(mm['accuracy_over_returned'])}.\"",
        f"Never report {pct(mm['accuracy_over_all'])} without the {pct(mm['coverage'])} coverage beside it.",
    ], fill="EAF7F1", edge="1BAF7A")
    b.p("Per-field full-policy numbers do not exist, and should not be invented for the paper. Per-field "
        "selection quality is defined only on the 142 genuinely reranked segments (`mars_selection_analysis`: "
        f"MARS {pct(ctx['sel']['mars'])} vs best generator {pct(ctx['sel']['best_generator'])} vs random "
        f"candidate {pct(ctx['sel']['random_candidate'])}, merged). That scope is correct for the question it "
        "answers, whether selection chose well. The full-policy figure answers a different question: what "
        "the system returns.")


def a7(b: Doc, sec, ctx):
    bl = bullets_of(sec["items"])
    b.label("Short answer")
    b.answer_line(
        "**Models and annotators saw identical transcript content, character for character, checked on all "
        "330 annotation-sheet rows; the annotation sheets additionally contained administrative identifiers "
        "(segment, child, corpus and session ID). The annotators did not see any model information: none of "
        "the 16 annotation or adjudication sheets contains a model name, confidence, candidate or MARS "
        "field.**")
    ans = [
        "Surrounding turns too: the full 2–8 turn window, with every speaker's turns (Q2), not the child's utterance alone.",
        "Yes: `*CHI:`, `*MOT:`, `*INV:` and so on, plus a guide explaining the codes (CHI child, INV investigator, MOT mother, FAT father).",
        "Only %spa, %act and %gpx, with a guide to each. %mor, %gra, %com and %tim are removed.",
        "None. No timestamps, dates, session IDs or header metadata appear in the prompt. The request carries fixed metadata (setting clinic, role interventionist, date 2026-08-28), which the provider attaches to the output record. It is never rendered into the prompt.",
        "None in the prompt. Child ID and session ID are fields on the request object, used only to label the output.",
        "None. The prompt gives no goals and no clinical context. `goal_ids` was empty for the research runs. The system prompt lists the fields and allowed values, and instructs: use null when the text does not clearly support a value; do not infer; do not diagnose, recommend treatment or make causal claims.",
    ]
    for q, a in zip(bl[:6], ans):
        b.sub(q)
        b.p(a)
    b.sub("Did the two specialists see exactly the same source representation?")
    b.p("The transcript cell of every annotation sheet is the `chat_text` string the model received, compared "
        "on all 330 rows: 330/330 identical. Unlike the model, the sheet also shows segment ID, child ID, "
        "corpus and session ID, because the importer keys on them. Each annotator received their own seeded "
        "row order, so fatigue did not fall on the same segments for both.")
    b.sub("During independent annotation and adjudication, did they see: " + "; ".join(x.rstrip(";.") for x in bl[6:]) + "?")
    if True:
        b.p("**Not shown.** The sheet columns are fixed: segment ID, child ID, corpus, session ID, transcript, "
            "five value columns, notes. The adjudication sheet adds only the two annotators' own values and "
            "the codebook rule. A search of all 16 sheets for model names, confidence, candidate and MARS "
            "fields returned nothing.")
    b.figure(FIG / "r3_fig9_inputs.png", "One text, two views: what the extraction model saw and what the annotators saw.")


def a8(b: Doc, sec, ctx):
    bl = bullets_of(sec["items"])
    b.label("Short answer")
    b.answer_line(
        "**Representation is not assessed per goal. Both checks always run over the child's full set of "
        "professionally accepted records in the review period, and no goal can be exempted.** Only goal "
        "coverage looks at goals, and only at their domains. Code: `src/backend/ml/adequacy_router.py` "
        "(`_check_contributor_representation`, `_check_setting_representation`, `_compute_level`).")
    ans = [
        "Yes, on every assessment, including when analytics are requested for a single goal. The router is not "
        "given the goal ID and does not filter records by goal.",
        "No. The `Goal` model has type, domain, title, description, criteria and parent, but no setting or "
        "contributor attribute, and nothing in the router reads such an attribute.",
        "Not applicable. There is no goal-specific exemption, so the check never becomes \"not applicable\".",
        "**Yes.** Data from a single setting or a single contributor role can never be rated ADEQUATE. That "
        "level requires every check to pass. The failure also withholds the matching comparison: "
        "`cross_setting_comparison` or `cross_contributor_comparison`. Alone, it does not make the level "
        "INSUFFICIENT: the level is LIMITED, and other analytics are still computed. Note also that "
        "\"contributors\" means distinct contributor **roles** (parent, teacher and so on), not people: two "
        "parents count as one.",
    ]
    for q, a in zip(bl[2:6], ans):
        b.sub(q)
        b.p(a)
    b.figure(FIG / "r3_fig10_adequacy.png", "What the router actually does. There is no goal-level branch.")
    b.boxed("Suggested manuscript wording", [
        "\"Representation is assessed over all accepted records for the child in the review period: at least two "
        "contributor roles and at least two settings. These checks are not goal-specific, and the prototype "
        "provides no mechanism to declare a goal single-setting or single-contributor. A failed "
        "representation check withholds the corresponding cross-comparison and prevents an ADEQUATE rating, "
        "but does not by itself withhold other analytics.\"",
    ], fill="EAF7F1", edge="1BAF7A")


def a9(b: Doc, sec, ctx):
    bl = bullets_of(sec["items"])
    b.label("Short answer")
    b.answer_line(
        "**Yes for all five. They pass: 120 tests across 10 suites, run 2026-09-26.** Two of the five, access "
        "to another child by a contributor and loss of access after revocation, had no real-database test "
        "until this round. Four integration tests were added in "
        "`tests/phase6/integration/test_authorization_e2e.py`. The revocation test was confirmed to fail when "
        "the `revoked_at` filter is removed.")
    rows = [
        [bl[0], "test_authorization_e2e::test_contributor_confined_to_assigned_child (new, real DB); test_security_rbac::test_cross_child_* ; test_children::test_get_child_not_authorized_returns_403", "pass"],
        [bl[1], "test_authorization_e2e::test_contributor_cannot_create_or_edit_goals (new); test_rbac::test_contributor_cannot_create_goal / _update_goal; test_goals::test_contributor_cannot_create_goal", "pass"],
        [bl[2], "test_structured_records::test_contributor_cannot_verify / _cannot_correct; test_rbac::test_contributor_cannot_correct_record; governance E2E S6 (live extractor, 2026-09-25)", "pass"],
        [bl[3], "test_authorization_e2e::test_revoked_contributor_loses_access_immediately (new): same token refused on the next request; test_security_rbac::test_deactivated_user_rejected", "pass"],
        [bl[4], "test_authorization_e2e::test_unrelated_professional_cannot_reach_analytics_or_reports (new); test_analytics_api::test_contributor_denied_analytics; test_reports_api::test_contributor_denied_generate / _list / _flag; test_rbac::test_contributor_cannot_view_analytics / _generate_report / _view_report", "pass"],
    ]
    b.table("Evidence for each authorization claim. Paths are under tests/. Run record: 120 passed, 0 failed, "
            "2026-09-26, commit in this round's release.",
            ["Claim (as asked)", "Tests", "Result"], rows, [1.9, 4.2, 0.6])
    b.p("These tests provide evidence for the claimed authorization behavior under the tested conditions; "
        "they do not prove the system secure. The mocked-database suites (`tests/phase3/test_rbac.py`, "
        "`tests/phase6/test_security_rbac.py`) show that each endpoint applies its role check. The new "
        "real-database tests show that the child-level decision holds against real data: in the test, a "
        "revoked contributor is refused on the very next request, using a token issued before revocation. "
        "This is consistent with enforcement per request in the API, not only in the UI or at login.")
    b.code(["uv run pytest tests/phase3/test_rbac.py tests/phase6/test_security_rbac.py tests/phase3/test_structured_records.py \\",
            "  tests/phase4/test_analytics_api.py tests/phase4/test_reports_api.py tests/phase3/test_children.py \\",
            "  tests/phase3/test_goals.py tests/phase3/test_observations.py tests/phase3/test_actors.py \\",
            "  tests/phase6/integration/test_authorization_e2e.py -q        # 120 passed"])


def a10(b: Doc, sec, ctx):
    b.label("Short answer")
    b.answer_line(
        "**Yes to both.** The reported κ intervals are 10,000-replicate percentile bootstraps (2.5th and 97.5th "
        "percentiles, linear interpolation). Each replicate draws 11 children with replacement and keeps "
        "every segment of each drawn child. Code: `reliability_stats.bootstrap_kappa` with `clusters`; seed "
        "20260915. A replicate whose κ is undefined (a resample with a single category) is dropped. This "
        "happened in three child-cluster intervals, all for modality: original frozen evaluation set (9,997 of 10,000 valid), and "
        "round-comparison round 1 and its difference (9,999 each). Every other interval has 10,000.")
    b.table("The two procedures, and where each appears. Both are stored side by side in every reliability artifact.",
            ["", "Primary", "Sensitivity analysis"],
            [["Unit resampled", "child (cluster)", "segment"],
             ["Draws per replicate", "11 children, all their segments", "N segments"],
             ["Replicates / seed", "10,000 / 20260915", "10,000 / 20260915"],
             ["Interval", "percentile 2.5–97.5", "percentile 2.5–97.5"],
             ["Artifact key", "cluster_bootstrap_ci", "bootstrap_ci"],
             ["Tables that print it", "κ tables (make_manuscript.py:689, 810)", "none"]],
            [1.8, 2.5, 2.4])
    b.boxed("One sentence in the manuscript mixes the two", [
        "make_manuscript.py:558 reads: \"Intervals are 10,000-replicate percentile bootstraps resampling "
        "segments\". The preceding paragraph and every κ table use the child-cluster interval. Suggested "
        "replacement: \"Intervals are 10,000-replicate child-cluster percentile bootstrap intervals, "
        "resampling children with replacement and retaining all segments belonging to each sampled child. "
        "Segment-level intervals are reported as a sensitivity analysis.\"",
        "Separately, the MARS paired differences (make_manuscript.py:1478) are segment-level bootstraps, and "
        "the caption says so. The model-performance intervals added in Q5 use the child-cluster procedure, "
        "so all intervals for annotation and model performance would then share one method.",
    ])


ANSWERS = [a1, a2, a3, a4, a5, a6, a7, a8, a9, a10]


# ---------------------------------------------------------------------------
# Context
# ---------------------------------------------------------------------------

def context() -> dict[str, Any]:
    r3, ss = load("round3_analysis.json"), load("segmentation_stats.json")
    pool = load("annotation_pool.json")["pool"]
    batch = load("annotation_batch.json")
    kept = {c["segment_id"] for c in batch["candidates"]}
    by_child = defaultdict(list)
    for p in pool:
        by_child[p["child_id"]].append(p["segment_id"])
    per_child = []
    for child in sorted(by_child):
        ids = by_child[child]
        last = max(i for i, s in enumerate(ids) if s in kept)
        per_child.append((child, last + 1, sum(s in kept for s in ids)))
    assert sum(e for _, e, _ in per_child) == batch["yield"]["segments_attempted"]

    sc = load("scored_against_gold_augmented.json")
    fields, m = sc["scoring_fields"], "qwen2.5:7b"
    cnt = defaultdict(int)
    for seg in sc["gold"].values():
        for f in fields:
            cnt[seg[m][f]] += 1
    whole_n = sum(all(seg[m][f] in ("correct", "correctly_null") for f in fields) for seg in sc["gold"].values())
    hm = load("holdout_metrics.json")["batch"]["per_model"][m]
    gold = load("gold_labels.json")["gold"]
    aug = {c["segment_id"]: c for c in load("annotation_batch_augmented.json")["candidates"]}
    unknown_hall = sum(
        1 for sid, g in gold.items() for f in fields
        if g["field_values"].get(f) is None
        and str((aug[sid]["candidates"].get(m) or {}).get("extracted_fields", {}).get(f) or "").lower() == "unknown")
    counts = {**cnt, "total": sum(cnt.values()), "whole_n": whole_n, "merged": hm["mean_field_accuracy"],
              "abstain": hm["mean_abstention_rate"], "hall": hm["unsupported_n"],
              "opp": hm["unsupported_opportunities"], "unknown_hall": unknown_hall}
    assert counts["hall"] == cnt["hallucinated"]

    msa = load("mars_selection_analysis.json")
    outcomes = load("mars_outcomes.json")
    base = load("scored_against_gold.json")["gold"]
    returned = [s for s, o in outcomes.items() if o.get("status") == "success"]
    best_on_returned = sum(all(base[s]["qwen2.5:7b"][f] in ("correct", "correctly_null") for f in fields)
                           for s in returned) / len(returned)
    return {"r3": r3, "ss": ss, "per_child": per_child, "counts_q7": counts,
            "mars_eval": load("mars_evaluation.json"), "n_short": msa["n_short_circuit"],
            "n_judged": msa["n_judged_segments"], "sel": msa["mean_accuracy"],
            "best_on_returned": best_on_returned}


def build() -> Path:
    FIG.mkdir(parents=True, exist_ok=True)
    ctx = context()
    fig_kappa(ctx["r3"]["kappa_composition"])
    fig_pipeline(ctx["ss"], ctx["r3"]["holdout_leakage"])
    fig_episode_anatomy()
    fig_episode_stats(ctx["ss"])
    fig_selection(ctx["per_child"])
    fig_state_matrix()
    fig_forest(ctx["r3"])
    fig_mars(ctx["mars_eval"], ctx["n_short"], ctx["n_judged"], ctx["best_on_returned"])
    fig_inputs()
    fig_adequacy()

    sections = read_questions()
    assert len(sections) == 10, len(sections)
    b = Doc()
    grey = RGBColor(0x52, 0x51, 0x4E)
    b.doc.add_heading("Response to Round 3 Review Comments", level=0)
    b.p("Answers the ten questions in 'Roun3-Review comments.docx', in their original order and wording. "
        "Branch model-finetuning, built 2026-09-26. Internal working document; not for submission.", colour=grey)
    b.p("Every number is read from a recorded artifact named in the table that shows it. Two artifacts were "
        "added this round, both computed from preserved outputs with no model re-run: `round3_analysis.json` "
        "and `segmentation_stats.json`. Test counts are run records and are labelled as such.", colour=grey)

    b.h("Summary", 1)
    leak = ctx["r3"]["holdout_leakage"]
    r3, kc, mm = ctx["r3"], ctx["r3"]["kappa_composition"], ctx["mars_eval"]["mars"]
    hp = r3["model_performance"]["holdout"]["paired_qwen3_14b_minus_qwen2_5_7b"]
    ep = r3["holdout_sensitivity"]["excluding_all_pool_episodes"]["paired_qwen3_14b_minus_qwen2_5_7b"]
    b.table("Each question and the answer in one line.", ["#", "Question", "Answer"], [
        ["1", f"Final merged κ = {kc['final_state_all_330']['mean']:.3f}", f"Mean of 4 per-field κ over all 330 in final pre-adjudication state (178 round-2 + 152 round-1). Not a before/after comparison: the retained 152 were at {kc['retained_152_round1_values']['mean']:.3f} already"],
        ["2", "Segmentation logic", "Rule-based 2–8 turn windows around child turns; all speakers; %spa/%act/%gpx only; 1-turn overlaps; 8 types; nothing excluded; deterministic"],
        ["3", "506 examined → 330", f"Per-child seeded shuffle (seed 42), walk in order, keep where all 3 generators admissible, stop at 30. **Found: {leak['n_in_pool']} of the {leak['n_holdout']} frozen evaluation segments are pool episodes** — conclusions unchanged on the {leak['n_holdout'] - leak['n_in_pool']}"],
        ["4", "Metric definitions", "Five judgements per field; four formulas with exact denominators; 'unknown' is a category; an omitted field is scored as null"],
        ["5", "Cluster CIs for models", f"Computed. qwen3:14b − qwen2.5:7b on the original frozen 132: merged {diff_str(hp['merged_agreement'])}; unsupported {diff_str(hp['unsupported_rate'])}. On the overlap-excluded 121: merged {diff_str(ep['merged_agreement'])}; unsupported {diff_str(ep['unsupported_rate'])}"],
        ["6", "Full-policy MARS", f"Exists: mars_evaluation.json. Coverage {pct(mm['coverage'])}; {pct(mm['accuracy_over_all'])} over all {mm['n_total']} (abstain = miss), {pct(mm['accuracy_over_returned'])} over the {mm['n_returned']} returned; qwen2.5:7b {pct(ctx['mars_eval']['best_generator_accuracy'])}"],
        ["7", "Model/annotator input", "Identical transcript content, 330/330; sheets additionally carried administrative IDs; model saw no IDs, goals or metadata; annotators saw no model information"],
        ["8", "Representation logic", "Always applied, not goal-relative, no exemption; single setting or role → never ADEQUATE; 'contributors' = roles"],
        ["9", "Authorization tests", "Test evidence for all five; 120 pass; 4 real-DB tests added (revocation, cross-child)"],
        ["10", "Bootstrap method", "Yes: child-cluster, 10,000 percentile. One manuscript sentence says 'segments' and needs fixing"],
    ], [0.35, 1.7, 4.65])
    b.label("Manuscript changes this round implies")
    b.bullets([
        f"κ = {kc['final_state_all_330']['mean']:.3f}: describe the mixed annotation state without causal "
        "'codebook effect' language; relabel the 'after codebook re-annotation' row (Q1).",
        f"Evaluation set: disclose the {leak['n_in_pool']} pool episodes; name the sets 'original frozen "
        f"evaluation set (n = {leak['n_holdout']})' and 'overlap-excluded evaluation set (n = "
        f"{leak['n_holdout'] - leak['n_in_pool']})', report both, with the {leak['n_holdout'] - leak['n_in_pool']} "
        "as primary; remove 'the entire original stratified pool … is excluded' (Q3).",
        "Metric definitions: omission is scored as null; 'unknown' is a category (Q4).",
        "Add child-cluster CIs to the model-performance tables (Q5).",
        "MARS: always pair whole-record agreement with coverage (Q6).",
        "Adequacy: representation is not goal-relative; 'contributors' are roles (Q8).",
        "Bootstrap sentence at make_manuscript.py:558 (Q10).",
    ])

    for sec, answer in zip(sections, ANSWERS):
        b.page_break()
        b.h(sec["heading"], 1)
        render_question(b, sec["items"])
        b.h("Response", 2)
        answer(b, sec, ctx)

    b.page_break()
    b.h("How this document was produced", 1)
    b.code(["uv run python -m tests.eval.extraction.round3_analysis      # no model calls",
            "uv run python -m tests.eval.extraction.segmentation_stats   # needs exp-data/ locally",
            "uv run python docs/paper/make_round3_response.py"])
    b.doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    print(build())
