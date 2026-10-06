"""Build the answers to the round-11 technical questions (R11-5, R11-6, R11-8).

Every request bullet in ``R11_Review_comments.docx`` is quoted in order and
answered in plain words, then as facts with their source. Numbers come from
``round11_analysis.json`` (no model calls), whose re-runs are asserted to
reproduce the recorded intervals. Facts only; no manuscript wording.

    uv run python -m tests.eval.extraction.round11_analysis
    uv run python docs/paper/make_round11_response.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

PAPER = Path(__file__).resolve().parent
sys.path.insert(0, str(PAPER))

import docx  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from make_round2_response import QUOTE, Doc  # noqa: E402

ROOT = PAPER.parents[1]
QUESTIONS = PAPER / "R11_Review_comments.docx"
OUT = PAPER / "Round11_Technical_Response.docx"
DATA = ROOT / "tests" / "eval" / "extraction" / "results" / "round11_analysis.json"
PLAIN_FILL, PLAIN_EDGE = "EAF7F1", "1BAF7A"


def read_bullets() -> dict[str, list[str]]:
    src = docx.Document(str(QUESTIONS))
    out: dict[str, list[str]] = {}
    cur = None
    for el in src.element.body.iterchildren():
        if el.tag.split("}")[1] != "p":
            continue
        p = Paragraph(el, src)
        t = p.text.strip()
        if t.startswith("R11-"):
            cur = t.split()[0]
            out[cur] = []
        elif t.startswith("Important constraints"):
            cur = None
        elif cur and p.style.name == "List Bullet":
            out[cur].append(t)
    return out


def ask(b: Doc, text: str) -> None:
    para = b.doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.keep_with_next = True
    r = para.add_run("• ")
    r.bold = True
    r.font.color.rgb = QUOTE
    b._runs(para, text, colour=QUOTE)


def plain(b: Doc, text: str) -> None:
    b.boxed("In plain words", [text], fill=PLAIN_FILL, edge=PLAIN_EDGE)


def r11_5(b: Doc, qs: list[str], d: dict) -> None:
    acc = d["r11_5_bootstrap_accounting"]
    rows, drops = [], []
    for name, fields in acc["kappa"].items():
        for f, v in fields.items():
            rows.append([f"{name} — {f}", str(v["requested"]), str(v["valid"]), str(v["dropped"]), f"{v['pct_dropped']:.2f}",
                         "; ".join(f"{n}× {r}" for r, n in v["reasons"].items()) or "—"])
            if v["dropped"]:
                drops.append(f"{name} — {f}")
    for f, v in acc["kappa_change"].items():
        rows.append([f"Change in kappa, round 1 → round 2 (178 re-issued) — {f}", str(v["requested"]), str(v["valid"]),
                     str(v["dropped"]), f"{v['pct_dropped']:.2f}", "; ".join(f"{n}× {r}" for r, n in v["reasons"].items()) or "—"])
    all_items = [v for fs in acc["kappa"].values() for v in fs.values()] + list(acc["kappa_change"].values())
    assert all(v["n_clusters"] == 11 for v in all_items)
    maxpct = max(v["pct_dropped"] for v in all_items)
    assert sum(v["dropped"] > 0 for v in all_items) == 4

    ask(b, qs[0])
    plain(b, "Every kappa interval with a child-based bootstrap that exists in the recorded results is listed below — 24 "
             "field intervals across six analyses, plus the four intervals on the change in kappa between rounds. The "
             "V11 manuscript and Supplementary Material 5 are not in the repository, so please match these rows to the "
             "intervals they actually print.")
    ask(b, qs[1])
    ask(b, qs[2])
    ask(b, qs[3])
    ask(b, qs[4])
    ask(b, qs[5])
    ask(b, qs[6])
    plain(b, f"Every interval asked for 10,000 replicates. Most used all 10,000. Four dropped a few (at most 3, i.e. "
             f"{maxpct:.2f}%), all in modality: in those replicates the drawn children happened to have only 'verbal' "
             "modality from both annotators, and kappa cannot be computed when everything is in one category.")
    b.table("Replicate accounting. 'Valid' was recomputed with the original inputs and seed (20260915) and matches the "
            "recorded value exactly (asserted). Source: `results/round11_analysis.json` (r11_5_bootstrap_accounting).",
            ["Analysis / field", "Requested", "Valid", "Dropped", "% dropped", "Reason"],
            rows, [2.4, 0.7, 0.6, 0.6, 0.6, 1.8])
    b.bullets([
        "Rule: `cohens_kappa()` (`tests/eval/extraction/agreement_report.py`) returns undefined when expected agreement "
        "≥ 1, i.e. every item in the resample has the same value from both annotators. Such replicates are skipped; the "
        "percentile interval is taken over the valid ones (`reliability_stats.bootstrap_kappa`, `delta_kappa`).",
        "For the change in kappa, a replicate is dropped if either round's kappa is undefined on it.",
        f"Intervals with dropped replicates: {', '.join(drops)}; and the modality change interval.",
    ])
    ask(b, qs[7])
    plain(b, "Confirmed: children were resampled (each drawn child brings all its segments), and every one of these "
             "analyses has 11 children.")
    b.bullets(["`unit = \"child\"` in every recorded interval; 11 distinct child labels in every analysis (asserted).",
               "Sources: reliability_round1.json, reliability_current.json, reliability_holdout.json, round6_analysis.json "
               "(121), round_comparison.json."])


def r11_6(b: Doc, qs: list[str], d: dict) -> None:
    v = d["r11_6_validator"]
    val, sc, res = v["validator"], v["scorer"], v["result_artifact"]
    assert all(h["diff_vs_previous"] in (None, "validator code identical (comment lines ignored)") for h in val["history"])
    assert sc["report_count_rule_at_run"] == sc["report_count_rule_at_creation"]
    assert res["at_run"]["prohibited_content_reports"] == 0 and res["at_run"]["n_llm_reports"] == 11
    run = next(h for h in res["history"] if h["commit"] == "9ecc330")

    b.label("A. Exact chronology")
    ask(b, qs[0])
    plain(b, f"The checker was written on 25 August 2026 (commit {val['created']['commit']}), three weeks before the reports.")
    b.bullets([f"`validate_report()` in `{val['file']}` — first commit {val['created']['commit']}, {val['created']['date']}. "
               f"Version used for the run: commit 9ecc330 ({run['date']}), code identical to the first commit."])
    ask(b, qs[1])
    plain(b, f"The number/reference scorer was committed on 16 September 2026 ({sc['history'][0]['commit']}).")
    b.bullets([f"`{sc['file']}` (`score_report()`, `grounded_numbers()`, `is_grounded()`) — first commit "
               f"{sc['history'][0]['commit']}, {sc['history'][0]['date']}; version used for the run 9ecc330."])
    ask(b, qs[2])
    plain(b, "The 11 model-written reports' pass/fail results were first saved on 17 September 2026 (9ecc330). An earlier "
             "version of the same file (16 September) contained no model-written reports.")
    b.table("Every commit of the result artefact.", ["Commit", "Date", "Content"],
            [[h["commit"], h["date"], "11 model-generated reports (the recorded run)" if h["commit"] == "9ecc330"
              else "0 model-generated reports; 11 template fallbacks"] for h in res["history"]], [1.0, 2.2, 3.5])
    b.p(f"File: `{res['file']}`.")
    ask(b, qs[3])
    plain(b, "No change. The checking code is the same in every commit from its creation to the run (and after). The "
             "line that turns errors into pass/fail, and the line that counts reports with prohibited content, are also "
             "unchanged.")
    b.table("validate_report code fingerprint per commit (comment lines ignored).", ["Commit", "Date", "Fingerprint", "vs previous"],
            [[h["commit"], h["date"][:10], h["validator_code_sha"], h["diff_vs_previous"] or "first version"] for h in val["history"]],
            [0.9, 1.1, 1.4, 3.3])
    b.bullets([
        "`validation_passed = len(validation_errors) == 0` (in `ReportingClient.generate_report`) is present in every "
        "commit up to the run.",
        "Scorer count of reports with prohibited content, identical at creation (7a4b725) and at the run (9ecc330): "
        "`sum(1 for e in llm_reports if any(\"Prohibited\" in err for err in e[\"validation_errors\"]))`.",
    ])

    b.label("B. Prohibited-clinical-content check")
    ask(b, qs[4])
    plain(b, "One function in one file does it, and the result is counted by the test script.")
    b.bullets([f"`{val['file']}` → `validate_report(sections, evidence_packet)`, called from "
               "`ReportingClient.generate_report()`; counted in `tests/eval/reporting/groundedness.py` (`run()` summary)."])
    ask(b, qs[5])
    plain(b, "It is a fixed list of six search patterns written into the code.")
    b.code([f"PROHIBITED_PATTERNS  (reporting_client.py @ 9ecc330)"] + [f"  {i + 1}. {p}" for i, p in enumerate(val["prohibited_patterns_at_run"])])
    ask(b, qs[6])
    plain(b, "Each pattern is searched for anywhere in the text, ignoring upper/lower case. Words must be whole words. In "
             "the two-part patterns (e.g. 'recommend … therapy') both parts must be on the same line of text.")
    b.bullets(["`re.findall(pattern, full_text, re.IGNORECASE)` for each pattern: case-insensitive regular-expression "
               "search, unanchored; `\\b` word boundaries; `.*` without `re.DOTALL`, so it does not cross a newline."])
    ask(b, qs[7])
    plain(b, "Only the main text of each section. Titles, the optional abstention note and the record-ID list were not looked at.")
    b.bullets(["`full_text = \" \".join(s.body for s in sections)`. Not inspected: `heading`, `abstention_note`, `evidence_refs`."])
    ask(b, qs[8])
    plain(b, "Pass = none of the six patterns found. Fail = any pattern found at least once (one error per pattern found). "
             "The reported result is the number of reports with at least one such error: 0 of 11.")
    b.bullets(["Fail: for each pattern with ≥1 match, append \"Prohibited content detected: pattern '<pattern>' matched\".",
               "Scorer: a report counts as containing prohibited content if any of its validation errors contains "
               "\"Prohibited\". Recorded: `prohibited_content_reports = 0` of `n_llm_reports = 11`."])
    ask(b, qs[9])
    plain(b, "None, apart from putting the section texts together with a space between them. Upper/lower case is handled "
             "by the search setting, not by changing the text.")
    b.bullets(["No lowercasing, punctuation stripping or whitespace normalisation of the text; case-insensitivity comes "
               "from `re.IGNORECASE`."])
    ask(b, qs[10])
    plain(b, "No configuration file is involved.")
    b.bullets([v["config_dependency"] + "."])
    b.boxed("Limit", ["The 11 report texts were not saved (the result file stores scores only; "
                      f"report_text_preserved = {res['at_run']['report_text_preserved']}). The check can be described "
                      "exactly but cannot be re-executed on those reports."])


def r11_8(b: Doc, qs: list[str], d: dict) -> None:
    u = d["r11_8_unaccounted_mass"]
    prim = u["episode_type_7/intended"]
    ask(b, qs[0])
    plain(b, f"{prim['overall_unaccounted_pct']:.1f}% of the target frame (the reported analysis: seven episode types, "
             "intended frame).")
    b.bullets([f"{prim['total_unaccounted']} of {prim['total_frame_episodes']} target-frame episodes = "
               f"{prim['overall_unaccounted_pct']:.2f}%."])
    ask(b, qs[1])
    ask(b, qs[2])
    plain(b, "Six children have some type in their target frame that none of their sampled episodes has; five have none. "
             "The largest is mars (41%: protest, response and play not sampled).")
    b.table("Reported analysis (7 episode types, intended frame). Source: round11_analysis.json (r11_8), reproduces "
            "round4_analysis.json unaccounted_intended_share_by_child (asserted).",
            ["Child", "Unaccounted mass (%)", "Missing episode type(s)", "Target-frame share of each / note"],
            [[c, f"{x['unaccounted_pct']:.2f}", ", ".join(x["missing_types"]) or "—",
              "; ".join(f"{t} {m['share_of_child_frame_pct']:.2f}% ({m['frame_episodes']}/{x['frame_episodes']})"
                        for t, m in x["missing_types"].items()) or f"all frame types sampled ({x['frame_episodes']} frame episodes)"]
             for c, x in prim["per_child"].items()], [1.2, 1.0, 1.7, 2.8])
    ask(b, qs[3])
    plain(b, "Confirmed: missing types were simply left out. Nothing was filled in or moved to another type.")
    b.bullets([u["handling"]])
    ask(b, qs[4])
    plain(b, "The overall figure divides all unaccounted target-frame episodes by all target-frame episodes, pooled over "
             "the 11 children.")
    b.bullets([f"Denominator = Σ over 11 children of target-frame episodes = {prim['total_frame_episodes']} (intended "
               f"frame). The unweighted mean of the per-child percentages would be {prim['mean_of_child_pct']:.2f}%."])
    ask(b, qs[5])
    plain(b, "The re-weighting was run four ways. The reported sensitivity result is the 7-type, intended-frame one (it is "
             "the 'intended' side of the reported frame effect); the others are given for completeness.")
    b.table("All four re-weighting variants in round4_analysis.json (poststratified_121).",
            ["Variant", "Target frame", "Unaccounted", "Overall %", "Children with non-zero mass"],
            [[k.split("/")[0], k.split("/")[1], f"{v['total_unaccounted']}/{v['total_frame_episodes']}", f"{v['overall_unaccounted_pct']:.2f}",
              str(sum(x["unaccounted_pct"] > 0 for x in v["per_child"].values()))] for k, v in u.items() if isinstance(v, dict)],
            [1.4, 1.1, 1.2, 0.9, 1.6])
    b.bullets(["The reported frame effect is (estimate weighted to the intended frame) − (estimate weighted to the analysed "
               "frame), 7 episode types (`frame_effect` under `episode_type_7`), as in the Round-4 response (R2-2)."])


def build() -> Path:
    d = json.loads(DATA.read_text())
    qs = read_bullets()
    assert [len(qs[k]) for k in ("R11-5", "R11-6", "R11-8")] == [8, 11, 6], {k: len(v) for k, v in qs.items()}
    b = Doc()
    grey = RGBColor(0x52, 0x51, 0x4E)
    b.doc.add_heading("Round 11 — Answers to the Technical Questions", level=0)
    b.p("Every request in 'R11_Review_comments.docx' (R11-5, R11-6, R11-8), in order, answered in plain words and then as "
        "facts with their source. No model was run and no historical artefact was changed. Built 2026-09-29. Internal "
        "working document; facts only.", colour=grey)
    b.h("At a glance", 1)
    acc = d["r11_5_bootstrap_accounting"]
    n_int = sum(len(v) for v in acc["kappa"].values()) + len(acc["kappa_change"])
    b.bullets([
        f"**R11-5:** {n_int} child-cluster kappa intervals, all 10,000 requested, 11 children each. Dropped replicates: "
        "0 in most; 1–3 in four modality intervals (single-category resamples).",
        "**R11-6:** validator 75eb1d5 (25 Aug), scorer 7a4b725 (16 Sep), recorded run 9ecc330 (17 Sep); no logic change "
        "in between. Six case-insensitive regex patterns over section body text only.",
        f"**R11-8:** unaccounted target-frame mass {d['r11_8_unaccounted_mass']['episode_type_7/intended']['overall_unaccounted_pct']:.1f}% "
        "overall (160/2,051); mars 41.2%, carl 21.3%, sid 17.4%; not imputed.",
    ])
    for key, fn, title in (("R11-5", r11_5, "R11-5 – Bootstrap replicate accounting"),
                           ("R11-6", r11_6, "R11-6 – Historical reporting-validator reproducibility"),
                           ("R11-8", r11_8, "R11-8 – Unaccounted episode-type mass in reweighting sensitivity")):
        b.page_break()
        b.h(title, 1)
        fn(b, qs[key], d)
    b.h("How this document was produced", 1)
    b.code(["uv run python -m tests.eval.extraction.round11_analysis   # no model calls",
            "uv run python docs/paper/make_round11_response.py"])
    b.doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    print(build())
