"""The manuscript must not contradict itself about what was completed.

A document assembled in sections drifts: a Methods paragraph written while an
experiment was pending survives into a build where the Results report that same
experiment as finished. That happened here -- Methods carried "annotation and
generator runs over it are not complete" in a build whose Results reported
mean kappa 0.910 over exactly those segments -- and a reviewer noticed before
any test did.

These checks derive the true state from the artefacts and assert the prose
agrees with it, across Methods, Results, Limitations and Not-measured together
rather than section by section.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[3]
RESULTS = Path(__file__).parent / "results"
MANUSCRIPT = ROOT / "docs/paper/Manuscript_Working_Draft.docx"


@pytest.fixture(scope="module")
def prose() -> str:
    if not MANUSCRIPT.exists():
        pytest.skip("manuscript not built")
    from docx import Document

    doc = Document(str(MANUSCRIPT))
    text = " ".join(p.text for p in doc.paragraphs)
    text += " " + " ".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    return re.sub(r"\s+", " ", text)


@pytest.fixture(scope="module")
def state() -> dict:
    """What the artefacts say actually happened."""
    holdout_batch = RESULTS / "holdout_batch.json"
    batch = json.loads(holdout_batch.read_text()) if holdout_batch.exists() else None
    generators_complete = []
    generators_partial = []
    if batch:
        total = len(batch["candidates"])
        for model in batch.get("models", []):
            n = sum(1 for e in batch["candidates"] if model in e["candidates"])
            (generators_complete if n == total else generators_partial).append(model)
    return {
        "holdout_frozen": (RESULTS / "holdout_set.json").exists(),
        "holdout_annotated": (RESULTS / "reliability_holdout.json").exists(),
        "holdout_gold": (RESULTS / "gold_holdout_labels.json").exists(),
        "generators_complete": generators_complete,
        "generators_partial": generators_partial,
        "main_gold": (RESULTS / "gold_labels.json").exists(),
    }


class TestHeldOutConsistency:
    def test_annotated_holdout_is_not_described_as_unannotated(
        self, prose, state
    ) -> None:
        if not state["holdout_annotated"]:
            pytest.skip("held-out round not annotated")
        for stale in (
            "annotation and generator runs over it are not complete",
            "held-out set is frozen (132 segments) but not yet annotated",
            "is not yet annotated",
        ):
            assert stale not in prose, f"stale claim survives: {stale!r}"

    def test_holdout_reliability_is_reported_when_measured(
        self, prose, state
    ) -> None:
        if not state["holdout_annotated"]:
            pytest.skip("held-out round not annotated")
        report = json.loads((RESULTS / "reliability_holdout.json").read_text())
        assert f"{report['mean_kappa']:.3f}" in prose

    def test_holdout_accuracy_is_not_claimed_without_gold(
        self, prose, state
    ) -> None:
        # Accuracy over the held-out sample requires an adjudicated reference.
        if state["holdout_gold"]:
            pytest.skip("held-out gold exists; accuracy may be reported")
        for forbidden in (
            "held-out accuracy",
            "accuracy on the held-out sample was",
            "unsupported-assertion rate on the held-out",
        ):
            assert forbidden.lower() not in prose.lower(), forbidden

    def test_generator_completion_is_described_accurately(
        self, prose, state
    ) -> None:
        if not state["generators_partial"]:
            pytest.skip("all held-out generator runs complete")
        # A partial run must be visible as partial somewhere in the prose.
        assert "in progress" in prose or "not complete" in prose

    def test_the_pending_marker_does_not_contradict_a_finished_experiment(
        self, prose, state
    ) -> None:
        if not state["holdout_annotated"]:
            pytest.skip("held-out round not annotated")
        # A [PENDING] tag adjacent to held-out text is the exact drift that
        # produced the contradiction a reviewer caught.
        for match in re.finditer(r"\[PENDING\][^.]*\.", prose):
            assert "held-out" not in match.group(0).lower(), match.group(0)


class TestClaimScope:
    def test_no_general_claim_about_model_size(self, prose) -> None:
        # Four local models, none at frontier scale, cannot support a general
        # claim that size does not explain the failure.
        for overreach in (
            "not attributable to model size",
            "not a consequence of model size",
            "model size does not",
        ):
            assert overreach not in prose.lower(), overreach

    def test_the_scoped_replacement_is_present(self, prose) -> None:
        assert "within the locally served models evaluated" in prose

    def test_no_causal_claim_from_the_uncontrolled_holdout_round(
        self, prose, state
    ) -> None:
        if not state["holdout_annotated"]:
            pytest.skip("held-out round not annotated")
        # The held-out round has no untreated control; its kappa must not be
        # presented as evidence the codebook caused the improvement.
        assert "no internal control" in prose
        assert "causal evidence for the codebook remains" in prose


class TestSubmissionGate:
    """The draft must not claim readiness it does not have."""

    def test_incomplete_sections_are_flagged_while_incomplete(self, prose) -> None:
        # References, prior-work comparison and the IRB determination are not
        # written; the draft must say so rather than read as finished.
        assert "TO COMPLETE BEFORE SUBMISSION" in prose

    def test_the_draft_declares_itself_not_submission_ready(self, prose) -> None:
        assert "NOT SUBMISSION READY" in prose


class TestDeploymentScope:
    """MARS must not be describable as part of the running application.

    The application wires a single extraction provider and the reranking
    provider is not reachable from its factory. A manuscript that reads as
    though reranking runs in the request path misrepresents the system, and
    this is the correction a reviewer flagged as the most serious one.
    """

    @staticmethod
    def _deployed_provider() -> str:
        from src.backend.config import settings

        return str(settings.extraction_provider_kwargs.get("model", ""))

    def test_the_application_still_uses_a_single_provider(self) -> None:
        # If this ever becomes a multi-generator roster, the prose below has
        # to change too -- so fail loudly rather than let the two drift.
        from src.backend.config import settings

        kwargs = settings.extraction_provider_kwargs
        assert "model" in kwargs, kwargs
        assert isinstance(kwargs["model"], str), "expected one model, not a roster"

    def test_reranking_is_not_reachable_from_the_provider_factory(self) -> None:
        import src.backend.ml.extraction as extraction

        source = Path(extraction.__file__).read_text()
        assert "mars" not in source.lower(), (
            "the factory now exposes reranking; the manuscript's scope "
            "statement must be revisited"
        )

    def test_the_manuscript_states_the_offline_scope(self, prose) -> None:
        assert "MARS is an evaluation method, not a deployed component" in prose
        assert "not intended for the application's request path" in prose

    def test_the_manuscript_names_the_deployed_provider(self, prose) -> None:
        assert self._deployed_provider() in prose

    def test_no_present_tense_deployment_phrasing(self, prose) -> None:
        # "MARS selects/returns/runs" reads as live behaviour; results are
        # reported in the past tense.
        assert not re.search(r"MARS (selects|returns|runs) ", prose)

    def test_limitations_records_the_scope_boundary(self, prose) -> None:
        assert (
            "offline evaluation method, not a deployed component, and is not "
            "intended to become one" in prose
        )

    def test_no_four_generator_reranking_is_claimed(self, prose) -> None:
        assert "no four-generator reranking run has been performed" in prose


class TestContributionHierarchy:
    """One primary contribution, stated in the places a reader looks for it.

    The work spans annotation method, extraction evaluation, reranking,
    reporting governance and an application, and an editor reading it as five
    co-equal contributions would call it a kitchen-sink paper. The hierarchy
    that resolves that -- the evaluation framework is primary, everything else
    is a component it measures -- has to survive edits to any one section.
    """

    def test_the_primary_contribution_is_named(self, prose) -> None:
        assert "The primary contribution is the evaluation framework" in prose

    def test_the_other_components_are_subordinated(self, prose) -> None:
        assert "objects measured, not parallel contributions" in prose

    def test_the_title_names_the_framework(self, prose) -> None:
        assert "Evaluation Framework" in prose

    def test_the_abstract_conclusion_leads_with_the_framework(
        self, prose
    ) -> None:
        assert "Conclusions. The contribution is the evaluation framework" in prose

    def test_the_conclusions_restate_the_hierarchy(self, prose) -> None:
        assert "The framework is the contribution; the pipeline is" in prose


class TestGovernanceBoundaries:
    """The two boundaries the results are used to justify.

    Both are claims about the running application, so both must be stated in
    the manuscript and hold in the code.
    """

    def test_the_verification_gate_is_stated(self, prose) -> None:
        assert "mandatory gate rather than an available action" in prose

    def test_the_gate_is_enforced_in_the_analytics_query(self) -> None:
        from src.backend.models.observation import RecordStatus
        from src.backend.services import analytics_service

        source = Path(analytics_service.__file__).read_text()
        assert "RecordStatus.VERIFIED, RecordStatus.CORRECTED" in source
        assert RecordStatus.VERIFIED.value == "verified"

    def test_generative_reporting_is_declared_experimental(self, prose) -> None:
        assert "not the operational reporting path" in prose

    def test_deterministic_reporting_is_the_configured_default(self) -> None:
        from src.backend.config import settings

        assert settings.reporting_mode == "deterministic"


class TestScopeOfTheClaims:
    """Four places the paper could overclaim, and the sentence that stops it.

    Each corresponds to a reviewer objection that is fatal if the manuscript
    reads as though the thing were established: extraction validated on real
    intervention notes, the whole extractor evaluated, the adequacy
    thresholds clinically grounded, or the workflow shown to help anyone.
    """

    def test_the_corpus_is_framed_as_a_proxy(self, prose) -> None:
        assert "proxy for the application's intended input, not a sample" in prose
        assert "component evaluation on proxy data" in prose

    def test_no_claim_of_validation_on_intervention_notes(self, prose) -> None:
        assert "Extraction is not validated on real intervention notes" in prose

    def test_the_evaluated_field_count_is_stated(self, prose) -> None:
        assert "asked for eight fields" in prose
        assert "extractor as a whole is therefore not evaluated" in prose

    def test_the_thresholds_are_not_called_clinical_cutoffs(
        self, prose
    ) -> None:
        assert (
            "configurable prototype governance thresholds, not validated "
            "clinical cutoffs" in prose
        )

    def test_benefit_claims_are_marked_undemonstrated(self, prose) -> None:
        assert "intended or potential, never demonstrated" in prose

    def test_no_bare_benefit_claim_survives_anywhere(self, prose) -> None:
        # A claim that the system reduces burden or saves time would need a
        # usability study that has not been run.
        for phrase in ("reduces review burden for", "saves time", "improves efficiency"):
            assert phrase not in prose, phrase


class TestAnnotatorClaims:
    """The Annotators section makes one factual claim; the rest is a stub.

    Full independent coverage of both batches is the part a reviewer can check
    against the artefacts, so it is checked here rather than trusted.
    """

    @staticmethod
    def _count(name: str) -> int:
        path = RESULTS / name
        if not path.exists():
            pytest.skip(f"{name} not present")
        return len(json.loads(path.read_text()))

    def test_both_annotators_covered_the_whole_main_batch(self, prose) -> None:
        a = self._count("gold_annotator_a.json")
        b = self._count("gold_annotator_b.json")
        assert a == b, (a, b)
        assert f"all {a} segments of the main batch" in prose

    def test_both_annotators_covered_the_whole_holdout_set(self, prose) -> None:
        a = self._count("gold_holdout_annotator_a.json")
        b = self._count("gold_holdout_annotator_b.json")
        assert a == b, (a, b)
        assert f"all {a} segments of the held-out set" in prose

    def test_the_qualifications_gap_is_flagged_not_invented(
        self, prose
    ) -> None:
        # Nothing about the annotators' credentials is known to this
        # repository, so the manuscript must ask for it rather than assert it.
        assert "professional role and qualification" in prose


class TestHeldOutResultsAreReported:
    """The successor to `test_holdout_accuracy_is_not_claimed_without_gold`.

    That test forbade claiming held-out accuracy before a reference standard
    existed, and retires itself the moment one does. The opposite failure then
    becomes possible: gold is adjudicated, the numbers exist in the artefacts,
    and the prose still says the round is pending. These check the second.
    """

    @pytest.fixture(scope="class")
    @classmethod
    def metrics(cls) -> dict:
        path = RESULTS / "holdout_metrics.json"
        if not path.exists():
            pytest.skip("held-out metrics not computed")
        return json.loads(path.read_text())

    def test_no_pending_language_survives(self, prose, state) -> None:
        if not state["holdout_gold"]:
            pytest.skip("held-out gold not adjudicated")
        for stale in (
            "pending adjudication",
            "awaits adjudication",
            "awaiting adjudication",
            "require the adjudicated reference standard",
        ):
            assert stale not in prose.lower(), stale

    def test_every_generator_delta_is_reported(self, prose, metrics) -> None:
        for model in metrics["models_compared"]:
            delta = metrics["delta"][model]["mean_field_accuracy"]
            for value in (delta["batch"], delta["holdout"]):
                assert f"{100 * value:.1f}%" in prose, (model, value)

    def test_the_selection_effect_direction_is_stated_correctly(
        self, prose, metrics
    ) -> None:
        # The claim in the prose is that every generator is *less* accurate on
        # the unselected sample. If a future re-run reverses that for any
        # model, the sentence becomes false and must be rewritten.
        deltas = [
            metrics["delta"][m]["mean_field_accuracy"]["delta"]
            for m in metrics["models_compared"]
        ]
        if all(d < 0 for d in deltas):
            assert "Every generator is less accurate" in prose
        else:
            pytest.fail(
                f"deltas are no longer uniformly negative ({deltas}); the "
                "manuscript's selection-effect sentence needs rewriting"
            )

    def test_both_sides_used_the_same_scoring_policy(self, metrics) -> None:
        # A comparison across different behavior-matching rules would compare
        # different quantities; the number would look fine and mean nothing.
        batch = json.loads((RESULTS / "scored_against_gold_augmented.json").read_text())
        holdout = json.loads((RESULTS / "scored_holdout.json").read_text())
        for key in ("scoring_fields", "behavior_match", "behavior_overlap_threshold"):
            assert batch[key] == holdout[key], key
        assert metrics["scoring_policy"]["scored_fields"] == batch["scoring_fields"]

    def test_the_holdout_gold_covers_every_double_annotated_segment(
        self, metrics
    ) -> None:
        gold = json.loads((RESULTS / "gold_holdout_labels.json").read_text())
        # Nothing was lost to an undecided row.
        assert gold["n_segments_in_gold"] == gold["n_segments_double_annotated"]
        assert metrics["holdout"]["n_segments"] == gold["n_segments_in_gold"]


class TestAgreementFraming:
    """Agreement, reported split, with the human level beside it.

    A single per-field percentage merged two abilities and hid that on
    `modality` the generators reach what the annotators reached. The split is
    now the primary reporting, so these pin both that the framing survives and
    that the numbers in the prose are the numbers in the artefacts.
    """

    @pytest.fixture(scope="class")
    @classmethod
    def split(cls) -> dict:
        path = RESULTS / "holdout_metrics.json"
        if not path.exists():
            pytest.skip("metrics not computed")
        return json.loads(path.read_text())["batch"]["per_model"]

    def test_results_are_framed_as_agreement_not_accuracy(self, prose) -> None:
        assert "Agreement with the reference standard" in prose
        assert "agreement names what was measured" in prose

    def test_the_two_cases_are_reported_separately(self, prose) -> None:
        assert "Agreement where the reference standard records a value" in prose
        assert "Agreement where the reference standard records no value" in prose

    def test_no_merged_per_field_figure_is_reported(self, prose) -> None:
        # The merged per-field percentage averaged the two abilities and
        # reported the average as one, so it is not reported at all. The mean
        # over the four fields survives only where the held-out comparison
        # needs a single scale, and is labelled as merged there.
        assert "Merged per-field figure" not in prose
        assert "No merged per-field figure is reported" in prose
        assert "Merged mean field accuracy" in prose

    def test_the_modality_result_matches_the_artefact(self, prose, split) -> None:
        best = "qwen2.5:7b"
        value = split[best]["agreement_where_value_recorded"]["modality"]["rate"]
        silence = split[best]["agreement_where_no_value_recorded"]["modality"]["rate"]
        assert f"{100 * value:.1f}%" in prose
        assert f"{100 * silence:.1f}%" in prose

    def test_the_human_reference_level_is_stated(self, prose) -> None:
        r1 = json.loads((RESULTS / "reliability_round1.json").read_text())
        rates = [
            v["observed_agreement"]
            for v in r1["fields"].values()
            if isinstance(v, dict) and "observed_agreement" in v
        ]
        lo, hi = round(100 * min(rates)), round(100 * max(rates))
        assert f"{lo}% to {hi}%" in prose, (lo, hi)

    def test_the_comparison_is_not_claimed_to_be_like_for_like(
        self, prose
    ) -> None:
        # Annotators were compared with each other; generators with the settled
        # value. Presenting those as one scale would overstate the comparison.
        assert "not a like-for-like comparison" in prose
