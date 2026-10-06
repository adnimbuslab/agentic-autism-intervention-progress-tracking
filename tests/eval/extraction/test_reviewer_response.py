"""The reviewer response must answer the review it claims to answer.

Two failure modes are worth a test here, and neither is caught by reading the
document. The first is drift between the response and the review: a comment
paraphrased, reordered or dropped leaves a document that answers a review
nobody wrote. The second is drift between the response and the data: the
response quotes held-out figures, governance counts and an adjudication split,
and a re-scored run would leave those numbers stale with nothing failing.

Both are checked against sources rather than against a copy: the reviewer's own
.docx, and the artefacts.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[3]
RESULTS = Path(__file__).parent / "results"
REPORTING = ROOT / "tests/eval/reporting/results"
REVIEW = ROOT / "docs/paper/JMIR_Strict_Reviewer_Critical_Issues_Prioritized.docx"
RESPONSE = ROOT / "docs/paper/Reviewer_Response.docx"


def _flat(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


@pytest.fixture(scope="module")
def review() -> dict:
    if not REVIEW.exists():
        pytest.skip("review document not present")
    from docx import Document

    doc = Document(str(REVIEW))
    issues: dict[int, dict] = {}
    current = None
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name == "Heading 2":
            number = re.match(r"^(\d+)\.\s*(.*)$", text)
            current = int(number.group(1)) if number else None
            if current is not None:
                issues[current] = {"title": number.group(2), "paragraphs": []}
        elif current is not None and para.style.name == "Normal":
            issues[current]["paragraphs"].append(text)
    return issues


@pytest.fixture(scope="module")
def holdout() -> dict:
    path = RESULTS / "holdout_metrics.json"
    if not path.exists():
        pytest.skip("holdout metrics not computed")
    return json.loads(path.read_text())


@pytest.fixture(scope="module")
def response() -> dict:
    if not RESPONSE.exists():
        pytest.skip("response not built")
    from docx import Document

    doc = Document(str(RESPONSE))
    headings = [
        p.text.strip() for p in doc.paragraphs
        if p.style.name.startswith("Heading")
    ]
    body = _flat(" ".join(p.text for p in doc.paragraphs))
    cells = _flat(" ".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    ))
    return {
        "headings": headings,
        "body": body,
        "cells": cells,
        "all": f"{body} {cells}",
        "images": sum(
            1 for rel in doc.part.rels.values() if "image" in rel.reltype
        ),
        "tables": len(doc.tables),
    }


def _section(response: dict, number: int) -> str:
    """The response text belonging to one issue, up to the next heading."""
    body = response["body"]
    start = body.index(f"Issue {number}.")
    following = [
        body.index(h, start + 1)
        for h in response["headings"]
        if h.startswith(("Issue ", "CRITICAL", "HIGH", "MODERATE", "What I would",
                         "Strict Reviewer"))
        and h not in (f"Issue {number}.",) and body.find(h, start + 1) > start
    ]
    return body[start:min(following)] if following else body[start:]


class TestEveryCommentIsAnswered:
    def test_every_numbered_issue_has_a_section(self, review, response) -> None:
        for number in review:
            assert any(h.startswith(f"Issue {number}.") for h in response["headings"]), (
                f"issue {number} unanswered"
            )

    def test_the_order_is_the_reviewers_order(self, review, response) -> None:
        numbers = [
            int(re.match(r"Issue (\d+)\.", h).group(1))
            for h in response["headings"] if h.startswith("Issue ")
        ]
        assert numbers == sorted(review)

    def test_each_title_is_the_reviewers_title(self, review, response) -> None:
        for number, issue in review.items():
            heading = next(
                h for h in response["headings"] if h.startswith(f"Issue {number}.")
            )
            assert heading == f"Issue {number}. {issue['title']}"

    @pytest.mark.parametrize("number", range(1, 13))
    def test_the_comment_is_quoted_verbatim(self, review, response, number) -> None:
        """Paraphrase is the failure this guards against.

        A response that restates the objection in its own words is answering
        something the reviewer did not write, and the substitution is invisible
        once the review document is closed.
        """
        if number not in review:
            pytest.skip(f"issue {number} absent from the review")
        for paragraph in review[number]["paragraphs"]:
            assert _flat(paragraph) in response["all"], f"issue {number} paraphrased"

    @pytest.mark.parametrize("number", range(1, 13))
    def test_each_response_carries_the_required_parts(self, response, number) -> None:
        section = _section(response, number)
        assert "Status:" in section, number
        assert "Response" in section, number
        assert "Limitations that remain" in section, number


class TestNumbersMatchTheArtefacts:
    def test_every_held_out_figure_is_the_artefacts_figure(
        self, response, holdout
    ) -> None:
        for model in holdout["models_compared"]:
            d = holdout["delta"][model]["mean_field_accuracy"]
            for side in ("batch", "holdout"):
                assert f"{100 * d[side]:.1f}%" in response["all"], f"{model} {side}"
            assert f"{100 * d['delta']:+.1f}" in response["all"], model

    def test_the_governance_counts_are_the_artefacts_counts(self, response) -> None:
        path = REPORTING / "reporting_groundedness.json"
        if not path.exists():
            pytest.skip("groundedness run not present")
        s = json.loads(path.read_text())["summary"]
        assert f"{s['total_numbers_ungrounded']} of {s['total_numbers_asserted']}" \
            in response["all"]
        assert f"{100 * s['pooled_ungrounded_rate']:.1f}%" in response["all"]

    def test_the_adjudication_split_is_recomputed_not_quoted(self, response) -> None:
        """The split that falsifies 'neither annotator dominated' on this round."""
        import sys

        sys.path.insert(0, str(ROOT / "docs/paper"))
        from make_reviewer_response import adjudication_balance

        balance = adjudication_balance()
        if balance is None:
            pytest.skip("held-out adjudication artefacts not present")
        assert f"{balance['a']} to {balance['b']}" in response["all"]
        assert f"{balance['p']:.3f}" in response["all"]

    def test_the_thresholds_are_the_configured_thresholds(self, response) -> None:
        from src.backend.config import settings

        assert str(settings.adequacy_min_observations) in response["cells"]
        assert f"{settings.adequacy_field_completeness_threshold:.2f}" \
            in response["cells"]
        assert f"{settings.adequacy_temporal_cluster_max_ratio:.2f}" \
            in response["cells"]


class TestTheDocumentIsComplete:
    def test_no_figure_failed_to_embed(self, response) -> None:
        assert "[missing figure" not in response["all"]
        assert response["images"] >= 5

    def test_tables_and_diagrams_are_present(self, response) -> None:
        assert response["tables"] >= 20

    def test_markup_was_rendered_not_printed(self, response) -> None:
        """Literal ** or ` in the output means a run was never converted."""
        assert "**" not in response["all"]
        assert "`" not in response["all"]

    def test_open_items_are_not_reported_as_progress(self, response) -> None:
        for phrase in ("IRB determination", "Annotator credentials", "Citation audit"):
            assert phrase in response["cells"], phrase
        assert "What remains open before submission" in response["headings"]
